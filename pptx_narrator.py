#!/usr/bin/env python3
"""PPTX-Narrator: automated narration of PowerPoint presenter notes.

Pipeline: note extraction -> technical-term scanning -> dictionary / SI-unit
normalization -> (optional) machine translation -> voice-cloned TTS
(GPT-SoVITS or Qwen3-TTS) -> ASR round-trip verification -> PPTX repackaging.
"""
import os
import re
import unicodedata
from collections import Counter
import difflib
import hashlib
import shlex
import traceback
import sys
import time
import shutil
import logging
import argparse
import tempfile
import zipfile
import csv
import json
import requests
from pptx import Presentation
from pydub import AudioSegment
import nltk
from pydub.effects import compress_dynamic_range, normalize

__version__ = "1.0.0"


def _ensure_nltk_data():
    """Download the NLTK corpora used by --scan on first use (not at import time)."""
    for corpus in ("stopwords", "words"):
        try:
            nltk.data.find(f"corpora/{corpus}")
        except LookupError:
            nltk.download(corpus, quiet=True)
            try:
                nltk.data.find(f"corpora/{corpus}")
            except LookupError:
                logger.error(
                    f"NLTK corpus '{corpus}' is not available and could not be downloaded. "
                    "Install it manually: python -m nltk.downloader stopwords words"
                )
                sys.exit(1)

# The weight paths are sent to the GPT-SoVITS API server and resolved there, relative to the
# server's own directory, not to this script.
MODELS_CONFIG = {
    "v2ProPlus": {
        "gpt": "./GPT_SoVITS/pretrained_models/s1v3.ckpt",
        "sovits": "./GPT_SoVITS/pretrained_models/v2Pro/s2Gv2ProPlus.pth"
    },
    "v4": {
        "gpt": "./GPT_SoVITS/pretrained_models/s1v3.ckpt",
        "sovits": "./GPT_SoVITS/pretrained_models/v4/s2Gv4.pth"
    },
    "v1_clear": {
        "gpt": "./GPT_SoVITS/pretrained_models/s1bert25hz-2kh-longer-epoch=68e-step=50232.ckpt",
        "sovits": "./GPT_SoVITS/pretrained_models/s2G488k.pth"
    }
}

# ==========================================
# Configuration & Constants
# ==========================================
# Workspace file names carry the language of the text: slide_N_<lang>.txt, with <lang>
# a language code (slide_3_ja.txt, slide_3_en.txt, slide_3_zh-CN.txt). Files
# written by pre-release versions (slide_N.txt for Japanese, slide_N_eng.txt for English)
# are still recognized when a workspace is read.
CJK_LANGS = {"ja", "zh", "yue", "ko"}

# Languages accepted by the TTS engines (ISO 639-1 base code -> value passed to the engine)
QWEN3_LANGUAGES = {
    "zh": "Chinese", "en": "English", "ja": "Japanese", "ko": "Korean", "de": "German",
    "fr": "French", "ru": "Russian", "pt": "Portuguese", "es": "Spanish", "it": "Italian",
}
GPT_SOVITS_LANGUAGES = {"zh": "zh", "en": "en", "ja": "ja", "ko": "ko", "yue": "yue"}

_TEXT_FILE_RE = re.compile(r"^slide_(\d+)(?:_([A-Za-z]{2,3}(?:-[A-Za-z0-9]{2,4})?))?\.txt$")


def normalize_lang(code):
    """Normalize a language code to the spelling used in file names (lower-case, with the
    region in capitals; the codes of earlier versions, such as 'iw' for Hebrew, are kept so
    that their workspaces are still read).

    'JA' -> 'ja', 'zh_cn' -> 'zh-CN', 'zh' -> 'zh-CN', 'eng' -> 'en', 'he' -> 'iw'.
    """
    if code is None:
        return None
    code = code.strip().replace("_", "-")
    if not code:
        return None
    if code.lower() == "auto":
        return "auto"
    base, _, region = code.partition("-")
    base = {"eng": "en", "jpn": "ja", "he": "iw", "jv": "jw"}.get(base.lower(), base.lower())
    if not region:
        return "zh-CN" if base == "zh" else base
    return f"{base}-{region.upper() if len(region) == 2 else region.title()}"


def base_lang(code):
    return code.split("-")[0].lower()


def lang_suffix(lang):
    return f"_{lang}"


def text_filename(slide_num, lang):
    return f"slide_{slide_num}{lang_suffix(lang)}.txt"


def audio_filename(slide_num, lang, model_label):
    return f"slide_{slide_num}{lang_suffix(lang)}.{model_label}.m4a"


def spoken_filename(slide_num, lang, model_label):
    return f"slide_{slide_num}{lang_suffix(lang)}.{model_label}.spoken.txt"


def qwen3_language(lang):
    return QWEN3_LANGUAGES.get(base_lang(lang))


def gpt_sovits_language(lang):
    return GPT_SOVITS_LANGUAGES.get(base_lang(lang))


# Language identification of the notes (--source-lang auto)
_LANGID_TO_FILE_CODE = {"he": "iw", "jv": "jw", "yue": "zh-TW", "wuu": "zh-CN"}
_language_identifier = None


def _get_language_identifier():
    """py3langid classifier (with normalized probabilities) over all its languages."""
    global _language_identifier
    if _language_identifier is None:
        from py3langid import langid
        _language_identifier = langid.LanguageIdentifier.from_modelpath(
            langid.MODEL_DIR / langid.MODEL_FILE, norm_probs=True)
    return _language_identifier


def _chinese_variant(text):
    """zh-CN if all Han characters exist in GB2312 (simplified), otherwise zh-TW."""
    han = "".join(re.findall(r'[一-鿿]', text))
    try:
        han.encode("gb2312")
        return "zh-CN"
    except UnicodeEncodeError:
        return "zh-TW"


def detect_language(text):
    """Identify the language of one note. Returns (language code, confident)."""
    letters = re.sub(r'[\W\d_]', '', text)
    if not letters:
        return "en", False
    kana = len(re.findall(r'[぀-ヿ]', letters))
    hangul = len(re.findall(r'[가-힯ᄀ-ᇿ]', letters))
    han = len(re.findall(r'[一-鿿]', letters))
    if kana and (kana + han) >= 0.3 * len(letters):
        return "ja", True
    if hangul and (hangul + han) >= 0.3 * len(letters):
        return "ko", True
    if han == len(letters):
        # Kanji/hanzi only: Chinese or a kanji-only Japanese note
        return _chinese_variant(text), False
    lang, prob = _get_language_identifier().classify(text)
    lang = _LANGID_TO_FILE_CODE.get(lang, lang)
    if lang == "zh":
        lang = _chinese_variant(text)
    return lang, (len(letters) >= 20 and prob >= 0.9)


def detect_note_languages(texts, known=None):
    """Identify the language of every note ({slide: text} -> {slide: code}).

    Notes that are too short or ambiguous to identify reliably ("Thank you.",
    kanji-only titles, formulas) are assigned the most frequent language among the
    confidently identified notes of the deck.
    """
    first = {n: detect_language(t) for n, t in texts.items()}
    votes = Counter(lang for lang, confident in first.values() if confident)
    votes.update((known or {}).values())
    majority = votes.most_common(1)[0][0] if votes else None
    result = {}
    for n, (lang, confident) in first.items():
        if not confident and majority and lang != majority:
            logger.info(f"Slide #{n}: short/ambiguous note ('{lang}'); using the deck language '{majority}'")
            lang = majority
        result[n] = lang
    return result


def find_source_text(workspace_dir, slide_num, source_lang, exclude_lang=None):
    """Return (lang, path) of the note text of a slide in the source language, or (None, None).

    With source_lang 'auto', any existing slide text is used (Japanese first, then
    English, then other languages), skipping the file of exclude_lang.
    """
    if source_lang and source_lang != "auto":
        if exclude_lang and lang_suffix(source_lang) == lang_suffix(exclude_lang):
            return None, None
        path = os.path.join(workspace_dir, text_filename(slide_num, source_lang))
        return (source_lang, path) if os.path.exists(path) else (None, None)
    candidates = []
    for name in os.listdir(workspace_dir):
        m = _TEXT_FILE_RE.match(name)
        if m and int(m.group(1)) == slide_num:
            candidates.append(normalize_lang(m.group(2)) if m.group(2) else "ja")
    candidates.sort(key=lambda lang: ({"ja": 0, "en": 1}.get(base_lang(lang), 2), lang))
    for lang in candidates:
        if exclude_lang and lang_suffix(lang) == lang_suffix(exclude_lang):
            continue
        return lang, os.path.join(workspace_dir, text_filename(slide_num, lang))
    return None, None

logging.basicConfig(level=logging.INFO, format='[%(asctime)s] %(levelname)s: %(message)s', datefmt='%H:%M:%S')
logger = logging.getLogger(__name__)

# ------------------------------------------
# Structured notes (translated narration + original note)
# ------------------------------------------
# When translated narration is written back into the slide notes, the note is laid out as
#
#   === pptx-narrator: narration [en] from [ja] #3f2a9c0d1e ===
#   <English narration>
#
#   === pptx-narrator: source [ja] ===
#   <original Japanese note>
#
# The hash identifies the version of the source note that was translated. --extract
# recognizes this layout: the source part becomes the note to translate, and the
# narration part is restored as the existing translation only if the source is unchanged.
_NARRATION_MARK_RE = re.compile(
    r"^[ \t]*=+[ \t]*pptx-narrator:[ \t]*narration[ \t]*\[([^\]\n]+)\][ \t]*from[ \t]*\[([^\]\n]+)\]"
    r"[ \t]*#([0-9a-fA-F]{6,40})([ \t]+spoken)?[ \t]*=+[ \t]*$", re.MULTILINE)
_SOURCE_MARK_RE = re.compile(
    r"^[ \t]*=+[ \t]*pptx-narrator:[ \t]*source[ \t]*\[([^\]\n]+)\][ \t]*=+[ \t]*$", re.MULTILINE)
TRANSLATION_MANIFEST = "translations.json"


def text_fingerprint(text):
    """Short hash of a note, insensitive to line-ending and trailing-space differences."""
    norm = unicodedata.normalize("NFC", text).replace("\r\n", "\n").replace("\r", "\n").replace("\v", "\n")
    norm = "\n".join(line.rstrip() for line in norm.strip().split("\n"))
    return hashlib.sha1(norm.encode("utf-8")).hexdigest()[:10]


# ------------------------------------------
# Notes holding texts of several languages
# ------------------------------------------
# When the texts of several languages are written into one note, each is preceded by
# a heading of the same form; no language is treated specially:
#
#   === pptx-narrator: [ja] ===
#   <Japanese text>
#
#   === pptx-narrator: [en] translated from [ja] 2026-09-25T14:02 #3f2a9c0d1e; edited 2026-09-26T10:15 ===
#   <English text>
#
# "translated from" is given when the text was made by translate (the hash identifies
# the version of the source it was made from); "edited" when it was changed after it was
# made. A note of one language has no heading. The earlier layout (narration/source) is
# still read.
_SECTION_MARK_RE = re.compile(r"^[ \t]*=+[ \t]*pptx-narrator:[ \t]*\[([^\]\n]+)\]([^\n]*?)[ \t]*=+[ \t]*$")
_FROM_RE = re.compile(r"translated from \[([^\]\n]+)\](?:[ \t]+(\d{4}-\d\d-\d\d\S*))?(?:[ \t]+#([0-9a-fA-F]{6,40}))?")
_EDITED_RE = re.compile(r"edited[ \t]+(\d{4}-\d\d-\d\d\S*)")


def _short_time(stamp):
    """2026-09-25T14:02:33+0900 -> 2026-09-25T14:02 (local time, to the minute)."""
    return stamp[:16] if stamp else ""


def section_heading(lang, info=None):
    info = info or {}
    text = f"=== pptx-narrator: [{lang}]"
    if info.get("source_lang"):
        text += f" translated from [{info['source_lang']}]"
        if info.get("translated_at"):
            text += " " + _short_time(info["translated_at"])
        if info.get("source_fingerprint"):
            text += f" #{info['source_fingerprint']}"
    if info.get("edited_at"):
        text += ("; " if info.get("source_lang") else " ") + f"edited {_short_time(info['edited_at'])}"
    return text + " ==="


def _heading_of(line):
    """(lang, info) if a line is a section heading (either layout), else None."""
    m = _SECTION_MARK_RE.match(line)
    if m and not m.group(1).startswith(("narration", "source")):
        rest, info = m.group(2), {}
        f = _FROM_RE.search(rest)
        if f:
            info.update(source_lang=normalize_lang(f.group(1)), translated_at=f.group(2) or "",
                        source_fingerprint=(f.group(3) or "").lower())
        e = _EDITED_RE.search(rest)
        if e:
            info["edited_at"] = e.group(1)
        return normalize_lang(m.group(1)), info
    m = _NARRATION_MARK_RE.match(line)
    if m and m.group(4):
        return None, {"spoken": True}
    if m:
        return normalize_lang(m.group(1)), {"source_lang": normalize_lang(m.group(2)),
                                            "source_fingerprint": m.group(3).lower(), "translated_at": ""}
    m = _SOURCE_MARK_RE.match(line)
    if m:
        return normalize_lang(m.group(1)), {}
    return None


def split_note_sections(lines):
    """Divide the lines of a note into [{"lang", "info", "lines"}]; a part before the first
    heading (or a note without headings) has lang None."""
    sections, cur = [], None
    for line in lines:
        h = _heading_of(line)
        if h is not None:
            cur = {"lang": h[0], "info": h[1], "lines": []}
            sections.append(cur)
            continue
        if cur is None:
            cur = {"lang": None, "info": {}, "lines": []}
            sections.append(cur)
        cur["lines"].append(line)
    return sections


def _section_text(lines):
    return "\n".join(lines).strip()


def _paragraph_text(p):
    """The text of one a:p, leaving out struck-through runs (see _text_frame_text)."""
    parts = []
    for child in p:
        if child.tag == _A_NS + "r":
            if _is_struck(child):
                continue
            t = child.find(_A_NS + "t")
            parts.append((t.text or "") if t is not None else "")
        elif child.tag == _A_NS + "fld":
            if _is_page_field(child):
                continue
            t = child.find(_A_NS + "t")
            parts.append((t.text or "") if t is not None else "")
        elif child.tag == _A_NS + "br":
            parts.append("\v")
    return "".join(parts)


def _new_paragraphs(text):
    """Plain a:p elements for a text (one per line; \\v is a line break)."""
    from pptx.oxml.xmlchemy import OxmlElement
    out = []
    for line in text.split("\n"):
        p = OxmlElement("a:p")
        for k, seg in enumerate(line.split("\v")):
            if k:
                p.append(OxmlElement("a:br"))
            if seg:
                r = OxmlElement("a:r")
                t = OxmlElement("a:t")
                t.text = seg
                r.append(t)
                p.append(r)
        out.append(p)
    return out


def _load_manifest(workspace_dir):
    path = os.path.join(workspace_dir, TRANSLATION_MANIFEST)
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            logger.warning(f"Could not read {path}: {e}")
    return {}


def record_translation(workspace_dir, slide_num, target_lang, source_lang, source_text, text=None):
    """Remember which version of the source note a translation was made from, and the
    translation as written, so that an edit of either can be told apart later."""
    manifest = _load_manifest(workspace_dir)
    entry = {"source_lang": source_lang, "source_fingerprint": text_fingerprint(source_text),
             "translated_at": time.strftime("%Y-%m-%dT%H:%M:%S%z")}
    if text is not None:
        entry["fingerprint"] = text_fingerprint(text)
    manifest.setdefault(str(slide_num), {})[target_lang] = entry
    with open(os.path.join(workspace_dir, TRANSLATION_MANIFEST), "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=1, sort_keys=True)


# ==========================================
# Note baseline (hand-edit detection for pack)
# ==========================================
# extract records the fingerprint of each slide's note as it read it, and pack
# records the fingerprint of any note it writes. A note whose fingerprint is
# neither was edited in the deck by hand since, and pack does not overwrite it
# unless --overwrite is given.
NOTE_BASELINE = "note_baseline.json"
PACK_TARGETS = ("all", "audio", "text")


def _load_note_baseline(workspace_dir):
    path = os.path.join(workspace_dir, NOTE_BASELINE)
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            logger.warning(f"Could not read {path}: {e}")
    return {}


def _save_note_baseline(workspace_dir, baseline):
    with open(os.path.join(workspace_dir, NOTE_BASELINE), "w", encoding="utf-8") as f:
        json.dump(baseline, f, ensure_ascii=False, indent=1, sort_keys=True)


def _note_is_known(entry, fingerprint):
    """True if the note is one that extract read or pack wrote."""
    return fingerprint in {entry.get("extracted"), entry.get("packed")} - {None}


_A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _is_struck(run):
    rpr = run.find(_A_NS + "rPr")
    return rpr is not None and rpr.get("strike") in ("sngStrike", "dblStrike")


# Invisible characters left out of a note: zero-width space, non-joiner and joiner, word
# joiner, and the byte-order mark / zero-width no-break space.
_ZERO_WIDTH = {ord(c): None for c in "\u200b\u200c\u200d\u2060\ufeff"}


def _is_page_field(fld):
    """A date or slide-number field that PowerPoint fills in, not text of the note."""
    kind = (fld.get("type") or "").lower()
    return kind.startswith("datetime") or kind == "slidenum"


# Placeholders of the notes page that PowerPoint fills in: not part of the note.
_PAGE_PLACEHOLDERS = {"DATE", "SLIDE_NUMBER", "HEADER", "FOOTER"}


def _is_page_placeholder(shape):
    try:
        return shape.is_placeholder and shape.placeholder_format.type is not None \
            and shape.placeholder_format.type.name in _PAGE_PLACEHOLDERS
    except (AttributeError, ValueError):
        return False


def _text_frame_text(text_frame, struck=None):
    """The text of a text frame as python-pptx gives it (paragraphs joined by \\n,
    line breaks as \\v), leaving out runs with a single or double strikethrough:
    struck-through text in a note is text the author has deleted. Struck runs are
    appended to `struck` when a list is given."""
    paragraphs = []
    for p in text_frame._txBody.iter(_A_NS + "p"):
        parts = []
        for child in p:
            tag = child.tag
            if tag == _A_NS + "r":
                t = child.find(_A_NS + "t")
                text = (t.text or "") if t is not None else ""
                if _is_struck(child):
                    if struck is not None and text:
                        struck.append(text)
                    continue
                parts.append(text)
            elif tag == _A_NS + "fld":
                if _is_page_field(child):
                    continue
                t = child.find(_A_NS + "t")
                parts.append((t.text or "") if t is not None else "")
            elif tag == _A_NS + "br":
                parts.append("\v")
        paragraphs.append("".join(parts))
    return "\n".join(paragraphs)


def _slide_note_raw(slide, struck=None):
    """The note of a slide as extract reads it: the text of every text shape in the
    notes page, excluding the date, slide-number, header and footer placeholders and
    fields, and struck-through text."""
    if not slide.has_notes_slide:
        return ""
    text_list = []
    for shape in slide.notes_slide.shapes:
        if shape.has_text_frame and not _is_page_placeholder(shape):
            found = []
            text = _text_frame_text(shape.text_frame, found).strip()
            if text and not text.isdigit():
                text_list.append(text)
                if struck is not None:
                    struck.extend(found)
    return "\n".join(text_list).strip()


def resolve_targets(target_arg):
    """pack's positional target -> set of {"audio", "text"} (omitted = all)."""
    if not target_arg or target_arg == "all":
        return {"audio", "text"}
    if target_arg in ("audio", "text"):
        return {target_arg}
    raise ValueError(f"Invalid target: {target_arg}. Use 'audio', 'text', or 'all'.")


# ==========================================
# Helper Functions
# ==========================================
def is_japanese(text):
    return bool(re.search(r'[\u3040-\u309F\u30A0-\u30FF\u4E00-\u9FFF]', text))

def parse_slide_ranges(range_str, max_slides):
    slides = set()
    if not range_str: return set(range(1, max_slides + 1))
    for part in range_str.split(','):
        part = part.strip()
        if not part: continue
        try:
            if '-' in part:
                s, e = part.split('-', 1)
                slides.update(range(int(s) if s else 1, (int(e) if e else max_slides) + 1))
            else: slides.add(int(part))
        except ValueError: pass
    return slides

def load_letter_map(json_path: str) -> dict:
    """Loads letter pronunciation mappings from a JSON file."""
    if not json_path or not os.path.exists(json_path):
        return {}
    try:
        with open(json_path, "r", encoding="utf-8") as f:
            logger.info(f"Loaded letter map from: {json_path}")
            return json.load(f)
    except Exception as e:
        logger.error(f"Failed to load letter map '{json_path}': {e}")
        return {}

def spell_out_letters(s: str, letter_map: dict = None) -> str:
    """Spells out characters using the supplied letter_map."""
    if not letter_map:
        return s
    return "".join(letter_map.get(ch.upper(), ch) for ch in s)

# ==========================================
# Dictionary & Term Scanning Logic
# ==========================================
def _is_single_letter(term):
    """A lone Latin letter or digit (e.g. the A of a base pair) is never a useful candidate."""
    return bool(re.fullmatch(r'[A-Za-z0-9]', term.strip()))


_KANJI_RE = re.compile(r'[一-鿿々]')


_COMPOUNDS_WARNED = False


def japanese_compounds(text):
    """Compounds of several words, with the reading a Japanese front end assembles for them.

    A compound the front end has to put together (二|本|鎖) is where a reading can go wrong:
    鎖 is クサリ on its own but サ in 二本鎖. Returns {compound: assembled reading}; the
    reading is a proposal to correct, not an answer.
    """
    try:
        import pyopenjtalk
    except ImportError:
        global _COMPOUNDS_WARNED
        if not _COMPOUNDS_WARNED:
            logger.warning("--scan-compounds needs pyopenjtalk, which is not installed "
                           "(pip install 'pptx-narrator[verify]'); no compounds are proposed.")
            _COMPOUNDS_WARNED = True
        return {}
    try:
        tokens = pyopenjtalk.run_frontend(text)
    except Exception:
        return {}
    out, run = {}, []

    def flush():
        if len(run) > 1:
            word = "".join(t["string"] for t in run)
            try:
                out[word] = unicodedata.normalize("NFKC", pyopenjtalk.g2p(word, kana=True))
            except Exception:
                pass
        run.clear()

    for token in tokens:
        if token.get("pos") == "名詞" and _KANJI_RE.search(token.get("string", "")):
            run.append(token)
        else:
            flush()
    flush()
    return out


def step_scan_and_update_dict(workspace_dir, dict_path, entries, requested_slides, target_lang,
                              source_lang="auto", for_translation=False, propose_compounds=False):
    """Collect candidate terms and append them to the dictionary file of the run.

    The scanned text is the text the dictionary will be applied to: the narration text in
    --target-lang when it exists, with provisional readings filled in; otherwise (or when
    combined with --translate) the notes to be translated, with blank replacements for review.
    """
    logger.info("--- [Option: Scan] Scanning text for technical terms ---")

    _ensure_nltk_data()
    from nltk.corpus import stopwords, words
    stop_words = set(stopwords.words('english'))
    common_english = set(w.lower() for w in words.words())

    narration_texts = [(os.path.join(workspace_dir, text_filename(n, target_lang)), target_lang)
                       for n in requested_slides
                       if os.path.exists(os.path.join(workspace_dir, text_filename(n, target_lang)))]
    if narration_texts and not for_translation:
        texts = narration_texts
    else:
        # nothing in --target-lang yet (or translating in this run): scan the notes to be translated
        texts = []
        for slide_num in requested_slides:
            src_lang, src_p = find_source_text(workspace_dir, slide_num, source_lang, exclude_lang=target_lang)
            if src_p is not None:
                texts.append((src_p, src_lang))
    if not texts:
        logger.info("No note texts to scan.")
        return

    jp_char = r'[぀-ヿ一-鿿]'
    latin = 'A-Za-zÀ-ÖØ-öø-ɏ'
    generic_pattern = rf'(?<![{latin}])[{latin}]{{2,}}(?![{latin}])|[゠-ヿ]{{3,}}'
    bypass_pattern = (
        r'\d+(?:\.\d+)?\s*(?:%|℃|°C|nm|μm|mm|cm|km|kg|mg|µg|ng|pg|ml|μl|kb|Mb|Gb|bp|kDa|Å)'
        r'|\b(?:I|II|III|IV|V|VI|VII|VIII|IX|X)\b'
    )

    existing_terms = {t.lower() for t, _, _ in entries or []}
    candidates = {}  # term -> is narration text (reading candidate)
    for p, file_lang in texts:
        with open(p, 'r', encoding='utf-8') as f:
            text = f.read()
        narration_text = lang_suffix(file_lang) == lang_suffix(target_lang)
        # English word lists only help for English words (English notes, or Latin words
        # embedded in Japanese). For other languages keep acronyms / mixed-case terms only.
        english_filter = base_lang(file_lang) in ("en", "ja")

        generic_found = re.findall(generic_pattern, text)
        bypass_found = set(re.findall(bypass_pattern, text))
        for m in re.finditer(rf'[A-Za-z0-9]+(?={jp_char})|(?<={jp_char})[A-Za-z0-9]+', text):
            tok = m.group()
            if re.search(r'[A-Za-z]', tok):
                bypass_found.add(tok)

        for term in generic_found:
            term_lower = term.lower()
            if _is_single_letter(term) or term_lower in existing_terms or term_lower in stop_words:
                continue
            if re.fullmatch(rf'[{latin}]+', term):
                acronym_like = term.isupper() or bool(re.search(r'[a-z][A-Z]', term))
                if len(term) < 3 and not term.isupper():
                    continue
                if english_filter and not term.isupper() and term_lower in common_english:
                    continue
                if not english_filter and not acronym_like:
                    continue
            candidates[term] = narration_text
            existing_terms.add(term_lower)

        for term in bypass_found:
            term_lower = term.lower()
            if _is_single_letter(term) or term_lower in existing_terms:
                continue
            candidates[term] = narration_text
            existing_terms.add(term_lower)

    if not candidates:
        logger.info("No new terms found.")
        return

    ROMAN_NUMERAL_READINGS = {
        "I": "いち", "II": "に", "III": "さん", "IV": "よん", "V": "ご",
        "VI": "ろく", "VII": "なな", "VIII": "はち", "IX": "きゅう", "X": "じゅう",
    }
    tgt_base = base_lang(target_lang)
    new_entries = []
    for term in sorted(candidates):
        repl = ""
        if candidates[term]:
            if tgt_base == "ja" and term in ROMAN_NUMERAL_READINGS:
                repl = ROMAN_NUMERAL_READINGS[term]
            elif re.match(r'^[A-Z]+$', term):
                repl = " ".join(term)
            elif tgt_base == "ja" and is_japanese(term):
                repl = term
        new_entries.append((term, repl))
        logger.info(f"New term: {term} -> {repl or '(blank)'}")

    append_dictionary_entries(dict_path, new_entries)

    if propose_compounds and base_lang(target_lang) == "ja":
        known = {t for t, _, _ in entries or []} | {t for t, _ in new_entries}
        proposals = {}
        for path, file_lang in texts:
            if base_lang(file_lang) != "ja":
                continue
            with open(path, "r", encoding="utf-8") as f:
                for word, reading in japanese_compounds(f.read()).items():
                    if word not in known:
                        proposals[word] = reading
        if proposals:
            with open(dict_path, "a", encoding="utf-8", newline="") as f:
                f.write("\n; Compounds of the notes with the reading a Japanese front end assembles for them.\n"
                        "; They are comments, so they do nothing until the ';' is removed; correct the reading\n"
                        "; first (鎖 is read クサリ on its own but サ in 二本鎖) or leave the line as it is.\n")
                for word in sorted(proposals):
                    f.write(f";{word},{proposals[word]},\n")
            logger.info(f"{len(proposals)} compound(s) proposed as comment lines in {dict_path}")
    logger.info(f"Added {len(new_entries)} entries to {dict_path}")

SI_PREFIXES = {
    "p": "ピコ", "n": "ナノ", "μ": "マイクロ", "µ": "マイクロ", "m": "ミリ",
    "c": "センチ", "k": "キロ", "K": "キロ", "M": "メガ", "G": "ギガ", "T": "テラ",
}
BASE_UNITS = {
    "bp": "えんきつい",
    "b": "ベース",
    "m": "メートル",
    "g": "グラム",
    "l": "リットル",
    "L": "リットル",
    "Da": "ダルトン",
}
SPECIAL_UNITS = {
    "%": "パーセント",
    "℃": "ど",
    "°C": "ど",
    "Å": "オングストローム",
    "rpm": "アールピーエム",
}

def _unit_reading(unit, extra_units, use_builtin=True):
    if unit in extra_units:
        return extra_units[unit]
    if not use_builtin:
        return None
    if unit in SPECIAL_UNITS:
        return SPECIAL_UNITS[unit]
    for split in range(1, len(unit)):
        prefix, base = unit[:split], unit[split:]
        if prefix in SI_PREFIXES and base in BASE_UNITS:
            return SI_PREFIXES[prefix] + BASE_UNITS[base]
    if unit in BASE_UNITS:
        return BASE_UNITS[unit]
    return None

def normalize_units(text, extra_units=None, letter_map=None, lang="ja"):
    """Rewrite number + unit symbol into a spoken reading.

    Built-in SI-prefix readings exist for Japanese only; for other languages, readings
    come from dictionary entries of Type 'unit' (and the optional letter map).
    """
    extra_units = extra_units or {}
    use_builtin = base_lang(lang) == "ja"
    if not use_builtin and not extra_units and not letter_map:
        return text
    sep = "" if base_lang(lang) in CJK_LANGS else " "

    def _replace(m):
        number, letters = m.group(1), m.group(2)
        reading = _unit_reading(letters, extra_units, use_builtin)
        if reading is not None:
            return number + sep + reading
        if letter_map:
            return number + sep + spell_out_letters(letters, letter_map)
        return number + letters if use_builtin else m.group(0)

    pattern = r'(\d+(?:\.\d+)?)\s*([A-Za-zμµÅ°%℃]+)(?![A-Za-z0-9_])'
    return re.sub(pattern, _replace, text)

# ------------------------------------------
# Dictionaries
# ------------------------------------------
# A dictionary is a list of plain string replacements, one per CSV row:
#
#     string,replacement,type
#     Gbp,ギガベースペア,unit
#     mRNA,メッセンジャーアールエヌエー,
#
# It is applied to the text processed in the run: with --translate, to the notes
# before translation (e.g. to fix how terms are translated); with --tts, to the
# narration text before synthesis (e.g. readings); with both, at both points.
# Translation and synthesis can be run separately to use different dictionaries. type 'unit' marks unit symbols
# that follow a number. The header row is optional. Files in the v1.x layout
# (Term,Japanese_Reading,English_Reading,Type) are read using the column of the
# narration language (ja or en).
DICT_HEADER = ["string", "replacement", "type"]
_HEADER_NAMES = {"string", "term", "replacement", "reading", "type"}


COMMENT_MARKER = ";"


def _strip_dictionary_comment(line):
    """Drop a ';' comment from a dictionary line.

    The marker starts a comment only at the start of the line or after a space, and not
    inside double quotes, so a term that contains or begins with ';' is still possible."""
    in_quotes = False
    for i, ch in enumerate(line):
        if ch == '"':
            in_quotes = not in_quotes
        elif ch == COMMENT_MARKER and not in_quotes and (i == 0 or line[i - 1].isspace()):
            return line[:i].strip()
    return line.strip()


def read_dictionary_file(path, lang=None):
    """Return [(string, replacement, type)] from one dictionary CSV."""
    with open(path, 'r', encoding='utf-8', newline='') as f:
        # A ';' at the start of a line or after a space begins a comment (see
        # _strip_dictionary_comment), so a line can explain or switch off an entry; a line
        # with nothing left in front of the comment is skipped.
        lines = [_strip_dictionary_comment(line) for line in f]
        rows = [r for r in csv.reader([ln for ln in lines if ln]) if any(c.strip() for c in r)]
    if not rows:
        return []
    head = [c.strip().lower() for c in rows[0]]
    if "japanese_reading" in head or "english_reading" in head:
        column = {"ja": "japanese_reading", "en": "english_reading"}.get(base_lang(lang)) if lang else None
        if column not in head:
            logger.warning(f"{path} is a v1.x dictionary without a column for '{lang}'; it is not used.")
            return []
        term_col, repl_col = 0, head.index(column)
        type_col = head.index("type") if "type" in head else None
        body = rows[1:]
    else:
        term_col, repl_col, type_col = 0, 1, 2
        body = rows[1:] if len(head) >= 2 and head[0] in _HEADER_NAMES and head[1] in _HEADER_NAMES else rows
    entries = []
    for r in body:
        if len(r) <= repl_col or not r[term_col].strip():
            continue
        if "\\" in r[term_col]:
            logger.warning(f"{path}: the entry '{r[term_col].strip()}' contains a backslash and is matched literally; "
                           "write the string as it appears in the notes.")
        typ = r[type_col].strip().lower() if type_col is not None and len(r) > type_col else ""
        entries.append((r[term_col].strip(), r[repl_col].strip(), "unit" if typ == "unit" else ""))
    return entries


def load_dictionaries(paths, lang=None):
    """Read the given dictionary files in order; later files override earlier entries."""
    merged = {}
    for p in paths or []:
        if not os.path.exists(p):
            logger.info(f"Dictionary {p} does not exist yet.")
            continue
        entries = read_dictionary_file(p, lang)
        for term, repl, typ in entries:
            merged[(term, typ)] = repl
        logger.info(f"Dictionary {p}: {len(entries)} entries")
    return [(term, repl, typ) for (term, typ), repl in merged.items()]


def append_dictionary_entries(path, new_entries):
    """Append (string, replacement) rows; create the file with a header if needed."""
    is_new = not os.path.exists(path) or os.path.getsize(path) == 0
    needs_newline = False
    if not is_new:
        with open(path, 'rb') as fb:
            fb.seek(-1, os.SEEK_END)
            needs_newline = fb.read(1) not in (b'\n', b'\r')
    with open(path, 'a', encoding='utf-8', newline='') as f:
        if needs_newline:
            f.write('\n')
        writer = csv.writer(f)
        if is_new:
            writer.writerow(DICT_HEADER)
        for term, repl in new_entries:
            writer.writerow([term, repl, ""])


# Apostrophes and primes come in several shapes: PowerPoint turns ' into ’ while
# notes pasted from a paper may use ′. A dictionary entry written with one of them should
# still match the others, so both text and terms are mapped to the plain forms.
_LOOKALIKES = {ord(c): "'" for c in "‘’‛′ʹ´＇"}
_LOOKALIKES.update({ord(c): '"' for c in "“”‟″＂"})
_LOOKALIKES.update({ord(c): "-" for c in "‐‑‒–−－"})


def normalize_lookalikes(text):
    """Map typographic quotes, primes and dashes to their plain ASCII forms."""
    return text.translate(_LOOKALIKES)


def _replace_term(text, term, replacement):
    if re.match(r'^[a-zA-Z0-9_ \-]+$', term):
        pattern = rf'(?<![A-Za-z0-9_]){re.escape(term)}(?![A-Za-z0-9_])'
        return re.subn(pattern, lambda _m: replacement, text)
    return text.replace(term, replacement), text.count(term)


def apply_dictionary(text, entries, lang, letter_map=None, builtin_units=True):
    """Apply the replacements (longest string first), then rewrite number + unit expressions.

    builtin_units=False (used before translation) disables the built-in Japanese unit
    readings and the letter map, so that only the dictionary's own unit entries apply.
    """
    text = normalize_lookalikes(unicodedata.normalize('NFC', text))
    units, terms = {}, {}
    for term, repl, typ in entries or []:
        term = normalize_lookalikes(unicodedata.normalize('NFC', term))
        repl = unicodedata.normalize('NFC', repl)
        if repl:
            (units if typ == "unit" else terms)[term] = repl
    for term, repl in sorted(terms.items(), key=lambda x: len(x[0]), reverse=True):
        text, _ = _replace_term(text, term, repl)
    if builtin_units:
        return normalize_units(text, units, letter_map=letter_map, lang=lang)
    return normalize_units(text, units, letter_map=None, lang="und")

# ==========================================
# Pipeline Steps
# ==========================================
def _text_langs_of_slide(workspace_dir, s_num):
    """Languages of the texts of one slide in the workspace, in language-code order."""
    langs = []
    for name in os.listdir(workspace_dir):
        m = _TEXT_FILE_RE.match(name)
        if m and int(m.group(1)) == s_num:
            langs.append(normalize_lang(m.group(2)) if m.group(2) else "ja")
    return sorted(langs, key=lang_suffix)


def _made_record(workspace_dir, s_num, lang, manifest, baseline):
    """How a text of the workspace was made: (heading info, fingerprint when made)."""
    rec = manifest.get(str(s_num), {}).get(lang) or {}
    made = (baseline.get(str(s_num), {}).get("langs") or {}).get(lang) or {}
    info = {}
    if rec.get("source_lang"):
        info.update(source_lang=rec["source_lang"], translated_at=rec.get("translated_at", ""),
                    source_fingerprint=rec.get("source_fingerprint", ""))
        return info, rec.get("fingerprint"), rec.get("edited_at") or ("?" if rec.get("edited") else "")
    return info, made.get("extracted"), made.get("edited_at", "")


def _section_info(workspace_dir, s_num, lang, text, manifest, baseline):
    """What the heading of a text says: how it was made, and when it was edited since."""
    info, made_fp, edited_at = _made_record(workspace_dir, s_num, lang, manifest, baseline)
    path = os.path.join(workspace_dir, text_filename(s_num, lang))
    if made_fp and made_fp != text_fingerprint(text) and os.path.exists(path):
        edited_at = time.strftime("%Y-%m-%dT%H:%M:%S%z", time.localtime(os.path.getmtime(path)))
    if edited_at and edited_at != "?":
        info["edited_at"] = edited_at
    return info


def step_extract_notes(pptx_path, workspace_dir, requested_slides, source_lang="auto",
                       update=False, overwrite=False):
    """Write the note of each slide into the workspace, one text per language.

    A note with texts of several languages gives one text per language. A text already in
    the workspace is replaced only with update (when it was not edited in the workspace
    since it was written) or overwrite. Returns the number of texts found in the deck.
    """
    logger.info("--- [Option: Extract] Extracting Notes ---")
    prs = Presentation(pptx_path)
    plain, headed = {}, {}
    raw_fingerprints = {}
    for slide_num in requested_slides:
        if slide_num > len(prs.slides): continue
        slide = prs.slides[slide_num - 1]

        if slide._element.get('show') == '0':
            logger.info(f"Slide #{slide_num} is a hidden slide (skipped)")
            continue

        struck = []
        raw_text = _slide_note_raw(slide, struck)
        if struck:
            logger.info(f"Slide #{slide_num}: left out {len(struck)} struck-through passage(s): "
                        + ", ".join(repr(t) for t in struck[:5]) + (" ..." if len(struck) > 5 else ""))
        raw_fingerprints[slide_num] = text_fingerprint(raw_text)

        txt = raw_text.translate(_ZERO_WIDTH).strip()
        if not txt:
            logger.info(f"Slide #{slide_num} has no notes (skipped)")
            continue
        parts = [dict(s, text=_section_text(s["lines"])) for s in split_note_sections(txt.split("\n"))
                 if s["lang"]]
        parts = [s for s in parts if s["text"]]
        if parts:
            headed[slide_num] = parts
        else:
            plain[slide_num] = txt

    if source_lang == "auto":
        languages = detect_note_languages(plain, known={n: p[0]["lang"] for n, p in headed.items()})
    else:
        # An explicit language is a selector. Do not relabel a note merely because
        # the caller requested a language; retain only notes whose detected language
        # agrees with the request. Short/ambiguous notes are accepted when the
        # detector cannot make a confident distinction and the requested language
        # is the only explicit choice.
        languages = {}
        for n, txt in plain.items():
            detected, confident = detect_language(txt)
            if confident and base_lang(detected) != base_lang(source_lang):
                logger.info(f"Slide #{n}: note detected as '{detected}', not requested '{source_lang}' (skipped)")
                continue
            languages[n] = source_lang

    items = [(n, languages[n], txt, {}) for n, txt in plain.items() if n in languages]
    for n, parts in headed.items():
        for part in parts:
            if source_lang != "auto" and base_lang(part["lang"]) != base_lang(source_lang):
                logger.info(f"Slide #{n}: the {part['lang']} part of the note is not the requested "
                            f"'{source_lang}' (skipped)")
                continue
            items.append((n, part["lang"], part["text"], part["info"]))
    items.sort(key=lambda it: (it[0], lang_suffix(it[1])))

    baseline = _load_note_baseline(workspace_dir)
    manifest = _load_manifest(workspace_dir)
    found, kept, translations_read = 0, [], False
    for n, lang, text, info in items:
        found += 1
        name = text_filename(n, lang)
        path = os.path.join(workspace_dir, name)
        new_fp = text_fingerprint(text)
        entry = baseline.setdefault(str(n), {})
        rec = (entry.get("langs") or {}).get(lang) or {}
        if os.path.exists(path) and text_fingerprint(_read_text(path)) != new_fp:
            cur_fp = text_fingerprint(_read_text(path))
            known = {rec.get("extracted"), rec.get("packed")} - {None}
            if n in plain:
                known |= {entry.get("extracted"), entry.get("packed")} - {None}
            unedited = cur_fp in known
            if not (overwrite or (update and unedited)):
                kept.append(n)
                if unedited:
                    logger.warning(f"Slide #{n}: the note in the deck has changed since {name} was written; "
                                   f"{name} is kept (add --update to take the note in again, or --overwrite).")
                else:
                    logger.warning(f"Slide #{n}: {name} differs from the note in the deck and was edited in the "
                                   f"workspace (or has no record of being written); it is kept "
                                   f"(add --overwrite to replace it with the note).")
                continue
            logger.info(f"Slide #{n}: {name} replaced with the note in the deck "
                        f"({'--overwrite' if overwrite else '--update'}).")
        elif os.path.exists(path):
            logger.info(f"Slide #{n}: {name} is already the same as the note ({lang}).")
        else:
            logger.info(f"Slide #{n}: note extracted to {name} (language: {lang}).")
        if not (os.path.exists(path) and text_fingerprint(_read_text(path)) == new_fp):
            with open(path, "w", encoding="utf-8") as f:
                f.write(text)
        made = {"extracted": new_fp}
        if info.get("edited_at"):
            made["edited_at"] = info["edited_at"]
        entry.setdefault("langs", {})[lang] = made
        entry["extracted"] = raw_fingerprints[n]
        if info.get("source_lang"):
            translations_read = True
            edited = bool(info.get("edited_at"))
            manifest.setdefault(str(n), {})[lang] = {
                "source_lang": info["source_lang"], "source_fingerprint": info.get("source_fingerprint", ""),
                "translated_at": info.get("translated_at", ""), "fingerprint": None if edited else new_fp,
                "edited": edited, "edited_at": info.get("edited_at", "")}
            logger.info(f"Slide #{n}: {name} was translated from {info['source_lang']}"
                        + (" and edited since" if edited else "") + " (from the heading in the note).")

    if items and baseline != _load_note_baseline(workspace_dir):
        _save_note_baseline(workspace_dir, baseline)
    if translations_read:
        with open(os.path.join(workspace_dir, TRANSLATION_MANIFEST), "w", encoding="utf-8") as f:
            json.dump(manifest, f, ensure_ascii=False, indent=1, sort_keys=True)
    counts = Counter(lang for _, lang, _, _ in items)
    if counts:
        logger.info("Note languages: " + ", ".join(f"{l} ({c} slides)" for l, c in counts.most_common()))
    if kept:
        logger.warning("Texts kept although the note differs, on slide(s): "
                       + ", ".join(map(str, sorted(set(kept)))) + ".")
    return found

# ------------------------------------------
# Translation
# ------------------------------------------
# Notes are translated by an instruction-tuned language model run locally through
# transformers (default Qwen/Qwen3-4B; see README for larger models). Each note is
# given whole, so the model sees the context of every sentence. Entries of the
# translation dictionary (translate --dict-file) that occur in a note are given to the
# model as instructions ("render X as Y"); the note itself is passed unchanged.
TRANSLATE_MODEL_DEFAULT = "Qwen/Qwen3-4B"

LANGUAGE_NAMES = {
    "ja": "Japanese", "en": "English", "de": "German", "fr": "French", "es": "Spanish",
    "it": "Italian", "pt": "Portuguese", "ru": "Russian", "ko": "Korean", "zh-CN": "Simplified Chinese",
    "zh-TW": "Traditional Chinese", "nl": "Dutch", "sv": "Swedish", "pl": "Polish", "tr": "Turkish",
    "vi": "Vietnamese", "th": "Thai", "id": "Indonesian", "ar": "Arabic", "iw": "Hebrew", "hi": "Hindi",
}


def language_name(code):
    return LANGUAGE_NAMES.get(code) or LANGUAGE_NAMES.get(base_lang(code)) or code


class LLMTranslator:
    """Translation with a causal language model that has a chat template."""

    def __init__(self, model_id=TRANSLATE_MODEL_DEFAULT, device="auto"):
        self.model_id = model_id
        self.device = resolve_torch_device(device)
        self._tok = self._model = None

    def _load(self):
        if self._model is None:
            import torch
            from transformers import AutoModelForCausalLM, AutoTokenizer
            dtype = torch.float32 if self.device == "cpu" else torch.bfloat16
            logger.info(f"Loading translation model {self.model_id} on {self.device}...")
            self._tok = AutoTokenizer.from_pretrained(self.model_id)
            self._model = AutoModelForCausalLM.from_pretrained(self.model_id, torch_dtype=dtype)
            self._model.to(self.device).eval()

    def translate(self, text, source, target, glossary=None):
        import torch
        self._load()
        system = (f"You translate the presenter notes of a lecture slide deck from {language_name(source)} "
                  f"into {language_name(target)}. The notes are read aloud as narration. Translate "
                  f"faithfully and completely, without adding or leaving out content, keep the paragraph "
                  f"breaks, and output only the translation.")
        if glossary:
            system += ("\nTranslate the following terms as given:\n"
                       + "\n".join(f"- {term} -> {rendering}" for term, rendering in glossary))
        messages = [{"role": "system", "content": system}, {"role": "user", "content": text}]
        try:  # Qwen3: answer directly, without a reasoning block
            prompt = self._tok.apply_chat_template(messages, tokenize=False, add_generation_prompt=True,
                                                   enable_thinking=False)
        except TypeError:
            prompt = self._tok.apply_chat_template(messages, tokenize=False, add_generation_prompt=True)
        ids = self._tok(prompt, return_tensors="pt").to(self.device)
        torch.manual_seed(0)
        with torch.no_grad():
            # sampling settings recommended on the Qwen3 model card for non-thinking use
            out = self._model.generate(**ids, max_new_tokens=max(256, 4 * len(text)), do_sample=True,
                                       temperature=0.7, top_p=0.8, top_k=20)
        answer = self._tok.decode(out[0][ids["input_ids"].shape[1]:], skip_special_tokens=True)
        return re.sub(r"(?s)^\s*<think>.*?</think>\s*", "", answer).strip()


def make_translator(model_id=TRANSLATE_MODEL_DEFAULT, device="auto"):
    return LLMTranslator(model_id, device)


def glossary_for(text, dictionary):
    """(term, rendering) pairs of the dictionary whose term occurs in the text, longest first."""
    seen, out = set(), []
    for term, repl, _typ in sorted(dictionary or [], key=lambda e: -len(e[0])):
        if term and repl and term not in seen and term in text:
            seen.add(term)
            out.append((term, repl))
    return out


def step_translate_notes(workspace_dir, requested_slides, source_lang, target_lang,
                         dictionary=None, overwrite=False, update=False, model=TRANSLATE_MODEL_DEFAULT,
                         device="auto"):
    """Translate the selected slides; an existing translation is kept unless asked otherwise.

    update: translate again where the source text has changed since the translation
            was made, unless the translation was edited since (that edit is kept).
    overwrite: translate again every selected slide.
    """
    logger.info(f"--- [Option: Translate] Translating notes into '{target_lang}' with {model} ---")
    translator = None
    sources_found = 0
    manifest = _load_manifest(workspace_dir)
    for slide_num in requested_slides:
        tgt_p = os.path.join(workspace_dir, text_filename(slide_num, target_lang))
        src_lang, src_p = find_source_text(workspace_dir, slide_num, source_lang, exclude_lang=target_lang)
        if src_p is None:
            continue
        sources_found += 1
        with open(src_p, "r", encoding="utf-8") as f:
            text = f.read().strip()
        if not text:
            logger.warning(f"Slide #{slide_num}: {os.path.basename(src_p)} is empty; nothing to translate.")
            continue
        existing = _read_text(tgt_p)
        if existing and not overwrite:
            record = (manifest.get(str(slide_num), {}).get(target_lang) or {})
            source_changed = record.get("source_fingerprint") != text_fingerprint(text)
            edited = bool(record.get("edited")) or record.get("fingerprint") not in (None, text_fingerprint(existing))
            name = os.path.basename(tgt_p)
            if not record:
                logger.warning(f"Slide #{slide_num}: {name} already exists and has no record of being "
                               f"translated from {os.path.basename(src_p)}; kept (add --overwrite to replace it).")
                continue
            if not source_changed:
                logger.info(f"Slide #{slide_num}: {name} is up to date with its source; kept.")
                continue
            if not update:
                logger.warning(f"Slide #{slide_num}: {name} already exists and its source has changed "
                               f"since; kept (add --update to translate it again, or --overwrite).")
                continue
            if edited:
                logger.warning(f"Slide #{slide_num}: the source of {name} has changed, but {name} was "
                               f"edited after it was translated; kept (add --overwrite to replace it).")
                continue
        if translator is None:
            translator = make_translator(model, device)
        glossary = glossary_for(text, dictionary)
        t0 = time.time()
        try:
            translated = translator.translate(text, src_lang, target_lang, glossary=glossary or None)
        except Exception as e:
            logger.error(f"Slide #{slide_num}: translation {src_lang} -> {target_lang} failed: {e}")
            continue
        if not translated.strip():
            logger.error(f"Slide #{slide_num}: the translation came back empty (not written).")
            continue
        with open(tgt_p, "w", encoding="utf-8") as out_f:
            out_f.write(translated.strip() + "\n")
        record_translation(workspace_dir, slide_num, target_lang, src_lang, text, text=translated)
        logger.info(f"Slide #{slide_num}: translated {src_lang} -> {target_lang} "
                    f"({time.time() - t0:.1f}s{', %d dictionary term(s)' % len(glossary) if glossary else ''}).")
    return sources_found

def step_generate_audio(
    workspace_dir,
    requested_slides,
    lang,
    ref_wav,
    ref_text_f,
    ref_lang,
    api_url,
    dictionaries,
    model_label,
    enable_drc=False,
    drc_threshold=-20.0,
    drc_ratio=3.0,
    letter_map=None,
):
    logger.info("--- [Option: TTS / engine=gpt_sovits] Generating Audio ---")
    text_lang = gpt_sovits_language(lang)
    prompt_lang = gpt_sovits_language(ref_lang)
    if text_lang is None or prompt_lang is None:
        logger.error(f"GPT-SoVITS supports only: {', '.join(GPT_SOVITS_LANGUAGES)} "
                     f"(narration '{lang}', reference '{ref_lang}')")
        return
    with open(ref_text_f, "r", encoding="utf-8") as f:
        ref_txt = f.read().strip()
    tts_url = api_url.rstrip("/") + "/tts"
    stats = SynthesisStats("gpt_sovits", api_url)

    for slide_num in requested_slides:
        txt_p = os.path.join(workspace_dir, text_filename(slide_num, lang))
        if not os.path.exists(txt_p):
            continue

        with open(txt_p, "r", encoding="utf-8") as f:
            spoken_text = apply_dictionary(f.read().strip(), dictionaries, lang, letter_map=letter_map)
        if not spoken_text:
            continue

        with open(os.path.join(workspace_dir, spoken_filename(slide_num, lang, model_label)),
                  "w", encoding="utf-8") as f:
            f.write(spoken_text)

        payload = {
            "text": spoken_text,
            "text_lang": text_lang,
            "ref_audio_path": os.path.abspath(ref_wav),
            "prompt_text": ref_txt,
            "prompt_lang": prompt_lang,
            "text_split_method": "cut5",
            "media_type": "wav",
        }

        started = time.time()
        wav_p = os.path.join(workspace_dir, "temp.wav")
        try:
            res = requests.post(tts_url, json=payload, timeout=600)
        except requests.RequestException as e:
            logger.error(f"Slide #{slide_num}: could not reach the GPT-SoVITS server at {api_url}: {e}")
            continue
        if res.status_code != 200:
            logger.error(f"Slide #{slide_num}: the GPT-SoVITS server returned [{res.status_code}]: {res.text}")
            continue
        try:
            with open(wav_p, "wb") as f:
                f.write(res.content)
            audio = AudioSegment.from_file(wav_p)
            if enable_drc:
                audio = compress_dynamic_range(audio, threshold=drc_threshold, ratio=drc_ratio)
                audio = normalize(audio)
            audio.export(os.path.join(workspace_dir, audio_filename(slide_num, lang, model_label)),
                         format="ipod")
            stats.add(slide_num, time.time() - started, len(audio))
        except Exception as e:
            logger.error(f"Slide #{slide_num}: the audio from GPT-SoVITS could not be converted to m4a: {e}")
        finally:
            if os.path.exists(wav_p):
                os.remove(wav_p)
    stats.report()

class SynthesisStats:
    """Collects synthesis times so that a run reports its cost (see --tts)."""

    def __init__(self, engine, device):
        self.engine, self.device = engine, device
        self.slides = 0
        self.seconds = 0.0
        self.audio_seconds = 0.0

    def add(self, slide_num, seconds, audio_ms):
        self.slides += 1
        self.seconds += seconds
        self.audio_seconds += audio_ms / 1000.0
        logger.info(f"Slide {slide_num}: {audio_ms / 1000.0:.1f} s of audio synthesized in {seconds:.1f} s "
                    f"({seconds / max(audio_ms / 1000.0, 1e-9):.2f} x audio time).")

    def report(self):
        if not self.slides:
            return
        logger.info(f"Synthesis summary ({self.engine}, {self.device}): {self.slides} slides, "
                    f"{self.audio_seconds / 60:.1f} min of audio in {self.seconds / 60:.1f} min "
                    f"({self.audio_seconds / max(self.seconds, 1e-9):.2f} x real time, "
                    f"{self.seconds / self.slides:.1f} s per slide).")


def split_into_chunks(text):
    """Split text into sentences: after 。！？, or after . ! ? followed by whitespace."""
    parts = re.split(r'(?<=[。！？])|(?<=[.!?])\s+', text)
    return [p.strip() for p in parts if p and p.strip()]

# Qwen3-TTS 12Hz models emit 12 codec frames per second of audio, so a cap on the
# number of new tokens is a cap on the duration.
QWEN3_CODEC_HZ = 12
RUNAWAY_FACTOR = 3.0


def chunk_token_budget(chunk_text, chars_per_sec, hz=QWEN3_CODEC_HZ, factor=RUNAWAY_FACTOR):
    """Most codec tokens a chunk of text can plausibly need, with room to spare.

    The same margin the anomaly check uses, so a generation is cut off only where it
    would have been rejected anyway.
    """
    expected_sec = max(len(chunk_text) / chars_per_sec, 1.0)
    return int(round(max(expected_sec * factor, 20.0) * hz)) + hz


_qwen3_model_cache = {}

def resolve_torch_device(device):
    """Resolve 'auto' to cuda:0, mps or cpu depending on what is available."""
    if device != "auto":
        return device
    try:
        import torch
    except ImportError:
        return "cpu"
    if torch.cuda.is_available():
        return "cuda:0"
    mps = getattr(torch.backends, "mps", None)
    if mps is not None and mps.is_available():
        return "mps"
    return "cpu"

def _load_qwen3_model(model_size, device):
    device = resolve_torch_device(device)
    key = (model_size, device)
    if key not in _qwen3_model_cache:
        import torch
        from qwen_tts import Qwen3TTSModel
        try:
            import transformers
            transformers.logging.set_verbosity_error()
        except Exception:
            pass
        model_name = f"Qwen/Qwen3-TTS-12Hz-{model_size}-Base"
        logger.info(f"Loading Qwen3-TTS model ({model_name}, device={device})...")
        base_kwargs = dict(device_map=device)
        if str(device).startswith("cuda"):
            import importlib.util
            if importlib.util.find_spec("flash_attn") is not None:
                base_kwargs["attn_implementation"] = "flash_attention_2"

        model = None
        last_err = None
        for dtype in (torch.bfloat16, torch.float16, torch.float32):
            try:
                model = Qwen3TTSModel.from_pretrained(model_name, dtype=dtype, **base_kwargs)
                logger.info(f"Successfully loaded model with dtype={dtype}")
                break
            except Exception as e:
                last_err = e
                logger.warning(f"Failed to load with dtype={dtype}: {e}")
        if model is None:
            raise last_err
        _qwen3_model_cache[key] = model
    return _qwen3_model_cache[key]

def step_generate_audio_qwen3(
    workspace_dir,
    requested_slides,
    lang,
    ref_wav,
    ref_text_f,
    dictionaries,
    model_label,
    qwen3_model_size,
    qwen3_device,
    enable_drc=False,
    drc_threshold=-20.0,
    drc_ratio=3.0,
    letter_map=None,
):
    logger.info("--- [Option: TTS / engine=qwen3] Generating Audio ---")
    language = qwen3_language(lang)
    if language is None:
        logger.error(f"Qwen3-TTS does not support '{lang}' (supported: {', '.join(QWEN3_LANGUAGES)})")
        return
    try:
        import numpy as np
        import soundfile as sf
    except ImportError as e:
        raise RuntimeError(f"Qwen3-TTS needs a library that is not installed ({e}); "
                           "pip install 'pptx-narrator[qwen3]'") from e

    with open(ref_text_f, "r", encoding="utf-8") as f:
        ref_txt = f.read().strip()

    model = _load_qwen3_model(qwen3_model_size, qwen3_device)
    stats = SynthesisStats(f"qwen3/{qwen3_model_size}", resolve_torch_device(qwen3_device))

    voice_clone_prompt = None
    try:
        logger.info("Analyzing reference audio (create_voice_clone_prompt)...")
        voice_clone_prompt = model.create_voice_clone_prompt(ref_audio=ref_wav, ref_text=ref_txt)
    except AttributeError:
        logger.warning(
            "This qwen_tts version lacks create_voice_clone_prompt; falling"
            " back to passing ref_audio/ref_text per invocation."
        )

    # rough speaking rate used to detect runaway generation
    chars_per_sec = 6.0 if base_lang(lang) in CJK_LANGS else 14.0

    for slide_num in requested_slides:
        txt_p = os.path.join(workspace_dir, text_filename(slide_num, lang))
        if not os.path.exists(txt_p):
            continue

        with open(txt_p, "r", encoding="utf-8") as f:
            spoken_text = apply_dictionary(f.read().strip(), dictionaries, lang, letter_map=letter_map)
        if not spoken_text:
            continue

        with open(os.path.join(workspace_dir, spoken_filename(slide_num, lang, model_label)),
                  "w", encoding="utf-8") as f:
            f.write(spoken_text)

        chunks = split_into_chunks(spoken_text)
        started = time.time()
        try:
            wavs = []
            sr = None

            def _gen_chunk(chunk_text):
                # Stop the model at a length the text cannot justify instead of letting a
                # runaway generation finish and then discarding it: the codec runs at a
                # fixed frame rate, so the budget follows from the expected duration.
                kw = {"max_new_tokens": chunk_token_budget(chunk_text, chars_per_sec)}
                if voice_clone_prompt is not None:
                    return model.generate_voice_clone(
                        text=chunk_text, language=language, voice_clone_prompt=voice_clone_prompt, **kw
                    )
                return model.generate_voice_clone(
                    text=chunk_text, language=language, ref_audio=ref_wav, ref_text=ref_txt, **kw
                )

            def _is_anomalous(chunk_text, wav, sample_rate):
                expected_sec = max(len(chunk_text) / chars_per_sec, 1.0)
                actual_sec = len(wav) / sample_rate
                return actual_sec > max(expected_sec * 3, 20)

            for chunk in chunks:
                w, sr = _gen_chunk(chunk)
                if _is_anomalous(chunk, w[0], sr):
                    logger.warning(
                        f"Slide {slide_num}: Chunk '{chunk[:20]}...' abnormally long"
                        f" ({len(w[0])/sr:.1f}s). Regenerating..."
                    )
                    w, sr = _gen_chunk(chunk)
                    if _is_anomalous(chunk, w[0], sr):
                        logger.error(
                            f"Slide {slide_num}: Chunk remains abnormally long"
                            f" ({len(w[0])/sr:.1f}s): '{chunk}'"
                        )
                wavs.append(w[0])
            combined = np.concatenate(wavs) if len(wavs) > 1 else wavs[0]

            wav_p = os.path.join(workspace_dir, "temp_qwen3.wav")
            sf.write(wav_p, combined, sr)

            audio = AudioSegment.from_file(wav_p)
            if enable_drc:
                audio = compress_dynamic_range(audio, threshold=drc_threshold, ratio=drc_ratio)
                audio = normalize(audio)
            audio.export(os.path.join(workspace_dir, audio_filename(slide_num, lang, model_label)),
                         format="ipod")
            stats.add(slide_num, time.time() - started, len(audio))
        except Exception as e:
            logger.error(f"Slide #{slide_num}: Qwen3-TTS could not generate the audio: {e}")
        finally:
            tmp_wav = os.path.join(workspace_dir, "temp_qwen3.wav")
            if os.path.exists(tmp_wav):
                os.remove(tmp_wav)
    stats.report()

_KANA_PUNCT_RE = re.compile(r"[\s、。，．,.!?！？「」『』（）()・…]")

def levenshtein_distance(a, b):
    """Minimum number of insertions, deletions and substitutions turning a into b."""
    if len(a) < len(b):
        a, b = b, a
    previous = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        current = [i]
        for j, cb in enumerate(b, 1):
            current.append(min(previous[j] + 1,
                               current[j - 1] + 1,
                               previous[j - 1] + (ca != cb)))
        previous = current
    return previous[-1]

def character_error_rate(reference, hypothesis):
    """CER = edit_distance(reference, hypothesis) / len(reference)."""
    if not reference:
        return 0.0 if not hypothesis else 1.0
    return levenshtein_distance(reference, hypothesis) / len(reference)

def kana_sequence_scores(kana_ref, kana_hyp):
    """Return (similarity, CER) for two katakana strings (punctuation/whitespace ignored)."""
    kana_ref = _KANA_PUNCT_RE.sub("", kana_ref)
    kana_hyp = _KANA_PUNCT_RE.sub("", kana_hyp)
    similarity = difflib.SequenceMatcher(None, kana_ref, kana_hyp, autojunk=False).ratio()
    return similarity, character_error_rate(kana_ref, kana_hyp)

def kana_scores(text_intended, text_asr):
    """Convert both texts to katakana with pyopenjtalk; return (similarity, CER, kana_intended, kana_asr)."""
    import pyopenjtalk
    kana_a = _KANA_PUNCT_RE.sub("", pyopenjtalk.g2p(text_intended, kana=True))
    kana_b = _KANA_PUNCT_RE.sub("", pyopenjtalk.g2p(text_asr, kana=True))
    similarity, cer = kana_sequence_scores(kana_a, kana_b)
    return similarity, cer, kana_a, kana_b

_NON_WORD_RE = re.compile(r"[\W_]+")

def normalize_for_comparison(text):
    """NFKC, case folding, and removal of punctuation, symbols and whitespace."""
    return _NON_WORD_RE.sub("", unicodedata.normalize("NFKC", text).casefold())

def difference_list(a, b, context=8, minimum=1):
    """The places where two sequences differ: (position, intended, recognized, before, after).

    This is what a reader of the report actually works with: the fragment that was meant
    and the fragment the ASR heard, in the order they occur, with a little context."""
    out = []
    for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        if tag != "equal" and max(i2 - i1, j2 - j1) >= minimum:
            out.append((i1, a[i1:i2], b[j1:j2], a[max(0, i1 - context):i1], a[i2:i2 + context]))
    return out


def difference_runs(a, b):
    """(number of differing stretches, characters in the longest one) for two sequences.

    The ratio-based scores say how much of a slide matched; these say how many places
    differ and how long the longest difference is, which is what points at a skipped
    phrase rather than at scattered recognition differences."""
    runs = [max(i2 - i1, j2 - j1)
            for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes()
            if tag != "equal"]
    return len(runs), max(runs, default=0)


def text_scores(text_intended, text_asr):
    """Character-level similarity and CER on normalized text; return (similarity, CER, norm_intended, norm_asr)."""
    a = normalize_for_comparison(text_intended)
    b = normalize_for_comparison(text_asr)
    similarity = difflib.SequenceMatcher(None, a, b, autojunk=False).ratio()
    return similarity, character_error_rate(a, b), a, b

def step_verify_audio(workspace_dir, requested_slides, lang, model_label,
                       asr_model_size="small", asr_device="cpu", threshold=0.85,
                       cer_threshold=None, max_difference=40, min_difference=4):
    logger.info("--- [Option: Verify] ASR round-trip check ---")
    is_ja = base_lang(lang) == "ja"
    try:
        from faster_whisper import WhisperModel
        if is_ja:
            import pyopenjtalk  # noqa: F401
    except ImportError as e:
        raise RuntimeError(f"verify needs a library that is not installed ({e}); "
                           "pip install 'pptx-narrator[verify]'") from e

    logger.info(f"Loading Whisper model ({asr_model_size}, device={asr_device})...")
    asr_model = WhisperModel(asr_model_size, device=asr_device, compute_type="int8")

    results, differences = [], []
    for slide_num in requested_slides:
        audio_p = os.path.join(workspace_dir, audio_filename(slide_num, lang, model_label))
        spoken_p = os.path.join(workspace_dir, spoken_filename(slide_num, lang, model_label))
        if not (os.path.exists(audio_p) and os.path.exists(spoken_p)):
            continue

        with open(spoken_p, "r", encoding="utf-8") as f:
            intended_text = f.read().strip()
        if not intended_text:
            continue

        try:
            segments, _ = asr_model.transcribe(audio_p, language=base_lang(lang))
            asr_text = "".join(seg.text for seg in segments).strip()
        except Exception as e:
            logger.error(f"Slide {slide_num}: transcription failed: {e}")
            continue

        try:
            if is_ja:
                score, cer, norm_intended, norm_asr = kana_scores(intended_text, asr_text)
            else:
                score, cer, norm_intended, norm_asr = text_scores(intended_text, asr_text)
        except Exception as e:
            logger.error(f"Slide {slide_num}: comparison failed: {e}")
            continue

        n_runs, worst_run = difference_runs(norm_intended, norm_asr)
        for pos, said, heard, before, after in difference_list(norm_intended, norm_asr, minimum=min_difference):
            differences.append((slide_num, max(len(said), len(heard)), pos, said, heard, before, after))

        # Latin-script words left in Japanese narration cannot be compared reliably as kana
        has_latin = is_ja and bool(re.search(r'[A-Za-z]{2,}', intended_text))
        long_difference = max_difference is not None and worst_run > max_difference
        failed = (score < threshold or long_difference
                  or (cer_threshold is not None and cer > cer_threshold))
        status = "ENGLISH" if has_latin else ("FLAGGED" if failed else "OK")

        results.append((slide_num, round(score, 4), round(cer, 4), status, n_runs, worst_run,
                        intended_text, asr_text, norm_intended, norm_asr))
        logger.info(f"Slide {slide_num}: similarity={score:.2f}, CER={cer:.2f}, "
                    f"{n_runs} difference(s), longest {worst_run} characters [{status}]")

    if not results:
        # The file names carry the engine and model, so looking for the wrong label is
        # the usual reason nothing is found. Say which labels the workspace does hold.
        present = sorted({m.group(1) for m in
                          (re.match(r"^slide_\d+(?:_[^.]+)?\.(.+)\.m4a$", n)
                           for n in os.listdir(workspace_dir)) if m})
        logger.info(f"No slides with both audio and spoken text for '{model_label}' in {lang} "
                    "-- nothing to verify.")
        if present:
            logger.info("The workspace holds audio for: " + ", ".join(present)
                        + ". Give the engine that produced it, e.g. --engine qwen3.")
        return

    report_p = os.path.join(workspace_dir, f"verify_report{lang_suffix(lang)}.{model_label}.csv")
    with open(report_p, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["slide", "similarity", "cer", "status", "differences", "longest_difference",
                    "intended_text", "asr_text", "intended_normalized", "asr_normalized"])
        for row in sorted(results, key=lambda r: r[1]):
            w.writerow(row)

    n_flagged = sum(1 for r in results if r[3] == "FLAGGED")
    n_latin = sum(1 for r in results if r[3] == "ENGLISH")
    criterion = f"similarity < {threshold}"
    if max_difference is not None:
        criterion += f", a difference longer than {max_difference} characters"
    if cer_threshold is not None:
        criterion += f" or CER > {cer_threshold}"
    logger.info(f"Done: {n_flagged}/{len(results)} slide(s) flagged for review ({criterion}).")
    if n_latin:
        logger.info(f"{n_latin} slide(s) contain un-converted Latin-script words and need a listen.")
    logger.info(f"Report saved to: {os.path.basename(report_p)} (sorted worst-first)")

    if differences:
        diff_p = os.path.join(workspace_dir, f"verify_differences{lang_suffix(lang)}.{model_label}.csv")
        with open(diff_p, "w", encoding="utf-8", newline="") as f:
            w = csv.writer(f)
            w.writerow(["slide", "length", "position", "intended", "recognized", "before", "after"])
            for row in sorted(differences, key=lambda r: (-r[1], r[0], r[2])):
                w.writerow(row)
        logger.info(f"{len(differences)} difference(s) of at least {min_difference} characters listed in: "
                    f"{os.path.basename(diff_p)} "
                    "(longest first)")

_P14_MEDIA_RE = re.compile(r'<(p14:media)\b([^>]*?)(/?)>(?:(.*?)</p14:media>)?', re.DOTALL)
_P14_PLAYBACK_CHILD_RE = re.compile(
    r'<p14:(?:trim|fade)\b[^>]*/>|<p14:(?:trim|fade)\b[^>]*>.*?</p14:(?:trim|fade)>|<p14:bmkLst\b.*?</p14:bmkLst>|<p14:bmkLst\b[^>]*/>',
    re.DOTALL)

def clear_media_playback_settings(slide_xml, media_rel_ids):
    """Remove trim, fade and bookmark settings from the p14:media elements of the given relationships.

    Returns (new_xml, number of media elements changed)."""
    changed = 0

    def _fix(m):
        nonlocal changed
        attrs, self_closing, body = m.group(2), m.group(3), m.group(4) or ""
        embed = re.search(r'r:embed="([^"]+)"', attrs)
        if self_closing or not embed or embed.group(1) not in media_rel_ids:
            return m.group(0)
        new_body = _P14_PLAYBACK_CHILD_RE.sub("", body)
        if new_body == body:
            return m.group(0)
        changed += 1
        return f"<p14:media{attrs}/>" if not new_body.strip() else f"<p14:media{attrs}>{new_body}</p14:media>"

    return _P14_MEDIA_RE.sub(_fix, slide_xml), changed

def _read_text(path):
    if path and os.path.exists(path):
        with open(path, "r", encoding="utf-8") as f:
            return f.read().strip()
    return ""

RECORDED_CHOICES = {"all": ("pointer", "events"), "pointer": ("pointer",), "events": ("events",), "none": ()}


def _slide_audio_matches(pkg_dir, slide_part, audio_path):
    """True if the slide already plays exactly this audio file."""
    slide_p = os.path.join(pkg_dir, *slide_part.split("/"))
    rels_p = os.path.join(os.path.dirname(slide_p), "_rels", os.path.basename(slide_p) + ".rels")
    if not os.path.exists(rels_p):
        return False
    with open(rels_p, encoding="utf-8") as f:
        rels_xml = f.read()
    want = _sha256_file(audio_path)
    for _, attrs in _rel_elements(rels_xml):
        if attrs.get("Type") not in (REL_AUDIO, REL_MEDIA) or attrs.get("TargetMode") == "External":
            continue
        target = os.path.normpath(os.path.join(os.path.dirname(slide_p), *attrs.get("Target", "").split("/")))
        if os.path.isfile(target) and _sha256_file(target) == want:
            return True
    return False


def _plain_note_lang(entry, text):
    """The language of a note without headings: the one extract recorded, else detected."""
    langs = list((entry.get("langs") or {}).keys())
    if len(langs) == 1:
        return langs[0]
    detected, confident = detect_language(text)
    return normalize_lang(detected) if confident else None


def _note_sections_xml(body):
    """The paragraphs of a notes text body, divided into sections by language."""
    sections, cur = [], None
    for p in body.findall(_A_NS + "p"):
        t = _paragraph_text(p)
        h = _heading_of(t)
        if h is not None:
            cur = {"lang": h[0], "info": h[1], "paras": [], "lines": [], "heading": True}
            sections.append(cur)
            continue
        if cur is None:
            cur = {"lang": None, "info": {}, "paras": [], "lines": [], "heading": False}
            sections.append(cur)
        cur["paras"].append(p)
        cur["lines"].append(t)
    return sections


def _pack_notes(prs, workspace_dir, requested_slides, lang, targets, update, overwrite, audio_lang,
                model_label=None):
    """Write the texts of the workspace into the notes, one section per language.

    For each language, the text of the workspace (Y) is compared with that part of the note
    in the deck (Z) and with what extract read or pack wrote (X): the same text is not
    written again; a text edited only in the workspace is written with --update or
    --overwrite; a note edited in the deck is written over only with --overwrite. A
    language the note does not have yet is added. Parts that are not written keep their
    formatting and struck-through text.
    """
    manifest = _load_manifest(workspace_dir)
    baseline = _load_note_baseline(workspace_dir)
    notes_written, protected, mismatched = [], [], []
    for i, slide in enumerate(prs.slides):
        n = i + 1
        if n not in requested_slides:
            continue
        ws_langs = [l for l in _text_langs_of_slide(workspace_dir, n)
                    if not lang or lang_suffix(l) == lang_suffix(lang)]
        if not ws_langs:
            if lang:
                logger.info(f"Slide #{n}: no {lang} text in the workspace (note left as is).")
            continue
        body = slide.notes_slide.notes_text_frame._txBody
        sections = _note_sections_xml(body)
        headed = any(s["heading"] for s in sections)
        entry = baseline.get(str(n), {})
        if not headed and sections and _section_text(sections[0]["lines"]):
            sections[0]["lang"] = _plain_note_lang(entry, _section_text(sections[0]["lines"]))
        sections = [s for s in sections if s["lang"] or _section_text(s["lines"])]
        by_lang = {lang_suffix(s["lang"]): s for s in sections if s["lang"]}
        written, added = [], []
        for l in ws_langs:
            name = text_filename(n, l)
            text = _read_text(os.path.join(workspace_dir, name))
            if not text:
                logger.warning(f"Slide #{n}: {name} is empty (not written).")
                continue
            y = text_fingerprint(text)
            rec_tr = manifest.get(str(n), {}).get(l) or {}
            if rec_tr.get("source_lang"):
                src_text = _read_text(os.path.join(workspace_dir, text_filename(n, rec_tr["source_lang"])))
                if src_text and text_fingerprint(src_text) != rec_tr.get("source_fingerprint"):
                    logger.warning(f"Slide #{n}: {name} was translated from an older version of "
                                   f"{text_filename(n, rec_tr['source_lang'])} (translate again with --update).")
            s = by_lang.get(lang_suffix(l))
            if s is None:
                if "text" in targets:
                    added.append({"lang": l, "info": {}, "paras": _new_paragraphs(text), "lines": [text],
                                  "heading": False, "written": True})
                    written.append((l, y))
                    logger.info(f"Slide #{n}: the {l} text ({name}) added to the note.")
                continue
            z = text_fingerprint(_section_text(s["lines"]))
            if z == y:
                logger.info(f"Slide #{n}: the {l} part of the note is already the same as {name}.")
                continue
            if "text" not in targets:
                mismatched.append(n)
                if lang_suffix(l) == lang_suffix(audio_lang or ""):
                    logger.warning(f"Slide #{n}: the note in the deck differs from the text the audio was "
                                   f"synthesized from ({name}); notes are not written with --data-type audio.")
                continue
            rec = (entry.get("langs") or {}).get(l) or {}
            known = {rec.get("extracted"), rec.get("packed")} - {None}
            if not headed:
                known |= {entry.get("extracted"), entry.get("packed")} - {None}
            if z in known:
                if not (update or overwrite):
                    mismatched.append(n)
                    logger.warning(f"Slide #{n}: {name} was edited in the workspace; the {l} part of the note "
                                   f"is not replaced (add --update or --overwrite).")
                    continue
            else:
                reason = ("there is no record of extract reading this note, so a hand edit cannot be ruled out"
                          if not entry else "the note was edited in the deck after extract")
                if not overwrite:
                    protected.append(n)
                    mismatched.append(n)
                    logger.warning(f"Slide #{n}: {reason}; the {l} part of the note is not overwritten, and it "
                                   f"differs from {name}. Re-extract after merging the edit, or use --overwrite "
                                   "to overwrite it.")
                    continue
                logger.warning(f"Slide #{n}: {reason}; overwriting it (--overwrite).")
            s["paras"] = _new_paragraphs(text)
            s["written"] = True
            written.append((l, y))
            logger.info(f"Slide #{n}: the {l} part of the note replaced with {name}.")
        # The text of the audio the slide plays goes on top, where the presenter reads.
        audio_first = None
        if audio_lang and "audio" in targets and os.path.exists(
                os.path.join(workspace_dir, audio_filename(n, audio_lang, model_label))):
            audio_first = lang_suffix(audio_lang)
        present = [s for s in sections + added if s["lang"]]
        misplaced = bool(audio_first and present and lang_suffix(present[0]["lang"]) != audio_first
                         and any(lang_suffix(s["lang"]) == audio_first for s in present)
                         and not any(s.get("written") for s in sections + added))
        if not written and not misplaced:
            continue
        # Newer parts above older ones: what this run wrote goes on top (the text changed
        # most recently first), and what it left keeps its place below, in the order the
        # note had -- which earlier runs laid out the same way. The text a presenter reads
        # first is therefore the one just written, typically the language of the narration.
        def _mtime(sec):
            p = os.path.join(workspace_dir, text_filename(n, sec["lang"]))
            return os.path.getmtime(p) if os.path.exists(p) else 0
        fresh = sorted([s for s in sections if s.get("written")] + added, key=_mtime, reverse=True)
        final = fresh + [s for s in sections if not s.get("written")]
        if audio_first:
            top = [s for s in final if s["lang"] and lang_suffix(s["lang"]) == audio_first]
            final = top + [s for s in final if s not in top]
        if misplaced:
            logger.info(f"Slide #{n}: the {audio_lang} part of the note moved to the top, above the other "
                        "languages, to go with the audio (its text is unchanged).")
        multi = sum(1 for s in final if s["lang"]) > 1
        for p in body.findall(_A_NS + "p"):
            body.remove(p)
        for k, s in enumerate(final):
            if multi and s["lang"]:
                # The heading describes the text under it: the text of the workspace where that
                # is what the note now holds, otherwise what the note's own heading said.
                ws_text = _read_text(os.path.join(workspace_dir, text_filename(n, s["lang"])))
                shown = _section_text(s["lines"]) if not s.get("written") else ws_text
                if ws_text and text_fingerprint(shown) == text_fingerprint(ws_text):
                    info = _section_info(workspace_dir, n, s["lang"], ws_text, manifest, baseline)
                else:
                    info = s["info"]
                body.extend(_new_paragraphs(section_heading(s["lang"], info)))
            body.extend(s["paras"])
            if multi and k < len(final) - 1:
                body.extend(_new_paragraphs(""))
        if body.find(_A_NS + "p") is None:
            body.extend(_new_paragraphs(""))
        entry = baseline.setdefault(str(n), {})
        for l, y in written:
            entry.setdefault("langs", {}).setdefault(l, {})["packed"] = y
        entry["packed"] = text_fingerprint(_slide_note_raw(slide))
        notes_written.append(n)
    if notes_written:
        _save_note_baseline(workspace_dir, baseline)
    return notes_written, protected, sorted(set(mismatched))


def _audio_languages(workspace_dir):
    """Languages of the narration audio in the workspace (any model)."""
    langs = set()
    for name in os.listdir(workspace_dir):
        m = re.match(r"^slide_\d+_([^.]+)\.(.+)\.m4a$", name)
        if m:
            langs.add(normalize_lang(m.group(1)))
    return sorted(langs)


AUDIO_SOURCES = "audio_sources.json"


def _load_audio_sources(workspace_dir):
    path = os.path.join(workspace_dir, AUDIO_SOURCES)
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            logger.warning(f"Could not read {path}: {e}")
    return {}


def record_audio_sources(workspace_dir, slides, lang, model_label, dictionary_paths, since):
    """Remember, for each audio file made in this run, the text (and dictionaries) it was
    made from, so that pack and verify can tell when the text has been edited since."""
    record = _load_audio_sources(workspace_dir)
    dict_fp = hashlib.sha1(b"".join(open(p, "rb").read() for p in (dictionary_paths or [])
                                    if p and os.path.exists(p))).hexdigest()[:10]
    made = []
    for n in slides:
        audio = audio_filename(n, lang, model_label)
        path = os.path.join(workspace_dir, audio)
        if not os.path.exists(path) or os.path.getmtime(path) < since:
            continue
        record[audio] = {"text": text_filename(n, lang),
                         "text_fingerprint": text_fingerprint(_read_text(os.path.join(workspace_dir, text_filename(n, lang)))),
                         "dictionary_fingerprint": dict_fp,
                         "made_at": time.strftime("%Y-%m-%dT%H:%M:%S%z", time.localtime(os.path.getmtime(path)))}
        made.append(n)
    if made:
        with open(os.path.join(workspace_dir, AUDIO_SOURCES), "w", encoding="utf-8") as f:
            json.dump(record, f, ensure_ascii=False, indent=1, sort_keys=True)
    return made


def audio_is_older_than_text(workspace_dir, s_num, lang, model_label, record=None):
    """True if the text of a slide was edited after its audio was made (None: no record)."""
    record = _load_audio_sources(workspace_dir) if record is None else record
    entry = record.get(audio_filename(s_num, lang, model_label))
    if not entry:
        return None
    text = _read_text(os.path.join(workspace_dir, text_filename(s_num, lang)))
    return text_fingerprint(text) != entry.get("text_fingerprint")


def step_pack_pptx(original_pptx, output_pptx, workspace_dir, requested_slides, lang, model_label,
                   targets=None, update=False, overwrite=False, audio_lang="same",
                   remove_recorded=("pointer", "events"), icon_outside=True, pause_ms=1000):
    """Write the texts (into the notes) and the narration audio of the workspace into a copy
    of the deck.

    lang: the language of the texts to write (None: every language in the workspace).
    audio_lang: the language of the audio to write ("same": lang; None: no audio).
    targets: subset of {"audio", "text"} (default: both).
    update / overwrite: see _pack_notes; update also leaves out audio the deck already has.
    """
    logger.info("--- [Option: Pack] Rebuilding PPTX ---")
    targets = set(targets or {"audio", "text"})
    if audio_lang == "same":
        audio_lang = lang
    prs = Presentation(original_pptx)
    notes_written, protected, mismatched = _pack_notes(prs, workspace_dir, requested_slides, lang, targets,
                                                       update, overwrite, audio_lang, model_label)

    tmp_pptx = os.path.join(workspace_dir, "tmp.pptx")
    prs.save(tmp_pptx)
    embedded, audio_same = 0, 0
    try:
        embedded, audio_same = _pack_audio(tmp_pptx, output_pptx, workspace_dir, requested_slides, targets,
                                           audio_lang, model_label, update, remove_recorded, icon_outside,
                                           pause_ms)
    finally:
        if os.path.exists(tmp_pptx):
            os.remove(tmp_pptx)
    _pack_summary(embedded, audio_same, notes_written, protected, mismatched, output_pptx)
    return embedded, notes_written, mismatched


def _pack_audio(tmp_pptx, output_pptx, workspace_dir, requested_slides, targets, audio_lang, model_label,
                update, remove_recorded, icon_outside, pause_ms):
    embedded, audio_same = 0, 0
    with tempfile.TemporaryDirectory() as tmpdir:
        with zipfile.ZipFile(tmp_pptx, 'r') as z:
            z.extractall(tmpdir)
        slide_w, slide_h = _slide_size(tmpdir)
        slide_parts = _slide_part_paths(tmpdir)
        if "audio" in targets and audio_lang:
            for s_num in requested_slides:
                if s_num > len(slide_parts):
                    continue
                name = audio_filename(s_num, audio_lang, model_label)
                m4a_p = os.path.join(workspace_dir, name)
                if not os.path.exists(m4a_p):
                    if os.path.exists(os.path.join(workspace_dir, text_filename(s_num, audio_lang))):
                        logger.warning(f"Slide #{s_num}: {name} was not found; no audio written.")
                    continue
                if audio_is_older_than_text(workspace_dir, s_num, audio_lang, model_label):
                    logger.warning(f"Slide #{s_num}: {text_filename(s_num, audio_lang)} was edited after {name} "
                                   f"was made; the audio does not say the text (synthesize it again).")
                if update and _slide_audio_matches(tmpdir, slide_parts[s_num - 1], m4a_p):
                    audio_same += 1
                    logger.info(f"Slide #{s_num}: the deck already has this audio (left as is, --update).")
                    continue
                if embed_slide_narration(tmpdir, s_num, m4a_p, len(AudioSegment.from_file(m4a_p)),
                                         slide_w, slide_h, slide_part=slide_parts[s_num - 1],
                                         remove_recorded=remove_recorded, icon_outside=icon_outside,
                                         pause_ms=pause_ms):
                    embedded += 1
                    logger.info(f"Slide #{s_num}: audio {name} written.")
                else:
                    logger.warning(f"Slide #{s_num}: audio {name} could not be written (see above).")
        _ensure_default_content_types(tmpdir, {"m4a": "audio/mp4", "png": "image/png"})
        _remove_unreferenced_media(tmpdir)
        archive_path = _zip_package(tmpdir, os.path.splitext(output_pptx)[0] + ".packing.zip")
        os.replace(archive_path, output_pptx)
    return embedded, audio_same


def _pack_summary(embedded, audio_same, notes_written, protected, mismatched, output_pptx):
    summary = f"Packed audio on {embedded} slide(s); wrote notes on {len(notes_written)} slide(s)."
    if audio_same:
        summary += f" {audio_same} slide(s) already had the same audio."
    logger.info(summary)
    if protected:
        logger.warning("Notes not overwritten because they were edited in the deck: "
                       + ", ".join(map(str, sorted(set(protected)))) + ".")
    if mismatched:
        logger.warning("Slides whose note does not match the texts of the workspace: "
                       + ", ".join(map(str, mismatched)) + ".")
    logger.info(f"Final output saved to: {output_pptx}")


# ------------------------------------------
# Embedding narration audio in slide XML
# ------------------------------------------
# The structure written here follows what PowerPoint itself writes for recorded slide
# narration: a p:pic with a:audioFile / p14:media, audio + media + icon relationships,
# a p:timing that starts the audio with the slide (isNarration="1"), and the slide
# transition's advance time (advTm) set to the audio length.
REL_AUDIO = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/audio"
REL_MEDIA = "http://schemas.microsoft.com/office/2007/relationships/media"
REL_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
_RELS_EMPTY = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
               '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>')
NARRATION_ICON = "pptx_narrator_audio.png"

_NARRATION_TIMING = (
    '<p:timing><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
    '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
    '<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0">'
    '<p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst><p:par><p:cTn id="4" fill="hold"><p:stCondLst>'
    '<p:cond delay="0"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="5" presetID="1" presetClass="mediacall" '
    'presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
    '<p:childTnLst><p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="1" fill="hold"/>'
    '<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd></p:childTnLst></p:cTn></p:par>'
    '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn>'
    '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
    '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
    '</p:seq><p:audio isNarration="1"><p:cMediaNode vol="80000" showWhenStopped="0"><p:cTn id="7" fill="hold" '
    'display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:endCondLst><p:cond evt="onStopAudio" '
    'delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:endCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/>'
    '</p:tgtEl></p:cMediaNode></p:audio></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'
)


def _rel_elements(rels_xml):
    return [(m.group(0), dict(re.findall(r'(\w+)="([^"]*)"', m.group(0))))
            for m in re.finditer(r'<Relationship\b[^>]*?/>', rels_xml)]


def _set_rel_target(rels_xml, rel_id, target):
    def _fix(m):
        attrs = dict(re.findall(r'(\w+)="([^"]*)"', m.group(0)))
        if attrs.get("Id") != rel_id:
            return m.group(0)
        return f'<Relationship Id="{rel_id}" Type="{attrs["Type"]}" Target="{target}"/>'
    return re.sub(r'<Relationship\b[^>]*?/>', _fix, rels_xml)


def _new_rel_id(rels_xml):
    used = {int(n) for n in re.findall(r'Id="rId(\d+)"', rels_xml)}
    n = 1
    while n in used:
        n += 1
    return f"rId{n}"


def _add_rel(rels_xml, rel_type, target):
    rel_id = _new_rel_id(rels_xml)
    rel = f'<Relationship Id="{rel_id}" Type="{rel_type}" Target="{target}"/>'
    return rels_xml.replace("</Relationships>", rel + "</Relationships>"), rel_id


def _slide_size(pkg_dir):
    try:
        with open(os.path.join(pkg_dir, "ppt", "presentation.xml"), encoding="utf-8") as f:
            m = re.search(r'<p:sldSz\b[^>]*\bcx="(\d+)"[^>]*\bcy="(\d+)"', f.read())
        if m:
            return int(m.group(1)), int(m.group(2))
    except OSError:
        pass
    return 12192000, 6858000


def _speaker_icon_png(size=64):
    """A small grey loudspeaker icon (only visible in the editor; hidden during the slide show)."""
    rows = []
    for y in range(size):
        row = bytearray([0])
        for x in range(size):
            u, v = x / size, y / size
            body = 0.18 <= u <= 0.36 and 0.38 <= v <= 0.62
            cone = 0.36 <= u <= 0.58 and abs(v - 0.5) <= 0.12 + (u - 0.36) * 1.1
            r = ((u - 0.58) ** 2 + (v - 0.5) ** 2) ** 0.5
            wave = u > 0.62 and abs(v - 0.5) < (u - 0.5) and (0.12 < r < 0.17 or 0.24 < r < 0.29)
            row += bytes((90, 90, 90, 255)) if (body or cone or wave) else bytes((0, 0, 0, 0))
        rows.append(bytes(row))
    import struct
    import zlib

    def chunk(tag, data):
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", size, size, 8, 6, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(b"".join(rows), 9)) + chunk(b"IEND", b""))


def _insert_after_transition_or_clrmap(slide_xml, fragment):
    m = re.search(r'<mc:AlternateContent\b(?:(?!</mc:AlternateContent>).)*?<p:transition\b.*?</mc:AlternateContent>',
                  slide_xml, re.DOTALL)
    if not m:
        m = re.search(r'<p:transition\b[^>]*/>|<p:transition\b.*?</p:transition>', slide_xml, re.DOTALL)
    if not m:
        m = re.search(r'</p:clrMapOvr>|<p:clrMapOvr\b[^>]*/>', slide_xml) or re.search(r'</p:cSld>', slide_xml)
    return slide_xml[:m.end()] + fragment + slide_xml[m.end():]


def _set_advance_time(slide_xml, dur_ms):
    if re.search(r'<p:transition\b', slide_xml):
        def _fix(m):
            tag = m.group(0)
            if 'advTm="' in tag:
                return re.sub(r'advTm="\d+"', f'advTm="{dur_ms}"', tag)
            return re.sub(r'^<p:transition\b', f'<p:transition advTm="{dur_ms}"', tag)
        return re.sub(r'<p:transition\b[^>]*>', _fix, slide_xml)
    return _insert_after_transition_or_clrmap(slide_xml, f'<p:transition advTm="{dur_ms}"/>')


def _slide_part_paths(pkg_dir):
    """Slide part paths (relative to the package root) in presentation order."""
    with open(os.path.join(pkg_dir, "ppt", "presentation.xml"), encoding="utf-8") as f:
        pres_xml = f.read()
    with open(os.path.join(pkg_dir, "ppt", "_rels", "presentation.xml.rels"), encoding="utf-8") as f:
        targets = {a.get("Id"): a.get("Target", "") for _, a in _rel_elements(f.read())}
    paths = []
    for rel_id in re.findall(r'<p:sldId\b[^>]*\br:id="([^"]+)"', pres_xml):
        target = targets.get(rel_id, "")
        target = target.lstrip("/") if target.startswith("/") else os.path.normpath(os.path.join("ppt", target))
        paths.append(target.replace(os.sep, "/"))
    return paths


# Data recorded while a slide show was presented: the laser-pointer path
# (p14:laserTraceLst, one time-stamped point per sample) and the media events of the
# recording (p14:showEvtLst: play, pause, seek, stop). Both are timed against the audio
# that was recorded with them, so they are meaningless once the audio is replaced.
_RECORDED_SHOW_DATA = {
    "pointer": ("<p14:laserTraceLst", "laser-pointer path"),
    "events": ("<p14:showEvtLst", "recorded playback events"),
}


def remove_recorded_show_data(slide_xml, kinds=("pointer", "events")):
    """Drop the p:ext elements holding the given kinds of recorded show data.

    Returns (new_xml, [descriptions of what was removed])."""
    removed = []
    markers = {_RECORDED_SHOW_DATA[k][0]: _RECORDED_SHOW_DATA[k][1] for k in kinds if k in _RECORDED_SHOW_DATA}

    out, pos = [], 0
    for m in re.finditer(r'<p:ext\b[^>]*[^/]>', slide_xml):
        if m.start() < pos:
            continue
        end, depth = m.end(), 1
        for t in re.finditer(r'<p:ext\b[^>]*[^/]>|</p:ext>', slide_xml[m.end():]):
            depth += 1 if t.group(0).startswith("<p:ext") else -1
            if depth == 0:
                end = m.end() + t.end()
                break
        element = slide_xml[m.start():end]
        what = next((w for marker, w in markers.items() if marker in element), None)
        if what:
            removed.append(what)
            out.append(slide_xml[pos:m.start()])
            pos = end
    out.append(slide_xml[pos:])
    slide_xml = "".join(out)
    if removed:
        slide_xml = re.sub(r'<p:extLst\s*>\s*</p:extLst>', "", slide_xml)
    return slide_xml, removed


ICON_GAP = 228600  # 0.25 inch between the slide edge and the icon parked next to it


def _picture_extent(pic_xml, default=812800):
    m = re.search(r'<a:ext\b[^>]*\bcx="(\d+)"[^>]*\bcy="(\d+)"', pic_xml)
    return (int(m.group(1)), int(m.group(2))) if m else (default, default)


def _set_picture_offset(pic_xml, x, y):
    """Move a picture; returns the picture unchanged if it carries no position."""
    if not re.search(r'<a:off\b', pic_xml):
        return pic_xml
    return re.sub(r'<a:off\b[^>]*/>', f'<a:off x="{x}" y="{y}"/>', pic_xml, count=1)


def embed_slide_narration(pkg_dir, s_num, audio_path, dur_ms, slide_w=12192000, slide_h=6858000, slide_part=None,
                          remove_recorded=("pointer", "events"), icon_outside=True, pause_ms=1000):
    """Put the narration audio into slide s_num of an unzipped PPTX package.

    Each slide gets its own media file, so slides that shared one audio clip (e.g. copied
    slides) no longer overwrite each other. An existing narration object is re-pointed to
    it and its trim, fade and bookmarks are removed; a slide without audio gets a new
    auto-playing narration object. Returns "replaced", "inserted" or None.
    """
    slide_part = slide_part or f"ppt/slides/slide{s_num}.xml"
    slide_p = os.path.join(pkg_dir, *slide_part.split("/"))
    rels_p = os.path.join(os.path.dirname(slide_p), "_rels", os.path.basename(slide_p) + ".rels")
    if not os.path.exists(slide_p):
        return None
    with open(slide_p, encoding="utf-8") as f:
        slide_xml = f.read()
    rels_xml = _RELS_EMPTY
    if os.path.exists(rels_p):
        with open(rels_p, encoding="utf-8") as f:
            rels_xml = f.read()

    slide_xml, removed = remove_recorded_show_data(slide_xml, remove_recorded)
    for what in removed:
        logger.info(f"Slide #{s_num}: removed the {what} of the previous recording.")
    if "<p:contentPart" in slide_xml:
        logger.warning(f"Slide #{s_num}: the slide contains ink annotations; they are kept and may no longer "
                       "match the new narration.")

    media_dir = os.path.join(pkg_dir, "ppt", "media")
    os.makedirs(media_dir, exist_ok=True)
    media_name = f"pptx_narrator_slide{s_num}.m4a"
    shutil.copy(audio_path, os.path.join(media_dir, media_name))
    target = f"../media/{media_name}"

    # Prefer the object PowerPoint marks as narration; otherwise take the first audio object.
    narration_ids = set(re.findall(r'<p:audio\b[^>]*isNarration="1".*?<p:spTgt spid="(\d+)"', slide_xml, re.DOTALL))
    pic = None
    for m in re.finditer(r'<p:pic\b.*?</p:pic>', slide_xml, re.DOTALL):
        if '<a:audioFile' in m.group(0) or '<p14:media' in m.group(0):
            pic_id = re.search(r'<p:cNvPr\b[^>]*\bid="(\d+)"', m.group(0))
            if pic is None or (pic_id and pic_id.group(1) in narration_ids):
                pic = m.group(0)
            if pic_id and pic_id.group(1) in narration_ids:
                break

    if pic is not None:
        rel_ids = set(re.findall(r'<a:audioFile\b[^>]*r:link="([^"]+)"', pic))
        rel_ids |= set(re.findall(r'<p14:media\b[^>]*r:embed="([^"]+)"', pic))
        for rel_id in rel_ids:
            rels_xml = _set_rel_target(rels_xml, rel_id, target)
        if icon_outside:
            cx, cy = _picture_extent(pic)
            moved = _set_picture_offset(pic, -(cx + ICON_GAP), 0)
            if moved != pic:
                slide_xml = slide_xml.replace(pic, moved, 1)
                pic = moved
                logger.info(f"Slide #{s_num}: moved the audio icon next to the slide, out of the visible area.")
        slide_xml, n_cleared = clear_media_playback_settings(slide_xml, rel_ids)
        if n_cleared:
            logger.info(f"Slide #{s_num}: removed trim/fade/bookmark settings of the previous audio.")
        status = "replaced"
    else:
        if '<p:timing' in slide_xml:
            logger.warning(f"Slide #{s_num}: the slide has animations but no audio object; narration not inserted "
                           "(insert any audio clip on this slide in PowerPoint and pack again).")
            return None
        icon_p = os.path.join(media_dir, NARRATION_ICON)
        if not os.path.exists(icon_p):
            with open(icon_p, "wb") as f:
                f.write(_speaker_icon_png())
        rels_xml, rid_media = _add_rel(rels_xml, REL_MEDIA, target)
        rels_xml, rid_audio = _add_rel(rels_xml, REL_AUDIO, target)
        rels_xml, rid_icon = _add_rel(rels_xml, REL_IMAGE, f"../media/{NARRATION_ICON}")
        spid = max([int(i) for i in re.findall(r'<p:cNvPr\b[^>]*\bid="(\d+)"', slide_xml)] + [1]) + 1
        size = 812800
        if icon_outside:
            x, y = -(size + ICON_GAP), 0
        else:
            x, y = max(slide_w - size - 215900, 0), max(slide_h - size - 215900, 0)
        pic_xml = (
            f'<p:pic><p:nvPicPr><p:cNvPr id="{spid}" name="Narration {spid}"/><p:cNvPicPr><a:picLocks noChangeAspect="1"/>'
            f'</p:cNvPicPr><p:nvPr><a:audioFile r:link="{rid_audio}"/><p:extLst><p:ext uri="{{DAA4B4D4-6D71-4841-9C94-3DE7FCFB9230}}">'
            f'<p14:media xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" r:embed="{rid_media}"/>'
            f'</p:ext></p:extLst></p:nvPr></p:nvPicPr><p:blipFill><a:blip r:embed="{rid_icon}"/><a:stretch><a:fillRect/>'
            f'</a:stretch></p:blipFill><p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{size}" cy="{size}"/></a:xfrm>'
            f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>'
        )
        slide_xml = slide_xml.replace("</p:spTree>", pic_xml + "</p:spTree>", 1)
        slide_xml = _set_advance_time(slide_xml, dur_ms + pause_ms)
        slide_xml = _insert_after_transition_or_clrmap(slide_xml, _NARRATION_TIMING.replace("{spid}", str(spid)))
        status = "inserted"

    if status == "replaced":
        slide_xml = _set_advance_time(slide_xml, dur_ms + pause_ms)
    with open(slide_p, "w", encoding="utf-8") as f:
        f.write(slide_xml)
    os.makedirs(os.path.dirname(rels_p), exist_ok=True)
    with open(rels_p, "w", encoding="utf-8") as f:
        f.write(rels_xml)
    logger.info(f"Slide #{s_num}: narration {status} ({dur_ms} ms).")
    return status


def _ensure_default_content_types(pkg_dir, defaults):
    path = os.path.join(pkg_dir, "[Content_Types].xml")
    with open(path, encoding="utf-8") as f:
        xml = f.read()
    for ext, ctype in defaults.items():
        if not re.search(rf'<Default\b[^>]*Extension="{ext}"', xml, re.IGNORECASE):
            xml = xml.replace("<Default ", f'<Default Extension="{ext}" ContentType="{ctype}"/><Default ', 1)
    with open(path, "w", encoding="utf-8") as f:
        f.write(xml)


def _remove_unreferenced_media(pkg_dir):
    media_dir = os.path.join(pkg_dir, "ppt", "media")
    if not os.path.isdir(media_dir):
        return
    referenced = set()
    for root, _, files in os.walk(pkg_dir):
        for name in files:
            if name.endswith(".rels"):
                with open(os.path.join(root, name), encoding="utf-8") as f:
                    referenced.update(os.path.basename(t) for t in re.findall(r'Target="([^"]+)"', f.read()))
    removed = [name for name in os.listdir(media_dir) if name not in referenced]
    for name in removed:
        os.remove(os.path.join(media_dir, name))
    if removed:
        ct_path = os.path.join(pkg_dir, "[Content_Types].xml")
        with open(ct_path, encoding="utf-8") as f:
            ct_xml = f.read()
        for name in removed:
            ct_xml = re.sub(rf'<Override\b[^>]*PartName="/ppt/media/{re.escape(name)}"[^>]*/>', "", ct_xml)
        with open(ct_path, "w", encoding="utf-8") as f:
            f.write(ct_xml)


def _zip_package(pkg_dir, zip_path):
    """Zip an unzipped OOXML package with [Content_Types].xml as the first entry."""
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(os.path.join(pkg_dir, "[Content_Types].xml"), "[Content_Types].xml")
        for root, _, files in os.walk(pkg_dir):
            for name in sorted(files):
                full = os.path.join(root, name)
                arc = os.path.relpath(full, pkg_dir).replace(os.sep, "/")
                if arc != "[Content_Types].xml":
                    z.write(full, arc)
    return zip_path

# ==========================================
# Main CLI
# ==========================================
DESCRIPTION = """PPTX-Narrator: automated narration of PowerPoint presenter notes

Usage:
  pptx-narrator WS COMMAND [INPUT] [OUTPUT] [OPTIONS]
  pptx-narrator WS COMMAND --help       the options of one command

WS is the workspace directory of one deck. INPUT and OUTPUT are the files outside
WS that a command reads or writes: the deck whose notes are taken in, the dictionary
of scan, and for pack the deck and the copy to save. The files inside WS are chosen
with --lang, --slides and the model options, not by path.

Options may also be written with underscores (--in_lang) and kept in a TOML
configuration file; see --config.
"""


# Configuration and execution-state files are deliberately separate.  The config
# is user-authored; the state is maintained by the program and records the last
# explicit inputs and their identities.
STATE_FILE = ".pptx_narrator_state.json"
RESOLVED_CONFIG_FILE = ".pptx_narrator_resolved.toml"


def _load_toml(path):
    if not path:
        return {}
    try:
        try:
            import tomllib
        except ModuleNotFoundError:
            import tomli as tomllib
    except ModuleNotFoundError:
        raise RuntimeError("TOML configuration requires Python 3.11+ or the 'tomli' package")
    with open(path, "rb") as f:
        return _normalize_config_keys(tomllib.load(f))


def _normalize_config_keys(data):
    """Accept hyphenated keys in the configuration file (dict-file as well as dict_file).

    TOML allows both spellings, and the command line accepts both, so the file should too.
    """
    if not isinstance(data, dict):
        return data
    return {str(k).replace("-", "_"): _normalize_config_keys(v) for k, v in data.items()}


def _deep_update(base, update):
    for key, value in update.items():
        if isinstance(value, dict) and isinstance(base.get(key), dict):
            _deep_update(base[key], value)
        else:
            base[key] = value
    return base


def _load_config(path):
    """Load explicit config, or the conventional local config when present."""
    if path:
        if not os.path.exists(path):
            raise RuntimeError(f"config file not found: {path}")
        return _load_toml(path), os.path.abspath(path)
    default = os.path.abspath("pptx_narrator.toml")
    if os.path.exists(default):
        return _load_toml(default), default
    return {}, None


def _load_state(path=STATE_FILE):
    if not os.path.exists(path):
        return {"version": 1, "commands": {}}
    try:
        with open(path, "r", encoding="utf-8") as f:
            state = json.load(f)
        state.setdefault("version", 1)
        state.setdefault("commands", {})
        return state
    except Exception as e:
        raise RuntimeError(f"could not read state file '{path}': {e}")


def _save_state(state, path=STATE_FILE):
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False, indent=2, sort_keys=True)
    os.replace(tmp, path)


def _sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _file_role(name):
    """What a workspace file is, for the record of a directory input."""
    if name.endswith(".spoken.txt"):
        return "spoken text"
    if name.endswith(".stale.txt"):
        return "set-aside narration"
    if name.startswith("verify_report"):
        return "verification report"
    if name.startswith("verify_differences"):
        return "verification differences"
    if name == TRANSLATION_MANIFEST:
        return "translation manifest"
    if _TEXT_FILE_RE.match(name):
        return "note text"
    if name.lower().endswith((".m4a", ".wav")):
        return "audio"
    if name.lower().endswith(".csv"):
        return "dictionary"
    if name.lower().endswith(".pptx"):
        return "deck"
    return "other"


def _input_snapshot(path, command):
    """Return a content-identity snapshot for a file or a command-relevant directory."""
    path = os.path.abspath(path)
    if os.path.isfile(path):
        return {
            "path": path,
            "kind": "file",
            "mtime": os.path.getmtime(path),
            "sha256": _sha256_file(path),
        }
    if os.path.isdir(path):
        extensions = {
            "scan": {".txt"},
            "translate": {".txt"},
            "synthesize": {".txt"},
            "verify": {".m4a", ".txt"},
        }.get(command, None)
        files = []
        for name in sorted(os.listdir(path)):
            full = os.path.join(path, name)
            if not os.path.isfile(full):
                continue
            if extensions is not None and os.path.splitext(name)[1].lower() not in extensions:
                continue
            files.append({
                "path": os.path.abspath(full),
                "role": _file_role(name),
                "mtime": os.path.getmtime(full),
                "sha256": _sha256_file(full),
            })
        return {"path": path, "kind": "directory", "files": files}
    raise RuntimeError(f"input not found: {path}")


def _snapshot_matches(saved, current):
    if not saved or saved.get("kind") != current.get("kind") or saved.get("path") != current.get("path"):
        return False
    if current["kind"] == "file":
        return saved.get("sha256") == current.get("sha256")
    old = {x["path"]: x.get("sha256") for x in saved.get("files", [])}
    new = {x["path"]: x.get("sha256") for x in current.get("files", [])}
    return old == new


def _record_input(state, command, snapshot, extra=None):
    entry = dict(extra or {})
    entry["input"] = snapshot
    state.setdefault("commands", {})[command] = entry


def _toml_scalar(value):
    if isinstance(value, bool):
        return "true" if value else "false"
    if value is None:
        return None
    if isinstance(value, str):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, list):
        vals = [_toml_scalar(x) for x in value]
        if any(x is None for x in vals):
            return None
        return "[" + ", ".join(vals) + "]"
    return None


def _write_toml_table(f, data, prefix=()):
    scalars = []
    tables = []
    for key, value in data.items():
        if isinstance(value, dict):
            tables.append((key, value))
        else:
            scalars.append((key, value))
    if prefix:
        f.write("[" + ".".join(prefix) + "]\n")
    for key, value in scalars:
        rendered = _toml_scalar(value)
        if rendered is not None:
            f.write(f"{key} = {rendered}\n")
    if scalars and tables:
        f.write("\n")
    for i, (key, value) in enumerate(tables):
        _write_toml_table(f, value, prefix + (key,))
        if i != len(tables) - 1:
            f.write("\n")


def _save_resolved_config(resolved, path=RESOLVED_CONFIG_FILE):
    with open(path, "w", encoding="utf-8") as f:
        f.write("# Resolved PPTX-Narrator configuration; all values are effective values.\n")
        _write_toml_table(f, resolved)


def _rel_to_workspace(path, workspace_dir):
    """A path inside the workspace, written relative to it.

    A workspace is moved and copied often, so what it records about its own
    contents must not depend on where it happens to sit. Anything outside it
    keeps its absolute path, because nothing else identifies it.
    """
    if not path or not workspace_dir:
        return path
    abs_path = os.path.abspath(path)
    abs_ws = os.path.abspath(workspace_dir)
    if abs_path == abs_ws:
        return "."
    if abs_path.startswith(abs_ws + os.sep):
        return os.path.relpath(abs_path, abs_ws)
    return abs_path


def _relativize_snapshot(snapshot, workspace_dir):
    """Rewrite the paths of an input record relative to the workspace."""
    if not snapshot:
        return snapshot
    out = dict(snapshot)
    out["path"] = _rel_to_workspace(out.get("path"), workspace_dir)
    if out.get("files"):
        out["files"] = [dict(f, path=_rel_to_workspace(f.get("path"), workspace_dir))
                        for f in out["files"]]
    return out


def _workspace_path(workspace_dir, filename):
    """Return the path of workspace-owned execution metadata."""
    return os.path.join(workspace_dir, filename)


def _append_history(workspace_dir, entry):
    """Append a compact, machine-readable record of one completed command."""
    path = _workspace_path(workspace_dir, ".pptx_narrator_history.jsonl")
    with open(path, "a", encoding="utf-8") as f:
        f.write(json.dumps(entry, ensure_ascii=False, sort_keys=True) + "\n")


def _config_for_command(config, command):
    common = config.get("common", {}) if isinstance(config.get("common", {}), dict) else {}
    section = config.get(command, {}) if isinstance(config.get(command, {}), dict) else {}
    out = {}
    _deep_update(out, common)
    _deep_update(out, section)
    return out


def _normalize_argv(argv):
    """Rewrite each --underscore_option token to its --hyphen-option spelling.

    Every option used to be registered twice, once per spelling, so that both
    were accepted; on the command line that put two matching option strings
    in front of argparse's abbreviation matching, and a short prefix such as
    --dic became ambiguous between --dict-file and --dict_file -- one option,
    "ambiguous" only because of how it was registered. Normalizing the argv
    instead means each option is registered once, and a prefix is ambiguous
    only when it is genuinely shared by two different options.
    """
    if argv is None:
        argv = sys.argv[1:]
    out = []
    for tok in argv:
        if tok.startswith("--") and tok != "--":
            name, eq, value = tok.partition("=")
            tok = name.replace("_", "-") + eq + value
        out.append(tok)
    return out


class ArgumentParser(argparse.ArgumentParser):
    """An ArgumentParser that accepts underscores in place of hyphens.

    Subparsers created with add_subparsers() default to the same class, so
    this covers every command without repeating the override.
    """

    def parse_known_args(self, args=None, namespace=None):
        return super().parse_known_args(_normalize_argv(args), namespace)


class HelpFormatter(argparse.RawTextHelpFormatter):
    """Raw-text help, without the redundant subparsers metavar line."""

    def _format_action(self, action):
        text = super()._format_action(action)
        if isinstance(action, argparse._SubParsersAction):
            # Drop the metavar line argparse prints above the list of commands:
            # the group is already titled, and the line says nothing.
            text = text.split("\n", 1)[1]
        return text


def _add(group, *names, **kwargs):
    """Register an option (the underscore spelling is accepted via _normalize_argv).

    A switch (store_true) also gets a --no-... counterpart, so that a value set to true
    in the configuration file can be taken back on the command line.
    """
    group.add_argument(*names, **kwargs)
    if kwargs.get("action") == "store_true":
        off = dict(kwargs)
        off["action"] = "store_false"
        off["help"] = argparse.SUPPRESS
        group.add_argument("--no-" + names[0][2:], **off)


def _add_common_options(parser, config_values):
    parser.add_argument("--config", default=None,
                        help="TOML configuration file (default: ./pptx_narrator.toml when present)")
    parser.add_argument("--version", action="version", version=f"%(prog)s {__version__}")
    # These are intentionally not parser defaults: config values and built-in defaults
    # must be distinguishable when resolving the three configuration layers.
    parser.set_defaults(_config_values=config_values)


def _apply_cli_config_defaults(namespace, config_values):
    """Apply command-section values only where argparse did not receive an explicit value."""
    for key, value in config_values.items():
        if hasattr(namespace, key) and getattr(namespace, key) is None:
            setattr(namespace, key, value)


COMMANDS = ("extract", "scan", "translate", "synthesize", "verify", "pack", "history")

SLIDES_HELP = "Slide selection, e.g. 1-5 or 1,3,5- (default: every slide)"


def _add_lang_options(p, in_help, out_help=None, lang_help=None):
    """--lang, --in-lang and (for translate) --out-lang.

    --lang is shorthand for giving --in-lang and --out-lang the same language, so
    that a command that reads and writes one language needs only one option.
    """
    _add(p, "--lang", dest="lang", type=normalize_lang, default=None,
         help=lang_help or "Language of the text (sets --in-lang" + (" and --out-lang)" if out_help else ")"))
    _add(p, "--in-lang", dest="in_lang", type=normalize_lang, default=None, help=in_help)
    if out_help:
        _add(p, "--out-lang", dest="out_lang", type=normalize_lang, default=None, help=out_help)


def _add_overwrite_options(p, update_help, overwrite_help):
    _add(p, "--update", dest="update", action="store_true", default=None, help=update_help)
    _add(p, "--overwrite", dest="overwrite", action="store_true", default=None, help=overwrite_help)


def _add_model_options(p, what):
    _add(p, "--engine", choices=["gpt_sovits", "qwen3"], default=None,
         help=f"TTS engine of the audio {what} (default: qwen3)")
    _add(p, "--model", default=None, help="GPT-SoVITS model (default: v2ProPlus)")
    _add(p, "--qwen3-model-size", dest="qwen3_model_size", choices=["0.6B", "1.7B"], default=None,
         help="Qwen3-TTS model size (default: 1.7B)")


def build_parser(config_values=None):
    parser = ArgumentParser(
        prog="pptx-narrator",
        usage="pptx-narrator WS COMMAND [INPUT] [OUTPUT] [options]",
        description=DESCRIPTION,
        epilog="WS comes first and COMMAND second. The options of a command may be given\n"
               "anywhere after COMMAND. Files inside WS are chosen with --lang, --slides and\n"
               "the model options, not by path.",
        formatter_class=HelpFormatter,
    )
    parser.add_argument("workspace", metavar="WS",
                        help="Workspace directory holding the texts and audio of one deck")
    _add_common_options(parser, config_values or {})

    sub = parser.add_subparsers(dest="command", metavar="COMMAND", title="Commands")
    # The commands are what a reader is looking for; show them above the
    # options that apply to all of them.
    parser._action_groups.insert(1, parser._action_groups.pop())

    # ---- extract ---------------------------------------------------------
    p = sub.add_parser("extract", help="Extract presenter notes from a PPTX into WS",
                       description="Extract the presenter notes of a PPTX into one text file per slide,\n"
                                   "named after the slide and the language of the note. WS is created\n"
                                   "when it does not exist.",
                       formatter_class=HelpFormatter)
    p.add_argument("deck", metavar="DECK", help="PPTX file whose notes are extracted")
    _add_lang_options(p, "Language(s) of notes to extract, e.g. ja or ja,en\n(default: every language found in the deck)",
                      lang_help="Same as --in-lang")
    _add_overwrite_options(p, "Take in again a note that changed in the deck, where its text\n"
                              "was not edited in WS since it was written",
                           "Take in every selected note, replacing texts edited in WS")
    _add(p, "--slides", dest="slides", default=None, help=SLIDES_HELP)

    # ---- scan ------------------------------------------------------------
    p = sub.add_parser("scan", help="Collect terms of the texts into a dictionary",
                       description="Scan the texts of WS for technical terms and write them into a\n"
                                   "dictionary, where their reading or translation can be corrected by hand.",
                       formatter_class=HelpFormatter)
    p.add_argument("dictionary", metavar="DICT",
                   help="Dictionary CSV to create, or to add to with --append")
    _add_lang_options(p, "Language of the texts to scan")
    _add(p, "--dict-file", dest="dict_file", action="append", default=None,
         help="Other dictionary CSV(s) whose terms need not be proposed again;\nrepeatable (read only)")
    _add(p, "--append", dest="append", action="store_true", default=None,
         help="Add new terms to an existing DICT")
    _add(p, "--overwrite", dest="overwrite", action="store_true", default=None,
         help="Make DICT again from scratch")
    _add(p, "--scan-compounds", dest="scan_compounds", action="store_true", default=None,
         help="With Japanese input, propose Japanese compounds in the dictionary")
    _add(p, "--slides", dest="slides", default=None, help=SLIDES_HELP)

    # ---- translate -------------------------------------------------------
    p = sub.add_parser("translate", help="Write the texts in another language",
                       description="Translate the texts of WS from --in-lang into --out-lang, leaving the\n"
                                   "result as a text file that can be reviewed before it is synthesized.",
                       formatter_class=HelpFormatter)
    _add_lang_options(p, "Language of the texts to translate", "Language of the translation")
    _add(p, "--dict-file", dest="dict_file", action="append", default=None,
         help="Translation dictionary CSV(s): terms of the notes and how to\ntranslate them, given to the model as instructions")
    _add(p, "--translate-model", dest="translate_model", default=None,
         help="Hugging Face id of the translation model (default: Qwen/Qwen3-4B;\nsee README for alternatives)")
    _add(p, "--translate-device", dest="translate_device", default=None,
         help="auto / cuda:0 / mps / cpu (default: auto)")
    _add_overwrite_options(p, "Translate again where the source text has changed since it was\n"
                              "translated (a translation edited since is kept)",
                           "Translate again every selected slide, replacing existing translations")
    _add(p, "--slides", dest="slides", default=None, help=SLIDES_HELP)

    # ---- synthesize ------------------------------------------------------
    p = sub.add_parser("synthesize", help="Generate voice-cloned narration",
                       description="Generate narration audio from the texts of WS, in a voice cloned from\n"
                                   "a short reference recording. Existing audio is replaced.",
                       formatter_class=HelpFormatter)
    _add_lang_options(p, "Language of the texts to read")
    _add(p, "--dict-file", dest="dict_file", action="append", default=None,
         help="Dictionary CSV(s) applied before synthesis")
    _add(p, "--letter-map", dest="letter_map", default=None,
         help="JSON mapping of letters to readings")
    _add(p, "--ref-wav", dest="ref_wav", default=None, help="Reference recording (.wav file)")
    _add(p, "--ref-text", dest="ref_text", default=None, help="Text file with the transcript of --ref-wav")
    _add(p, "--ref-lang", dest="ref_lang", type=normalize_lang, default=None,
         help="Language of reference recording for GPT-SoVITS")
    _add_model_options(p, "to generate")
    _add(p, "--api-url", dest="api_url", default=None, help="GPT-SoVITS API server URL")
    _add(p, "--qwen3-device", dest="qwen3_device", default=None, help="auto / cuda:0 / mps / cpu")
    _add(p, "--enable-drc", dest="enable_drc", action="store_true", default=None,
         help="Even out the loudness of the generated audio (default: off)")
    _add(p, "--drc-threshold", dest="drc_threshold", type=float, default=None,
         help="Level in dBFS above which --enable-drc compresses (default: -20.0)")
    _add(p, "--drc-ratio", dest="drc_ratio", type=float, default=None,
         help="Compression ratio used by --enable-drc (default: 3.0)")
    _add(p, "--slides", dest="slides", default=None, help=SLIDES_HELP)

    # ---- verify ----------------------------------------------------------
    p = sub.add_parser("verify", help="Check the generated narration with ASR",
                       description="Transcribe the generated narration and compare it with the text it came\n"
                                   "from, to report the slides most likely to be misread.",
                       formatter_class=HelpFormatter)
    _add_lang_options(p, "Language of the narration")
    _add_model_options(p, "to check")
    _add(p, "--asr-model", dest="asr_model", default=None,
         help="faster-whisper model used for the check (default: small)")
    _add(p, "--asr-device", dest="asr_device", default=None,
         help="auto / cuda / cpu (default: cpu)")
    _add(p, "--verify-threshold", dest="verify_threshold", type=float, default=None,
         help="Flag a slide whose similarity falls below this (0-1, default: 0.85)")
    _add(p, "--min-difference", dest="min_difference", type=int, default=None,
         help="Shortest difference still listed, in characters (default: 4)")
    _add(p, "--max-difference", dest="max_difference", type=int, default=None,
         help="Flag a slide with one stretch of disagreement longer than this\n(characters, default: 40; 0 disables)")
    _add(p, "--cer-threshold", dest="cer_threshold", type=float, default=None,
         help="Also flag a slide whose character error rate exceeds this\n(default: off)")
    _add(p, "--slides", dest="slides", default=None, help=SLIDES_HELP)

    # ---- pack ------------------------------------------------------------
    p = sub.add_parser("pack", help="Write the narration of WS into a copy of a PPTX",
                       description="Write the texts (into the notes) and the generated audio of WS into a\n"
                                   "copy of DECK, saved as OUT, which then plays by itself and can be\n"
                                   "exported as a video. DECK itself is not changed.",
                       formatter_class=HelpFormatter)
    p.add_argument("deck", metavar="DECK", help="PPTX file to write into (it is not changed)")
    p.add_argument("out_deck", metavar="OUT", help="PPTX file to save the result as")
    _add(p, "--lang", dest="lang", type=normalize_lang, default=None, help="Language of the narration to write")
    _add(p, "--data-type", dest="data_type", choices=PACK_TARGETS, default=None,
         help="What to write: text (into the notes), audio, or all (default: all)")
    _add_model_options(p, "to write")
    _add_overwrite_options(p, "Write where WS differs from the deck; also allows an existing OUT",
                           "Write everything selected, even over notes edited in the deck;\n"
                           "also allows an existing OUT")
    _add(p, "--slide-pause", dest="slide_pause", type=float, default=None,
         help="Seconds between the end of the narration and the automatic\nslide advance (default: 1.0)")
    _add(p, "--keep-audio-icon", dest="keep_audio_icon", action="store_true", default=None,
         help="Leave the audio icon on the slide instead of parking it outside\nthe visible area (default: off)")
    _add(p, "--remove-recorded", dest="remove_recorded", choices=["all", "pointer", "events", "none"],
         default=None,
         help="Settings of a previous recording to remove: trim/fade/bookmarks,\nlaser-pointer path, playback events (default: all)")
    _add(p, "--slides", dest="slides", default=None, help=SLIDES_HELP)

    # ---- history --------------------------------------------------------
    p = sub.add_parser("history", help="List the commands run in WS",
                       description="List the commands run in WS, oldest first, with the files they read\n"
                                   "and wrote.",
                       formatter_class=HelpFormatter)
    _add(p, "--dates", dest="dates", action="store_true", default=None, help="Show when each command was run")
    _add(p, "--slides", dest="slides", default=None, help=argparse.SUPPRESS)

    # --config is read from anywhere in the command line; accept it after the
    # command as well, which is where people naturally write it.
    for _sp in sub.choices.values():
        _sp.add_argument("--config", default=None, help=argparse.SUPPRESS)

    return parser


def _defaults():
    return {
        "extract": {"slides": None, "update": False, "overwrite": False},
        "scan": {"scan_compounds": False, "slides": None, "dict_file": None,
                 "append": False, "overwrite": False},
        "translate": {"dict_file": None, "slides": None, "update": False, "overwrite": False,
                      "translate_model": TRANSLATE_MODEL_DEFAULT, "translate_device": "auto"},
        "synthesize": {
            "engine": "qwen3", "ref_lang": "ja", "api_url": "http://127.0.0.1:9880/", "slides": None,
            "model": "v2ProPlus", "qwen3_model_size": "1.7B", "qwen3_device": "auto",
            "enable_drc": False, "drc_threshold": -20.0, "drc_ratio": 3.0, "dict_file": None,
            "letter_map": None,
        },
        "verify": {
            "model": "v2ProPlus", "engine": "qwen3", "qwen3_model_size": "1.7B",
            "asr_model": "small", "asr_device": "cpu", "verify_threshold": 0.85,
            "min_difference": 4, "max_difference": 40, "cer_threshold": None, "slides": None,
        },
        "history": {"dates": False},
        "pack": {
            "model": "v2ProPlus", "engine": "qwen3", "qwen3_model_size": "1.7B",
            "data_type": "all", "slides": None, "update": False, "overwrite": False,
            "slide_pause": 1.0, "keep_audio_icon": False, "remove_recorded": "all",
        },
    }


# Positional arguments name files; they are recorded as paths, not as settings.
_POSITIONAL_KEYS = {"command", "workspace", "deck", "out_deck", "dictionary", "config"}
# Configuration keys of earlier versions that no longer mean anything.
_RETIRED_CONFIG_KEYS = {"workspace": None, "out": None, "target": "data_type", "input": None,
                        "retranslate": "overwrite", "forceupdate": "overwrite"}
_RETIRED_CONFIG_WHY = {
    "workspace": " (the workspace is now the first argument: pptx-narrator WS COMMAND ...)",
    "out": " (the deck pack writes is now its second argument: pack DECK OUT)",
    "input": " (the files a command reads are given on its command line)",
}
# Settings that are paths, recorded relative to the workspace when inside it.
_PATH_SETTINGS = ("dict_file", "letter_map", "ref_wav", "ref_text")


def _merge_effective(command, args, config, parser):
    values = _defaults().get(command, {}).copy()
    cfg = _config_for_command(config, command)
    for key, replacement in _RETIRED_CONFIG_KEYS.items():
        if key in cfg:
            cfg.pop(key)
            where = f" in {_CONFIG_PATH}" if _CONFIG_PATH else ""
            why = _RETIRED_CONFIG_WHY.get(key, "")
            logger.warning(f"The configuration key '{key}'{where} is no longer used{why}. "
                           + (f"Use '{replacement}' instead. " if replacement else "")
                           + "Remove it from the file to stop this message.")
    _deep_update(values, cfg)
    for key, value in vars(args).items():
        if key.startswith("_") or key in _POSITIONAL_KEYS:
            continue
        # argparse uses None for omitted optional values; boolean flags also use None here.
        if value is None:
            continue
        # An empty value on the command line clears what the configuration file set,
        # since there is otherwise no way to take a parameter back: --dict-file ''.
        if value == "" or (isinstance(value, list) and value and not any(value)):
            values[key] = None
            continue
        values[key] = value
    return values


def _resolve_lang(command, values, args, parser):
    """Turn --lang into --in-lang (and --out-lang for translate)."""
    lang = values.pop("lang", None)
    explicit = [flag for flag, key in (("--in-lang", "in_lang"), ("--out-lang", "out_lang"))
                if getattr(args, key, None) is not None]
    if lang is None:
        return values
    if getattr(args, "lang", None) is not None and explicit:
        parser.error(f"--lang cannot be combined with {' or '.join(explicit)}")
    if getattr(args, "lang", None) is None and explicit:
        return values  # --in-lang/--out-lang on the command line win over lang in the configuration
    values["in_lang"] = lang
    if command == "translate":
        values["out_lang"] = lang
    return values


def _record_effective(effective, workspace_dir):
    """The settings as recorded: paths inside the workspace relative to it."""
    out = dict(effective)
    for key in _PATH_SETTINGS:
        value = out.get(key)
        if isinstance(value, list):
            out[key] = [_rel_to_workspace(v, workspace_dir) if v else v for v in value]
        elif value:
            out[key] = _rel_to_workspace(value, workspace_dir)
    return out


def _audio_model_labels(workspace_dir, lang):
    """Return model labels of audio files available for one narration language."""
    pattern = re.compile(r"^slide_\d+_" + re.escape(lang) + r"\.(.+)\.m4a$")
    return sorted({m.group(1) for name in os.listdir(workspace_dir)
                   if (m := pattern.match(name))})


def _select_slides(slides, selection, parser):
    """Narrow the slides of a workspace to a --slides selection."""
    if not selection:
        return slides
    wanted = parse_slide_ranges(selection, max(slides) if slides else 0)
    chosen = [n for n in slides if n in wanted]
    if not chosen:
        parser.error(f"no slide of this workspace matches --slides {selection!r}"
                     + (f" (it has {', '.join(str(n) for n in slides)})" if slides else ""))
    return chosen


def _text_slides(workspace_dir, lang):
    return [slide for slide in sorted(_slides_from_workspace(workspace_dir))
            if os.path.exists(os.path.join(workspace_dir, text_filename(slide, lang)))]


def _model_label(v):
    return v["model"] if v["engine"] == "gpt_sovits" else f"qwen3-{v['qwen3_model_size']}"


def _validate_and_normalize(command, v, parser):
    if command == "extract":
        if v.get("in_lang"):
            raw = v["in_lang"]
            if isinstance(raw, str):
                raw = raw.split(",")
            v["in_lang"] = [normalize_lang(x) for x in raw if normalize_lang(x)]
    elif command == "translate":
        if not v.get("in_lang") or not v.get("out_lang"):
            parser.error("translate requires --in-lang and --out-lang")
        if v["in_lang"] == "auto" or v["out_lang"] == "auto":
            parser.error("--in-lang/--out-lang cannot be 'auto'")
        if lang_suffix(v["in_lang"]) == lang_suffix(v["out_lang"]):
            parser.error("translate requires different --in-lang and --out-lang")
    elif command in {"scan", "synthesize", "verify"}:
        if not v.get("in_lang"):
            parser.error(f"{command} requires --lang")
    if command in {"scan", "synthesize", "verify", "pack"}:
        if v.get("in_lang") == "auto":
            parser.error("--lang cannot be 'auto'")
    if command in {"scan", "translate", "synthesize"} and isinstance(v.get("dict_file"), str):
        v["dict_file"] = [v["dict_file"]]
    if command in {"extract", "scan", "translate", "pack"} and v.get("update") and v.get("overwrite"):
        parser.error("--update and --overwrite cannot be combined")
    if command == "scan" and v.get("append") and v.get("overwrite"):
        parser.error("--append and --overwrite cannot be combined")
    for flag, key in (("--ref-wav", "ref_wav"), ("--ref-text", "ref_text"), ("--letter-map", "letter_map"),
                      ("--dict-file", "dict_file")):
        if command not in {"scan", "translate", "synthesize"}:
            break
        values = v.get(key) if isinstance(v.get(key), list) else [v.get(key)]
        for path in values or []:
            if path and not os.path.isfile(path):
                parser.error(f"{flag}: no such file: {path}")
    if command == "synthesize":
        missing = [flag for flag, key in (("--ref-wav", "ref_wav"), ("--ref-text", "ref_text")) if not v.get(key)]
        if missing:
            parser.error("synthesize requires " + " and ".join(missing))
        if v["engine"] == "qwen3" and qwen3_language(v["in_lang"]) is None:
            parser.error(f"Qwen3-TTS does not support '{v['in_lang']}' (supported: {', '.join(QWEN3_LANGUAGES)})")
        if v["engine"] == "gpt_sovits":
            for flag, code in (("--lang", v["in_lang"]), ("--ref-lang", v["ref_lang"])):
                if gpt_sovits_language(code) is None:
                    parser.error(f"GPT-SoVITS does not support {flag} '{code}' (supported: {', '.join(GPT_SOVITS_LANGUAGES)})")
            if v["model"] not in MODELS_CONFIG:
                parser.error(f"unknown GPT-SoVITS model '{v['model']}' (available: {', '.join(MODELS_CONFIG)})")
    if command in {"verify", "pack"} and v["engine"] == "qwen3":
        v["model"] = f"qwen3-{v['qwen3_model_size']}"
    return v


_CONFIG_PATH = None


def _command_parser(parser, command):
    """The parser of one command, whose errors show that command's usage and name."""
    for act in parser._actions:
        if isinstance(act, argparse._SubParsersAction):
            return act.choices.get(command, parser)
    return parser


def _check_command_line(boot, parser):
    """Catch the two likely slips of the WS-first command line before argparse does."""
    if boot.workspace in COMMANDS and boot.command not in COMMANDS:
        parser.error(f"the workspace comes first and the command second:\n"
                     f"  pptx-narrator <workspace> {boot.workspace} ...")
    if boot.command and boot.command not in COMMANDS:
        close = difflib.get_close_matches(boot.command, COMMANDS, n=1)
        parser.error(f"unknown command '{boot.command}'"
                     + (f"; did you mean '{close[0]}'?" if close else f" (commands: {', '.join(COMMANDS)})"))


def main(argv=None):
    # Parse WS and the command first without loading a config, so that --config can be honored.
    bootstrap = argparse.ArgumentParser(add_help=False)
    bootstrap.add_argument("--config")
    bootstrap.add_argument("workspace", nargs="?")
    bootstrap.add_argument("command", nargs="?")
    boot, _ = bootstrap.parse_known_args(_normalize_argv(argv))
    try:
        config, config_path = _load_config(boot.config)
    except RuntimeError as e:
        raise SystemExit(str(e))

    parser = build_parser(config)
    if not boot.command:
        # -h, --help and --version are answers in themselves.
        asked = set(sys.argv[1:] if argv is None else argv) & {"-h", "--help", "--version"}
        if asked:
            parser.parse_args(argv)  # argparse prints it and exits 0
        if boot.workspace in COMMANDS:
            _check_command_line(boot, parser)
        parser.print_help()
        raise SystemExit(1)
    _check_command_line(boot, parser)
    args = parser.parse_args(argv)
    command = args.command
    global _CONFIG_PATH
    _CONFIG_PATH = config_path
    # From here on, an error is reported with the usage of the command it concerns.
    top_parser, parser = parser, _command_parser(parser, command)
    ws = os.path.abspath(args.workspace)
    effective = _merge_effective(command, args, config, parser)
    effective = _resolve_lang(command, effective, args, parser)
    effective = _validate_and_normalize(command, effective, parser)

    deck = out = dict_out = None
    if command in {"extract", "pack"}:
        deck = os.path.abspath(args.deck)
        if not os.path.isfile(deck) or not deck.lower().endswith(".pptx"):
            parser.error(f"{command} DECK must be a PPTX file: {args.deck}")
    if command == "extract":
        if os.path.exists(ws) and not os.path.isdir(ws):
            parser.error(f"WS is not a directory: {args.workspace}")
        os.makedirs(ws, exist_ok=True)
    elif not os.path.isdir(ws):
        parser.error(f"workspace does not exist: {args.workspace}\n\nCreate it by extracting a deck into it:\n"
                     f"  pptx-narrator {args.workspace} extract <deck>.pptx")
    if command == "history":
        show_history(ws, dates=effective.get("dates"))
        return
    log_handler = _open_workspace_log(ws, argv)
    logger.info(f"Workspace: {ws}")
    try:
        _run_command(command, args, effective, config, config_path, parser, ws)
    except SystemExit:
        raise
    except KeyboardInterrupt:
        logger.error(f"{command} was interrupted.")
        _record_failure(ws, command, "interrupted")
        raise SystemExit(130)
    except Exception as e:
        logger.error(f"{command} failed: {_describe_error(e)}")
        log_handler.stream.write(traceback.format_exc())
        log_handler.flush()
        _record_failure(ws, command, f"{type(e).__name__}: {e}")
        print(f"pptx-narrator {command}: the details are in {os.path.join(ws, LOG_FILE)}", file=sys.stderr)
        raise SystemExit(1)
    finally:
        logger.removeHandler(log_handler)
        log_handler.close()


def _describe_error(e):
    """One line saying what went wrong, in terms of the files and tools involved."""
    if isinstance(e, FileNotFoundError):
        return f"file not found: {e.filename or e}"
    if isinstance(e, PermissionError):
        return f"permission denied: {e.filename or e}"
    if isinstance(e, zipfile.BadZipFile):
        return f"not a valid PPTX (zip) file: {e}"
    return f"{type(e).__name__}: {e}"


def _record_failure(ws, command, reason):
    _append_history(ws, {"command": command, "generated_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
                         "workspace_at_run": ws, "status": "failed", "error": reason})


def _run_command(command, args, effective, config, config_path, parser, ws):
    before = _workspace_files(ws)
    deck = os.path.abspath(args.deck) if command in {"extract", "pack"} else None
    out = dict_out = None
    audio_lang = effective.get("in_lang")
    if command == "pack" and "audio" in resolve_targets(effective.get("data_type")) and not audio_lang:
        # The texts of every language can go into one note, but a slide plays one audio.
        audio_langs = _audio_languages(ws)
        if len(audio_langs) > 1:
            parser.error("audio of several languages is in this workspace (" + ", ".join(audio_langs)
                         + "); choose one with --lang.")
        audio_lang = audio_langs[0] if audio_langs else None
        if audio_lang:
            logger.info(f"Audio: {audio_lang} (the only language with audio in this workspace).")
    if command in {"verify", "pack"} and audio_lang:
        configured = _config_for_command(config, command)
        model_is_explicit = any(getattr(args, key, None) is not None
                                for key in ("engine", "model", "qwen3_model_size"))
        model_is_configured = any(key in configured for key in ("engine", "model", "qwen3_model_size"))
        labels = _audio_model_labels(ws, audio_lang)
        if not model_is_explicit and not model_is_configured and len(labels) > 1:
            parser.error("multiple audio model labels are present in this workspace: "
                         + ", ".join(labels) + ". Specify --engine (and model size if applicable).")

    if command == "extract":
        prs = Presentation(deck)
        req_slides = sorted(parse_slide_ranges(effective.get("slides"), len(prs.slides)))
        langs = effective.get("in_lang")
        flags = dict(update=effective["update"], overwrite=effective["overwrite"])
        if not langs:
            step_extract_notes(deck, ws, req_slides, "auto", **flags)
        else:
            # Extract each requested language independently. A PPTX may legitimately
            # contain notes in several languages; each requested language is therefore
            # a selector, not a claim that every note in the deck has that language.
            # A requested language the deck does not contain is an error: the run would
            # otherwise report success while producing nothing for that language.
            missing = [lang for lang in langs if not step_extract_notes(deck, ws, req_slides, lang, **flags)]
            if missing:
                parser.error("no note in " + ", ".join(missing) + " was found in " + os.path.basename(deck))
    elif command == "scan":
        dict_out = os.path.abspath(args.dictionary)
        exists = os.path.exists(dict_out)
        if exists and not (effective["append"] or effective["overwrite"]):
            parser.error(f"{args.dictionary} already exists. Add --append to add the new terms to it,\n"
                         f"or --overwrite to make it again from scratch.")
        available_slides = _text_slides(ws, effective["in_lang"])
        if not available_slides:
            parser.error(f"no {effective['in_lang']} slide text was found in workspace: {ws}")
        slides = _select_slides(available_slides, effective.get("slides"), parser)
        set_aside = None
        if exists and effective["overwrite"]:
            # Kept until the new dictionary is written, so that a failed scan loses nothing.
            set_aside = dict_out + ".pptx_narrator_previous"
            os.replace(dict_out, set_aside)
        try:
            readable = ([dict_out] if os.path.exists(dict_out) else []) + list(effective.get("dict_file") or [])
            dictionaries = load_dictionaries(readable, effective["in_lang"])
            step_scan_and_update_dict(ws, dict_out, dictionaries, slides, effective["in_lang"],
                                      source_lang=effective["in_lang"], for_translation=False,
                                      propose_compounds=effective["scan_compounds"])
        except BaseException:
            if set_aside:
                os.replace(set_aside, dict_out)
            raise
        if set_aside:
            if os.path.exists(dict_out):
                os.remove(set_aside)
                logger.info(f"Made {os.path.basename(dict_out)} again from scratch (--overwrite).")
            else:
                os.replace(set_aside, dict_out)
                logger.warning(f"No term was found; {os.path.basename(dict_out)} is left as it was.")
    elif command == "translate":
        slides = _select_slides(sorted(_slides_from_workspace(ws)), effective.get("slides"), parser)
        dictionaries = load_dictionaries(effective.get("dict_file"), effective["in_lang"])
        if not step_translate_notes(ws, slides, effective["in_lang"], effective["out_lang"],
                                    dictionary=dictionaries, overwrite=effective["overwrite"],
                                    update=effective["update"],
                                    model=effective["translate_model"], device=effective["translate_device"]):
            parser.error(f"no {effective['in_lang']} text was found in workspace: {ws}")
    elif command == "synthesize":
        slides = _select_slides(_text_slides(ws, effective["in_lang"]), effective.get("slides"), parser)
        if not slides:
            parser.error(f"no {effective['in_lang']} text was found in workspace: {ws}")
        dictionaries = load_dictionaries(effective.get("dict_file"), effective["in_lang"])
        letter_map_data = load_letter_map(effective.get("letter_map"))
        if effective["engine"] == "gpt_sovits":
            config_model = MODELS_CONFIG[effective["model"]]
            base_url = effective["api_url"].rstrip("/")
            logger.info(f"Switching GPT-SoVITS weights to {effective['model']}...")
            try:
                requests.get(f"{base_url}/set_gpt_weights", params={"weights_path": config_model["gpt"]}, timeout=300)
                requests.get(f"{base_url}/set_sovits_weights", params={"weights_path": config_model["sovits"]}, timeout=300)
            except requests.RequestException as e:
                parser.error(f"Could not reach the GPT-SoVITS API server at {effective['api_url']}: {e}")
        model_label = _model_label(effective)
        started = time.time() - 1
        if effective["engine"] == "qwen3":
            step_generate_audio_qwen3(ws, slides, effective["in_lang"], effective["ref_wav"],
                                      effective["ref_text"], dictionaries, model_label,
                                      effective["qwen3_model_size"], effective["qwen3_device"],
                                      enable_drc=effective["enable_drc"], drc_threshold=effective["drc_threshold"],
                                      drc_ratio=effective["drc_ratio"], letter_map=letter_map_data)
        else:
            step_generate_audio(ws, slides, effective["in_lang"], effective["ref_wav"],
                                effective["ref_text"], effective["ref_lang"], effective["api_url"],
                                dictionaries, model_label, enable_drc=effective["enable_drc"],
                                drc_threshold=effective["drc_threshold"], drc_ratio=effective["drc_ratio"],
                                letter_map=letter_map_data)
        made = record_audio_sources(ws, slides, effective["in_lang"], model_label,
                                    effective.get("dict_file"), started)
        missing = [n for n in slides if n not in made]
        if missing:
            logger.warning("No audio was made for slide(s): " + ", ".join(map(str, missing)) + ".")
    elif command == "verify":
        model_label = _model_label(effective)
        slides = _select_slides(
            [slide for slide in _text_slides(ws, effective["in_lang"])
             if os.path.exists(os.path.join(ws, audio_filename(slide, effective["in_lang"], model_label)))],
            effective.get("slides"), parser)
        if not slides:
            available = _audio_model_labels(ws, effective["in_lang"])
            parser.error(f"verify found no {effective['in_lang']} audio for model '{model_label}'. "
                         + (f"Available model labels: {', '.join(available)}." if available else ""))
        for n in slides:
            if audio_is_older_than_text(ws, n, effective["in_lang"], model_label):
                logger.warning(f"Slide #{n}: {text_filename(n, effective['in_lang'])} was edited after its audio "
                               f"was made; the check compares the audio with the text it was made from.")
        step_verify_audio(ws, slides, effective["in_lang"], model_label,
                          effective["asr_model"], effective["asr_device"], effective["verify_threshold"],
                          effective["cer_threshold"], max_difference=effective["max_difference"] or None,
                          min_difference=effective["min_difference"])
    elif command == "pack":
        out = os.path.abspath(args.out_deck)
        if not out.lower().endswith(".pptx"):
            parser.error(f"pack OUT must be a .pptx file name: {args.out_deck}")
        if out == deck or (os.path.exists(out) and os.path.samefile(out, deck)):
            parser.error("pack OUT must differ from DECK (DECK itself is never changed)")
        if os.path.exists(out) and not (effective["update"] or effective["overwrite"]):
            parser.error(f"{args.out_deck} already exists. Add --update or --overwrite to replace it.")
        prs = Presentation(deck)
        slides = sorted(parse_slide_ranges(effective.get("slides"), len(prs.slides)))
        model_label = _model_label(effective)
        targets = resolve_targets(effective.get("data_type"))
        if "audio" in targets:
            if not audio_lang:
                logger.warning("No audio was found in this workspace; only the texts are written.")
            elif not any(os.path.exists(os.path.join(ws, audio_filename(s, audio_lang, model_label)))
                         for s in slides):
                available = _audio_model_labels(ws, audio_lang)
                logger.warning(f"No {audio_lang} audio for model '{model_label}' was found"
                               + (f" (available model labels: {', '.join(available)})" if available else "")
                               + "; no audio is written.")
        step_pack_pptx(deck, out, ws, slides, effective.get("in_lang"), model_label,
                       targets=targets, update=effective["update"], overwrite=effective["overwrite"],
                       audio_lang=audio_lang,
                       remove_recorded=RECORDED_CHOICES[effective["remove_recorded"]],
                       icon_outside=not effective["keep_audio_icon"],
                       pause_ms=int(round(effective["slide_pause"] * 1000)))

    written = _files_written(ws, before)
    _suggest_next(command, args, effective, ws, deck=deck, out=out, dict_out=dict_out)
    _report(command, ws, written, out=out, dict_out=dict_out)
    _record_run(command, ws, effective, config_path, deck=deck, out=out, dict_out=dict_out, written=written)


LOG_FILE = "pptx_narrator.log"
_RECORD_FILES = {STATE_FILE, RESOLVED_CONFIG_FILE, ".pptx_narrator_history.jsonl", LOG_FILE}


def _open_workspace_log(ws, argv):
    """Everything this run logs also goes to the log file of the workspace."""
    handler = logging.FileHandler(os.path.join(ws, LOG_FILE), encoding="utf-8")
    handler.setFormatter(logging.Formatter("[%(asctime)s] %(levelname)s: %(message)s", "%Y-%m-%d %H:%M:%S"))
    logger.addHandler(handler)
    words = sys.argv[1:] if argv is None else argv
    logger.info("==== pptx-narrator " + " ".join(shlex.quote(w) for w in words) + f" (version {__version__})")
    return handler


def _say(text):
    """Text for the reader (not a log line): to stderr, and into the log file."""
    print(text, file=sys.stderr)
    for h in logger.handlers:
        if isinstance(h, logging.FileHandler):
            h.stream.write(text + "\n")
            h.flush()


def _workspace_files(ws):
    out = {}
    for name in os.listdir(ws):
        path = os.path.join(ws, name)
        if os.path.isfile(path) and name not in _RECORD_FILES:
            st = os.stat(path)
            out[name] = (st.st_mtime_ns, st.st_size)
    return out


def _files_written(ws, before):
    after = _workspace_files(ws)
    return {"created": sorted(n for n in after if n not in before),
            "changed": sorted(n for n in after if n in before and after[n] != before[n]),
            "removed": sorted(n for n in before if n not in after)}


def _q(path):
    return shlex.quote(path)


def _model_options(effective):
    if effective.get("engine") == "gpt_sovits":
        return f" --engine gpt_sovits --model {effective['model']}"
    return f" --engine qwen3 --qwen3-model-size {effective.get('qwen3_model_size', '1.7B')}"


def _suggest_next(command, args, effective, ws, deck=None, out=None, dict_out=None):
    """What can be run next, ready to copy, with the values of this run written out."""
    w = "pptx-narrator " + _q(args.workspace)
    lang = effective.get("out_lang") if command == "translate" else effective.get("in_lang")
    langs = [lang] if isinstance(lang, str) and lang else (
        lang if isinstance(lang, list) and lang else sorted({l for n in _slides_from_workspace(ws)
                                                             for l in _text_langs_of_slide(ws, n)}))
    lines = []
    for l in langs:
        if command in {"extract", "translate"}:
            lines.append(f'{w} scan "dictionary_{l}.csv" --lang {l}'
                         + (" --append" if os.path.exists(f"dictionary_{l}.csv") else ""))
            lines.append(f'{w} synthesize --lang {l} --ref-wav "ref.wav" --ref-text "ref.txt"')
        elif command == "scan":
            lines.append(f"{w} synthesize --lang {l} --dict-file {_q(args.dictionary)} "
                         f'--ref-wav "ref.wav" --ref-text "ref.txt"')
        elif command == "synthesize":
            lines.append(f"{w} verify --lang {l}{_model_options(effective)}")
            lines.append(f'{w} pack "deck.pptx" "narrated.pptx" --lang {l}{_model_options(effective)}')
        elif command == "verify":
            lines.append(f'{w} pack "deck.pptx" "narrated.pptx" --lang {l}{_model_options(effective)}')
    if not lines:
        return
    _say("\nNext, for example (replace the names in quotes with your own):")
    for line in lines:
        _say("  " + line)


def _report(command, ws, written, out=None, dict_out=None):
    """What this run did, as the last thing it says."""
    _say(f"\nReport of {command} (workspace {ws}):")
    for key, label in (("created", "created"), ("changed", "rewritten"), ("removed", "removed")):
        if written[key]:
            _say(f"  {label} in the workspace: " + ", ".join(written[key]))
    if not any(written.values()):
        _say("  nothing was written in the workspace")
    if dict_out:
        _say(f"  dictionary: {dict_out}")
    if out:
        _say(f"  deck written: {out}")
    _say(f"  log: {os.path.join(ws, LOG_FILE)}")


def show_history(ws, dates=False):
    """List the commands run in a workspace, from its history file."""
    path = os.path.join(ws, ".pptx_narrator_history.jsonl")
    if not os.path.exists(path):
        print(f"No command has been recorded in {ws}.")
        return
    with open(path, encoding="utf-8") as f:
        for k, line in enumerate(f, 1):
            try:
                e = json.loads(line)
            except ValueError:
                continue
            eff = e.get("effective", {})
            parts = [e.get("command", "?")]
            for key in ("deck", "dictionary", "out"):
                if e.get("paths", {}).get(key):
                    parts.append(e["paths"][key])
            for key, flag in (("in_lang", "--in-lang"), ("out_lang", "--out-lang")):
                if eff.get(key):
                    v = eff[key]
                    parts.append(f"{flag} {','.join(v) if isinstance(v, list) else v}")
            if eff.get("slides"):
                parts.append(f"--slides {eff['slides']}")
            files = e.get("written", {})
            n_files = len(files.get("created", [])) + len(files.get("changed", []))
            when = (e.get("generated_at", "")[:16].replace("T", " ") + "  ") if dates else ""
            print(f"{k:3d}  {when}{' '.join(parts)}" + (f"  ({n_files} file(s) written)" if files else ""))


def _record_run(command, ws, effective, config_path, deck=None, out=None, dict_out=None, written=None):
    """Record what was run in the workspace.

    Everything inside the workspace is recorded relative to it, so that the
    workspace can be moved or copied without invalidating its own record;
    anything outside it keeps its absolute path, because nothing else identifies it.
    """
    paths = {key: _rel_to_workspace(path, ws)
             for key, path in (("deck", deck), ("out", out), ("dictionary", dict_out)) if path}
    input_path = deck if command in {"extract", "pack"} else ws
    try:
        snapshot = _input_snapshot(input_path, command)
    except RuntimeError:
        snapshot = {"path": input_path, "kind": "file" if os.path.isfile(input_path) else "directory"}
    rel_snapshot = _relativize_snapshot(snapshot, ws)
    recorded = _record_effective(effective, ws)
    generated_at = time.strftime("%Y-%m-%dT%H:%M:%S%z")

    # (No state file is written: nothing is carried over from one run to the next.)

    metadata = {
        "software_version": __version__,
        "generated_at": generated_at,
        "command": command,
        "workspace": ".",
        "workspace_at_run": ws,
        "paths_relative_to": "workspace",
        "config_file": _rel_to_workspace(config_path, ws) if config_path else "",
        "input_path": rel_snapshot.get("path", ""),
        "input_kind": rel_snapshot.get("kind", ""),
    }
    metadata.update({f"{key}_path": value for key, value in paths.items()})
    if rel_snapshot.get("kind") == "file":
        metadata["input_sha256"] = rel_snapshot.get("sha256", "")
    else:
        metadata["input_files_json"] = json.dumps(rel_snapshot.get("files", []), ensure_ascii=False, sort_keys=True)
    _save_resolved_config({"metadata": metadata, command: recorded}, _workspace_path(ws, RESOLVED_CONFIG_FILE))
    _append_history(ws, {
        "command": command,
        "generated_at": generated_at,
        "workspace_at_run": ws,
        "paths": paths,
        "input": rel_snapshot,
        "effective": recorded,
        "written": written or {},
        "status": "success" if (out or dict_out or any((written or {}).values())) else "nothing written",
    })


def _slides_from_workspace(workspace_dir):
    """Return slide numbers represented by language-tagged text/audio files in a workspace."""
    slides = set()
    if not os.path.isdir(workspace_dir):
        return slides
    for name in os.listdir(workspace_dir):
        m = _TEXT_FILE_RE.match(name)
        if m:
            slides.add(int(m.group(1)))
            continue
        m = re.match(r"^slide_(\d+)(?:_[^.]+)?\.[^.]+$", name)
        if m:
            slides.add(int(m.group(1)))
    return slides


if __name__ == "__main__":
    main()
