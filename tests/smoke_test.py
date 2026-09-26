"""Smoke tests for PPTX-Narrator (no TTS/ASR models or network needed; back-ends are mocked).

Run from the repository root:  python tests/smoke_test.py
Requires: python-pptx, pydub (+ FFmpeg), numpy, soundfile, py3langid.
"""
import os, sys, types, tempfile, csv, zipfile, re, io, contextlib, argparse, json, shutil
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
import pptx_narrator as pn
from pptx import Presentation
from pydub import AudioSegment
import numpy as np


def ok(cond, msg):
    print(("PASS " if cond else "FAIL ") + msg)
    if not cond:
        sys.exit(1)


def write(path, text):
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)


def read(path):
    with open(path, encoding="utf-8") as f:
        return f.read()


d = tempfile.mkdtemp()

# ---------------------------------------------------------------- languages
ok(pn.normalize_lang("ZH_cn") == "zh-CN" and pn.normalize_lang("zh") == "zh-CN" and pn.normalize_lang("eng") == "en"
   and pn.normalize_lang("he") == "iw" and pn.normalize_lang("Auto") == "auto", "normalize_lang")
ok(pn.text_filename(3, "ja") == "slide_3_ja.txt" and pn.text_filename(3, "en") == "slide_3_en.txt" and pn.text_filename(3, "de") == "slide_3_de.txt",
   "text file names carry the language")
ok(pn.audio_filename(1, "ja", "v4") == "slide_1_ja.v4.m4a" and pn.spoken_filename(1, "en", "v4") == "slide_1_en.v4.spoken.txt",
   "audio/spoken filenames")
ok(pn.qwen3_language("de") == "German" and pn.qwen3_language("zh-TW") == "Chinese" and pn.qwen3_language("nl") is None, "Qwen3 languages")
ok(pn.gpt_sovits_language("yue") == "yue" and pn.gpt_sovits_language("de") is None, "GPT-SoVITS languages")
ok([pn.detect_language(t)[0] for t in ["今日はDNAの話です", "안녕하세요 여러분", "На этом слайде показаны результаты.",
                                       "Heute sprechen wir über die mRNA und ihre Rolle.", "这张幻灯片显示了实验结果。", "這張投影片顯示了實驗結果。"]]
   == ["ja", "ko", "ru", "de", "zh-CN", "zh-TW"], "language identification")
langs = pn.detect_note_languages({1: "Willkommen zur heutigen Vorlesung über Genomeditierung.", 2: "Fragen?", 3: "Vielen Dank.",
                                  4: "Hier vergleichen wir die Effizienz von drei Methoden in Zellen."})
ok(set(langs.values()) == {"de"}, f"short notes follow the deck language {langs}")
ok(pn.detect_note_languages({1: "本日はゲノム編集について説明します。", 2: "結論"}) == {1: "ja", 2: "ja"}, "kanji-only note in a Japanese deck")
ok(pn.split_into_chunks("Hello world. It costs 3.5 mg! Next?") == ["Hello world.", "It costs 3.5 mg!", "Next?"], "sentence split (Latin)")
ok(pn.split_into_chunks("今日は晴れ。明日は雨！") == ["今日は晴れ。", "明日は雨！"], "sentence split (CJK)")

# ---------------------------------------------------------------- dictionaries
readings_ja = os.path.join(d, "readings_ja.csv")
write(readings_ja, "string,replacement,type\nGbp,ギガベースペア,unit\nDNA,ディーエヌエー,\nCRISPR-Cas9,クリスパーキャスナイン,\n")
readings_de = os.path.join(d, "readings_de.csv")
write(readings_de, "mRNA,Boten-RNA,\nkDa,Kilodalton,unit\nX1,a\\1b,\n")          # no header
terms_ja_de = os.path.join(d, "terms_ja_de.csv")
write(terms_ja_de, "string,replacement\n塩基対,Basenpaare\n3 Gbp,drei Gigabasenpaare\n")
legacy = os.path.join(d, "legacy_dict.csv")
write(legacy, "Term,Japanese_Reading,English_Reading,Type\nGbp,ギガベースペア,gigabase pairs,unit\nRNA,アールエヌエー,R N A,\n")

R_ja = pn.load_dictionaries([readings_ja, legacy], "ja")
ok(pn.apply_dictionary("DNA 3Gbp 5mg RNA", R_ja, "ja") == "ディーエヌエー 3ギガベースペア 5ミリグラム アールエヌエー",
   "replacements + built-in units + v1.x file (Japanese column)")
ok(pn.apply_dictionary("RNA is 3 Gbp and 5 mg", pn.load_dictionaries([legacy], "en"), "en") == "R N A is 3 gigabase pairs and 5 mg",
   "v1.x file, English column")
ok(pn.load_dictionaries([legacy], "de") == [], "v1.x file has no column for other languages")
R_de = pn.load_dictionaries([readings_de], "de")
ok(pn.apply_dictionary("Die mRNA hat 2 kDa. X1", R_de, "de") == "Die Boten-RNA hat 2 Kilodalton. a\\1b", "headerless file, literal backslash")
ok(len(pn.load_dictionaries([readings_ja, os.path.join(d, "missing.csv")], "ja")) == 3, "missing dictionary file tolerated")
ok(pn.apply_dictionary("mRNA 5 mg", [], "fr") == "mRNA 5 mg", "no dictionary -> unchanged")
T = pn.load_dictionaries([terms_ja_de])
ok(pn.apply_dictionary("塩基対は3 Gbpで5 mg", T, "ja", builtin_units=False) == "Basenpaareはdrei Gigabasenpaareで5 mg",
   "pre-translation replacement without built-in unit readings")
both = pn.load_dictionaries([terms_ja_de, readings_de], "de")
pre = pn.apply_dictionary("塩基対とmRNA", both, "ja", builtin_units=False)
ok(pre == "BasenpaareとBoten-RNA" and pn.apply_dictionary("Basenpaare und Boten-RNA, 2 kDa", both, "de") == "Basenpaare und Boten-RNA, 2 Kilodalton",
   "one dictionary applied before translation and again before synthesis")

# ---------------------------------------------------------------- extraction
prs = Presentation(); layout = prs.slide_layouts[5]
for t in ["今日はDNAの話です。", "Today we talk about DNA and its structure in the cell.", "안녕하세요 여러분 DNA"]:
    s = prs.slides.add_slide(layout); s.notes_slide.notes_text_frame.text = t
deck = os.path.join(d, "deck.pptx"); prs.save(deck)
ws = os.path.join(d, "ws"); os.makedirs(ws)
pn.step_extract_notes(deck, ws, [1, 2, 3], "auto")
ok(sorted(os.listdir(ws)) == ["note_baseline.json", "slide_1_ja.txt", "slide_2_en.txt", "slide_3_ko.txt"], f"auto extraction {sorted(os.listdir(ws))}")
ws2 = os.path.join(d, "ws2"); os.makedirs(ws2)
pn.step_extract_notes(deck, ws2, [2], "de")
ok(os.listdir(ws2) == [], "--in-lang is a selector: a note of another language is left out")
ws2b = os.path.join(d, "ws2b"); os.makedirs(ws2b)
pn.step_extract_notes(deck, ws2b, [1, 2, 3], "en")
ok(sorted(os.listdir(ws2b)) == ["note_baseline.json", "slide_2_en.txt"], "--in-lang keeps only the notes of that language")
ok(pn.step_extract_notes(deck, ws2b, [1, 2, 3], "fr") == 0
   and pn.step_extract_notes(deck, ws2b, [1, 2, 3], "en") == 1,
   "extract reports what this run wrote, not what the workspace already held")
ok(pn._file_role("slide_3_ja.txt") == "note text"
   and pn._file_role("slide_3_ja.v4.spoken.txt") == "spoken text"
   and pn._file_role("slide_3_ja.v4.m4a") == "audio"
   and pn._file_role("verify_report_ja.v4.csv") == "verification report",
   "a directory record says what each file is")
ok(pn.find_source_text(ws, 1, "auto", exclude_lang="de")[0] == "ja" and pn.find_source_text(ws, 3, "auto", exclude_lang="de")[0] == "ko"
   and pn.find_source_text(ws, 2, "auto", exclude_lang="en") == (None, None), "find_source_text")

# ---------------------------------------------------------------- translation
class FakeTr:
    calls = []
    glossaries = []
    def __init__(self, model_id=None, device=None):
        self.model_id = model_id
    def translate(self, text, source, target, glossary=None):
        if target == "xx":
            raise ValueError("unsupported language")
        FakeTr.calls.append(text)
        FakeTr.glossaries.append(glossary)
        return f"[{source}->{target}] {text}"

pn.make_translator = FakeTr
pn.step_translate_notes(ws, [1, 2, 3], "auto", "de")
ok(read(os.path.join(ws, "slide_1_de.txt")).startswith("[ja->de]") and read(os.path.join(ws, "slide_2_de.txt")).startswith("[en->de]")
   and read(os.path.join(ws, "slide_3_de.txt")).startswith("[ko->de]"), "translate ja/en/ko -> de")
write(os.path.join(ws, "slide_1_ja.txt"), "塩基対の話です。")
pn.step_translate_notes(ws, [1], "auto", "de", dictionary=T)
ok("Basenpaare" not in read(os.path.join(ws, "slide_1_de.txt")),
   "an existing translation whose source has changed is kept without --update or --overwrite")
n_calls = len(FakeTr.calls)
write(os.path.join(ws, "slide_2_de.txt"), "von Hand verbessert")
write(os.path.join(ws, "slide_2_en.txt"), "The source has changed.")
write(os.path.join(ws, "slide_3_ko.txt"), "원문이 바뀌었습니다.")
pn.step_translate_notes(ws, [1, 2, 3], "auto", "de", update=True)
ok(read(os.path.join(ws, "slide_1_de.txt")).startswith("[ja->de] 塩基対")
   and read(os.path.join(ws, "slide_2_de.txt")) == "von Hand verbessert"
   and read(os.path.join(ws, "slide_3_de.txt")).startswith("[ko->de] 원문")
   and len(FakeTr.calls) == n_calls + 2,
   "--update translates again where the source changed, but keeps a translation edited by hand")
n_calls = len(FakeTr.calls)
pn.step_translate_notes(ws, [1, 3], "auto", "de", update=True)
ok(len(FakeTr.calls) == n_calls, "--update leaves a translation that is up to date with its source")
write(os.path.join(ws, "slide_4_ja.txt"), "新しいスライド")
write(os.path.join(ws, "slide_4_de.txt"), "von Hand geschrieben")
pn.step_translate_notes(ws, [4], "auto", "de", update=True)
ok(read(os.path.join(ws, "slide_4_de.txt")) == "von Hand geschrieben",
   "--update keeps a text that translate did not make")
os.remove(os.path.join(ws, "slide_4_ja.txt")); os.remove(os.path.join(ws, "slide_4_de.txt"))
pn.step_translate_notes(ws, [1], "auto", "de", dictionary=T, overwrite=True)
ok(FakeTr.calls[-1] == "塩基対の話です。" and ("塩基対", "Basenpaare") in (FakeTr.glossaries[-1] or [])
   and read(os.path.join(ws, "slide_1_ja.txt")) == "塩基対の話です。",
   "the whole note goes to the translator unchanged; dictionary terms occurring in it go as instructions")
ok(pn.glossary_for("mRNAとRNA", [("RNA", "RNA", ""), ("mRNA", "Boten-RNA", ""), ("DNA", "DNA", "")])
   == [("mRNA", "Boten-RNA"), ("RNA", "RNA")], "only terms occurring in the note, longest first")
write(os.path.join(ws2, "slide_2_de.txt"), "Heute sprechen wir über Basenpaare.")
pn.step_translate_notes(ws2, [2], "de", "ja")
ok(read(os.path.join(ws2, "slide_2_ja.txt")).startswith("[de->ja]"), "translate de -> ja")
pn.step_translate_notes(ws2, [2], "de", "xx")
ok(not os.path.exists(os.path.join(ws2, "slide_2_xx.txt")), "unsupported target language handled")

# ---------------------------------------------------------------- scan
pn._ensure_nltk_data = lambda: None
sys.modules["nltk.corpus"] = types.SimpleNamespace(
    stopwords=types.SimpleNamespace(words=lambda *a: ["the", "is", "about", "we", "and", "its", "in"]),
    words=types.SimpleNamespace(words=lambda: ["today", "talk", "structure", "cell", "slide"]))
ws3 = os.path.join(d, "ws3"); os.makedirs(ws3)
write(os.path.join(ws3, "slide_1_ja.txt"), "今日はDNAとPCRとGFPの話です。")
write(os.path.join(ws3, "slide_2_ja.txt"), "CRISPRについて。")
write(os.path.join(ws3, "slide_2_de.txt"), "Heute zeigen wir, wie CRISPR und mRNA in Zellen wirken, 5 kDa groß.")
terms_file = os.path.join(d, "new_terms.csv")
pn.step_scan_and_update_dict(ws3, terms_file, [], [1, 2], "de", for_translation=True)
rows = list(csv.reader(open(terms_file, encoding="utf-8")))
ok(rows[0] == ["string", "replacement", "type"] and sorted(r[0] for r in rows[1:]) == ["CRISPR", "DNA", "GFP", "PCR"]
   and all(r[1] == "" for r in rows[1:]), f"scan before translation: notes scanned, blank replacements {rows[1:]}")
pn.step_scan_and_update_dict(ws3, readings_de, pn.load_dictionaries([readings_de], "de"), [1, 2], "de")
rows = list(csv.reader(open(readings_de, encoding="utf-8")))
added = {r[0]: r[1] for r in rows[3:]}
ok(added.get("CRISPR") == "C R I S P R" and "kDa" not in added and "DNA" not in added,
   f"scan for synthesis: only the German narration text is scanned, with readings {added}")
pn.step_scan_and_update_dict(ws3, readings_ja, pn.load_dictionaries([readings_ja], "ja"), [1], "ja")
ok(["PCR", "P C R", ""] in list(csv.reader(open(readings_ja, encoding="utf-8"))), "scan of Japanese narration text")

# ---------------------------------------------------------------- synthesis (engines mocked)
logs = []
class Capture(pn.logging.Handler):
    def emit(self, record):
        logs.append(record.getMessage())
pn.logger.addHandler(Capture())
reftxt = os.path.join(d, "ref.txt"); write(reftxt, "x")
pn.step_generate_audio_qwen3(ws, [1], "nl", "r.wav", reftxt, R_de, "qwen3-1.7B", "1.7B", "cpu")
ok(any("does not support 'nl'" in m for m in logs), "Qwen3: unsupported language rejected before loading the model")
pn.step_generate_audio(ws, [1], "de", "r.wav", reftxt, "ja", "http://127.0.0.1:1/", R_de, "v4")
ok(any("GPT-SoVITS supports only" in m for m in logs), "GPT-SoVITS: unsupported language rejected")

sent = {}
class Resp:
    status_code = 500
    text = "mock"
pn.requests = types.SimpleNamespace(post=lambda url, json, timeout: (sent.update(json), Resp())[1], RequestException=Exception)
pn.step_generate_audio(ws, [3], "ko", "r.wav", reftxt, "ja", "http://x/", R_de, "v4")
ok(sent.get("text_lang") == "ko" and sent.get("prompt_lang") == "ja" and os.path.exists(os.path.join(ws, "slide_3_ko.v4.spoken.txt")),
   "GPT-SoVITS request languages")

seen = {}
class FakeModel:
    def create_voice_clone_prompt(self, ref_audio, ref_text):
        return "P"
    def generate_voice_clone(self, text, language, **kw):
        seen.setdefault("languages", set()).add(language)
        seen["max_new_tokens"] = kw.get("max_new_tokens")
        return [np.zeros(2400, dtype="float32")], 24000
pn._load_qwen3_model = lambda size, device: FakeModel()
pn.step_generate_audio_qwen3(ws, [1, 2], "de", "r.wav", reftxt, R_de, "qwen3-1.7B", "1.7B", "auto")
ok(seen.get("languages") == {"German"} and os.path.exists(os.path.join(ws, "slide_1_de.qwen3-1.7B.m4a")), "Qwen3 German synthesis")
ok(seen.get("max_new_tokens") is not None
   and pn.chunk_token_budget("a" * 140, 14.0) < pn.chunk_token_budget("a" * 700, 14.0)
   and pn.chunk_token_budget("", 14.0) >= 20 * pn.QWEN3_CODEC_HZ,
   "generation is capped at a length the text can justify")

# ---------------------------------------------------------------- verification (ASR mocked)
fw = types.ModuleType("faster_whisper")
class Segment:
    def __init__(self, text):
        self.text = text
class WhisperModel:
    def __init__(self, *a, **k):
        pass
    def transcribe(self, path, language=None):
        seen["asr_language"] = language
        return [Segment(" ganz anderer Text")], None
fw.WhisperModel = WhisperModel
sys.modules["faster_whisper"] = fw
pn.step_verify_audio(ws, [1, 2], "de", "qwen3-1.7B", threshold=0.5, cer_threshold=0.2)
report = list(csv.reader(open(os.path.join(ws, "verify_report_de.qwen3-1.7B.csv"), encoding="utf-8")))
ok(report[0][:4] == ["slide", "similarity", "cer", "status"] and seen["asr_language"] == "de" and report[1][3] == "FLAGGED",
   "verification report for German")
ok(pn.text_scores("Hello, World!", "hello world")[1] == 0 and pn.text_scores("Héllo", "hello")[1] > 0, "text normalization for comparison")
pj = types.ModuleType("pyopenjtalk")
pj.g2p = lambda t, kana=True: {"今日は晴れです。": "キョーワ、ハレデス。", "今日は晴れでした": "キョーワハレデシタ"}.get(t, t)
sys.modules["pyopenjtalk"] = pj
ok([round(x, 3) if isinstance(x, float) else x for x in pn.kana_scores("今日は晴れです。", "今日は晴れでした")]
   == [0.824, 0.25, "キョーワハレデス", "キョーワハレデシタ"], "kana similarity and CER")

ok(pn.apply_dictionary("5\u2019末端と5\u2032末端", [("5'", "ごだっしゅ", "")], "ja") == "ごだっしゅ末端とごだっしゅ末端",
   "typographic apostrophes and primes match a plain dictionary entry")

cmt = os.path.join(d, "commented.csv")
write(cmt, "; readings for the DNA lecture\nstring,replacement,type\nDNA,ディーエヌエー,   ; only in the first slide\n;PCR,ピーシーアール,\nGbp,ギガベースペア,unit\n   ; indented\n")
ok(pn.read_dictionary_file(cmt, "ja") == [("DNA", "ディーエヌエー", ""), ("Gbp", "ギガベースペア", "unit")],
   "; starts a comment anywhere in a dictionary line")
quoted = os.path.join(d, "quoted.csv")
write(quoted, '"#1",ナンバーワン,  ; a label in the figure\nC#,シーシャープ,\n; C,シー,\n')
ok([t for t, _, _ in pn.read_dictionary_file(quoted, "ja")] == ["#1", "C#"],
   "a # inside a term is kept now that only ; starts a comment")

ok(pn.difference_runs("abcdefghij", "abXXXXghij") == (1, 4) and pn.difference_runs("abc", "abc") == (0, 0),
   "difference runs count the places that differ and the longest one")
ok(pn.difference_list("0123456789abcdefghij", "0123456789XXXXefghij", context=3, minimum=2)
   == [(10, "abcd", "XXXX", "789", "efg")], "differences are listed with what was meant, what was heard and context")
_kana = "アイウエオカキクケコサシスセソタチツテトナニヌネノハヒフヘホマミムメモ" * 4
ok(pn.kana_sequence_scores(_kana, _kana[:300] + _kana[350:])[0] > 0.85, "a long note survives a 50-character gap by similarity")

c = pn.japanese_compounds("二本鎖の話")
ok(c == {} or "二本鎖" in c, "compounds are proposed with the reading the front end assembles (skipped without pyopenjtalk)")

# ---------------------------------------------------------------- packing
# Deck as PowerPoint writes it: slide 1 has no audio; slide 2 has a recorded narration with trim;
# slide 3 is a copy of slide 2 that shares its media file (its audio link is "NULL").
P14 = 'xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main"'
def audio_pic(link, embed, trim):
    return ('<p:pic><p:nvPicPr><p:cNvPr id="4" name="Audio 3"/><p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr>'
            f'<p:nvPr><a:audioFile r:link="{link}"/><p:extLst><p:ext uri="{{DAA4B4D4-6D71-4841-9C94-3DE7FCFB9230}}">'
            f'<p14:media {P14} r:embed="{embed}">{trim}</p14:media></p:ext></p:extLst></p:nvPr></p:nvPicPr>'
            '<p:blipFill><a:blip r:embed="rId90"/><a:stretch><a:fillRect/></a:stretch></p:blipFill><p:spPr><a:xfrm><a:off x="0" y="0"/>'
            '<a:ext cx="812800" cy="812800"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>')
transition = ('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><mc:Choice '
              f'{P14} Requires="p14"><p:transition spd="slow" p14:dur="2000" advTm="4289"/></mc:Choice><mc:Fallback>'
              '<p:transition spd="slow" advTm="4289"/></mc:Fallback></mc:AlternateContent>')
REL = '<Relationship Id="{}" Type="{}" Target="{}"{}/>'
media_rels = {
    2: [("rId91", pn.REL_AUDIO, "../media/media1.m4a", ""), ("rId92", pn.REL_MEDIA, "../media/media1.m4a", "")],
    3: [("rId91", pn.REL_AUDIO, "NULL", ' TargetMode="External"'), ("rId92", pn.REL_MEDIA, "../media/media1.m4a", "")],
}
z_in = zipfile.ZipFile(deck)
packed = os.path.join(d, "deck_audio.pptx")
z_out = zipfile.ZipFile(packed, "w", zipfile.ZIP_DEFLATED)
for item in z_in.infolist():
    data = z_in.read(item.filename)
    m = re.match(r"ppt/slides/(_rels/)?slide(\d)\.xml(\.rels)?$", item.filename)
    if item.filename == "[Content_Types].xml":
        data = data.replace(b"<Default ", b'<Default Extension="m4a" ContentType="audio/mp4"/><Default Extension="png" ContentType="image/png"/><Default ', 1)
    elif m and int(m.group(2)) in media_rels:
        n = int(m.group(2))
        if m.group(1):
            rels = "".join(REL.format(*r) for r in media_rels[n] + [("rId90", pn.REL_IMAGE, "../media/image1.png", "")])
            data = data.replace(b"</Relationships>", rels.encode() + b"</Relationships>")
        else:
            trim = '<p14:trim st="700.2867" end="1499.8366"/>'
            timing = pn._NARRATION_TIMING.replace("{spid}", "4")
            data = data.replace(b"</p:spTree>", audio_pic("rId91", "rId92", trim).encode() + b"</p:spTree>")
            data = re.sub(rb"(</p:clrMapOvr>)", lambda mm: mm.group(1) + transition.encode() + timing.encode(), data, count=1)
            recorded = ('<p:extLst><p:ext uri="{3A86A75C-4F4B-4683-9AE1-C65F6400EC91}"><p14:laserTraceLst ' + P14 + '>'
                        '<p14:tracePtLst><p14:tracePt t="48796" x="6062662" y="3259137"/></p14:tracePtLst>'
                        '</p14:laserTraceLst></p:ext><p:ext uri="{E180D4A7-C9FB-4DFB-919C-405C955672EB}">'
                        '<p14:showEvtLst ' + P14 + '><p14:playEvt time="12722" objId="4"/><p14:seekEvt time="38839" '
                        'objId="4" seek="10379"/></p14:showEvtLst></p:ext><p:ext uri="{BB962C8B}">'
                        '<p14:creationId ' + P14 + ' val="1"/></p:ext></p:extLst>')
            data = data.replace(b"</p:sld>", recorded.encode() + b"</p:sld>")
    z_out.writestr(item, data)
AudioSegment.silent(duration=300).export(os.path.join(d, "m.m4a"), format="ipod")
z_out.writestr("ppt/media/media1.m4a", open(os.path.join(d, "m.m4a"), "rb").read())
z_out.writestr("ppt/media/image1.png", pn._speaker_icon_png(8))
z_out.close()
for n, dur in [(1, 1800), (2, 2500), (3, 3200)]:
    AudioSegment.silent(duration=dur).export(os.path.join(ws, f"slide_{n}_de.v4.m4a"), format="ipod")
out_deck = os.path.join(d, "out.pptx")
orig_note1 = pn._slide_note_raw(Presentation(packed).slides[0])
pn.step_pack_pptx(packed, out_deck, ws, [1, 2, 3], "de", "v4")
notes = Presentation(out_deck).slides[0].notes_slide.notes_text_frame.text
secs = [s_ for s_ in pn.split_note_sections(notes.split("\n")) if s_["lang"]]
ok([s_["lang"] for s_ in secs] == ["de", "ja"] and notes.startswith("=== pptx-narrator: [de] translated from [ja]")
   and pn._section_text(secs[1]["lines"]) == orig_note1
   and pn._section_text(secs[0]["lines"]).startswith("[ja->de]")
   and secs[0]["info"]["source_lang"] == "ja"
   and secs[0]["info"]["source_fingerprint"] == pn.text_fingerprint("塩基対の話です。"),
   "write-back: the language just written is on top, where the presenter reads; the note keeps its text below")
notes2 = Presentation(out_deck).slides[1].notes_slide.notes_text_frame.text
ok("=== pptx-narrator: [en] ===" in notes2 and "=== pptx-narrator: [de] translated from [en]" in notes2,
   "write-back for an English-note slide: no language is treated specially")

zo = zipfile.ZipFile(out_deck)
names = zo.namelist()
sx = {n: zo.read(f"ppt/slides/slide{n}.xml").decode() for n in (1, 2, 3)}
rx = {n: zo.read(f"ppt/slides/_rels/slide{n}.xml.rels").decode() for n in (1, 2, 3)}
import xml.dom.minidom
for name in names:
    if name.endswith((".xml", ".rels")):
        xml.dom.minidom.parseString(zo.read(name))
ok(names[0] == "[Content_Types].xml", "packed file is a well-formed package")
adv = {n: [int(a) for a in re.findall(r'advTm="(\d+)"', sx[n])] for n in (1, 2, 3)}
ok(all(adv[n] and all(abs(a - dur - 1000) < 100 for a in adv[n]) for n, dur in [(1, 1800), (2, 2500), (3, 3200)]),
   f"slide advance times set to the audio length plus the pause {adv}")
ok('<a:audioFile r:link=' in sx[1] and 'isNarration="1"' in sx[1] and "pptx_narrator_slide1.m4a" in rx[1]
   and re.search(r"</p:clrMapOvr><p:transition[^>]*/><p:timing>", sx[1]), "narration inserted into a slide without audio")
ok("pptx_narrator_slide2.m4a" in rx[2] and "pptx_narrator_slide3.m4a" in rx[3] and "NULL" not in rx[3]
   and "ppt/media/media1.m4a" not in names, "copied slides get their own audio; the unused old clip is removed")
ok("p14:trim" not in sx[2] and "p14:trim" not in sx[3] and sx[2].count("<p:pic>") == 1, "trim of the previous recording removed")
offs = {n: re.findall(r'<a:off x="(-?\d+)" y="(-?\d+)"/>', sx[n]) for n in (1, 2, 3)}
ok(all(any(int(x) < 0 for x, _ in offs[n]) for n in (1, 2, 3)),
   f"audio icon parked outside the slide area {offs}")
pkg2 = tempfile.mkdtemp()
os.makedirs(os.path.join(pkg2, "ppt", "slides"))
write(os.path.join(pkg2, "ppt", "slides", "slide1.xml"),
      '<p:sld><p:cSld><p:spTree>' + audio_pic("rId91", "rId92", "") + '</p:spTree></p:cSld></p:sld>')
pn.embed_slide_narration(pkg2, 1, os.path.join(d, "m.m4a"), 300, icon_outside=False)
ok('<a:off x="0" y="0"/>' in read(os.path.join(pkg2, "ppt", "slides", "slide1.xml")),
   "--keep-audio-icon leaves the icon where it is")
ok(all("laserTraceLst" not in sx[n] and "showEvtLst" not in sx[n] for n in (2, 3)) and "creationId" in sx[2],
   "laser-pointer path and recorded playback events removed, other slide extensions kept")
kept = pn.remove_recorded_show_data(open(os.path.join(d, "recorded.xml"), encoding="utf-8").read()
                                    if False else '<p:sld><p:extLst><p:ext uri="{x}"><p14:laserTraceLst/></p:ext>'
                                    '<p:ext uri="{y}"><p14:showEvtLst/></p:ext></p:extLst></p:sld>', ("pointer",))
ok("laserTraceLst" not in kept[0] and "showEvtLst" in kept[0] and kept[1] == ["laser-pointer path"],
   "--remove-recorded pointer keeps the recorded playback events")
ok(len(Presentation(out_deck).slides) == 3, "packed deck opens with python-pptx")

logs.clear()
pkg = tempfile.mkdtemp()
os.makedirs(os.path.join(pkg, "ppt", "slides"))
write(os.path.join(pkg, "ppt", "slides", "slide1.xml"), '<p:sld><p:cSld><p:spTree></p:spTree></p:cSld><p:timing></p:timing></p:sld>')
ok(pn.embed_slide_narration(pkg, 1, os.path.join(d, "m.m4a"), 300) is None and any("animations" in m for m in logs),
   "slide with animations but no audio object is left alone with a warning")

xml_trim = ('<p14:media r:embed="rId2"><p14:trim st="1200" end="800"/><p14:fade in="500"/><p14:bmkLst><p14:bmk name="a" time="1"/></p14:bmkLst></p14:media>'
            '<p14:media r:embed="rId9"><p14:trim st="10"/></p14:media>')
cleared, n_cleared = pn.clear_media_playback_settings(xml_trim, {"rId2"})
ok(n_cleared == 1 and cleared.startswith('<p14:media r:embed="rId2"/>') and '<p14:trim st="10"/>' in cleared,
   "trim/fade/bookmarks of the replaced audio removed, other media untouched")

# ---------------------------------------------------------------- notes in the layout of earlier versions
def compose_structured_note(narration_lang, narration_text, source_lang, source_text,
                            fingerprint=None, spoken=False):
    """A note as earlier versions wrote it (narration above, source below)."""
    fingerprint = fingerprint or pn.text_fingerprint(source_text)
    head = (f"=== pptx-narrator: narration [{narration_lang}] from [{source_lang}] "
            f"#{fingerprint}{' spoken' if spoken else ''} ===")
    return (f"{head}\n{narration_text.strip()}\n\n"
            f"=== pptx-narrator: source [{source_lang}] ===\n{source_text.strip()}")

note = compose_structured_note("en", "Today we talk about DNA.", "ja", "今日はDNAの話です。")
secs = [s_ for s_ in pn.split_note_sections(note.split("\n")) if s_["lang"]]
ok([s_["lang"] for s_ in secs] == ["en", "ja"] and pn._section_text(secs[0]["lines"]) == "Today we talk about DNA."
   and pn._section_text(secs[1]["lines"]) == "今日はDNAの話です。" and secs[0]["info"]["source_lang"] == "ja",
   "a note in the layout of earlier versions is read into one text per language")
ok([s_["lang"] for s_ in pn.split_note_sections(["普通のノート"])] == [None], "an ordinary note has no headings")
ok(pn.text_fingerprint("a\r\nb  \n") == pn.text_fingerprint("a\nb"), "fingerprint ignores line endings and trailing spaces")

# extract the packed deck again: source part is extracted, unchanged translation restored
ws4 = os.path.join(d, "ws4"); os.makedirs(ws4)
pn.step_extract_notes(out_deck, ws4, [1, 2, 3], "auto")
ok(read(os.path.join(ws4, "slide_1_ja.txt")) == orig_note1 and read(os.path.join(ws4, "slide_1_de.txt")).startswith("[ja->de]"),
   "re-extraction: one text per language of the note")
ok(pn._load_manifest(ws4)["1"]["de"]["source_fingerprint"] == pn.text_fingerprint("塩基対の話です。"), "manifest written on restore")
FakeTr.calls.clear()
pn.step_translate_notes(ws4, [1], "auto", "de", dictionary=T)
ok(FakeTr.calls == [], "restored translation is not translated again")

# edit the source part inside PowerPoint -> narration is stale
prs_e = Presentation(out_deck)
tf = prs_e.slides[0].notes_slide.notes_text_frame
tf.text = tf.text.replace(orig_note1, "塩基対とRNAの話です。")
edited = os.path.join(d, "edited.pptx"); prs_e.save(edited)
ws5 = os.path.join(d, "ws5"); os.makedirs(ws5)
write(os.path.join(ws5, "slide_1_de.txt"), "old narration in the workspace")
logs.clear()
pn.step_extract_notes(edited, ws5, [1], "auto")
ok(read(os.path.join(ws5, "slide_1_de.txt")) == "old narration in the workspace"
   and any("slide_1_de.txt differs from the note in the deck and was edited in the workspace" in m for m in logs),
   "extract keeps a text of the workspace that it did not write, with a warning")
pn.step_extract_notes(edited, ws5, [1], "auto", overwrite=True)
ok(read(os.path.join(ws5, "slide_1_de.txt")).startswith("[ja->de]") and not os.path.exists(os.path.join(ws5, "slide_1_de.stale.txt")),
   "--overwrite replaces it with the note; nothing is set aside")
pn.step_translate_notes(ws5, [1], "auto", "de", dictionary=T, update=True)
ok("RNA" in read(os.path.join(ws5, "slide_1_de.txt")), "a translation whose source changed is translated again with --update")
write(os.path.join(ws5, "slide_1_de.txt"), "von Hand")
logs.clear()
pn.step_extract_notes(edited, ws5, [1], "auto", update=True)
ok(read(os.path.join(ws5, "slide_1_de.txt")) == "von Hand",
   "--update of extract keeps a text edited in the workspace")

# spoken-form narration blocks are not restored as translations
spoken_note = compose_structured_note("de", "Boten-RNA", "ja", "mRNAの話", spoken=True)
prs_s = Presentation(); s_ = prs_s.slides.add_slide(prs_s.slide_layouts[5]); s_.notes_slide.notes_text_frame.text = spoken_note
sp = os.path.join(d, "spoken.pptx"); prs_s.save(sp)
ws6 = os.path.join(d, "ws6"); os.makedirs(ws6)
pn.step_extract_notes(sp, ws6, [1], "auto")
ok(sorted(os.listdir(ws6)) == ["note_baseline.json", "slide_1_ja.txt"], "spoken narration block ignored on extraction")

# pack warns when the translation is older than the source note
logs.clear()
write(os.path.join(ws, "slide_1_ja.txt"), "塩基対の話を変更しました。")
pn.step_pack_pptx(packed, os.path.join(d, "out2.pptx"), ws, [1], "de", "v4")
note_old = Presentation(os.path.join(d, "out2.pptx")).slides[0].notes_slide.notes_text_frame.text
ok(any("was translated from an older version of slide_1_ja.txt" in m for m in logs)
   and "#" + pn.text_fingerprint("塩基対の話です。") in note_old,
   "stale translation flagged at pack time and marked by the fingerprint of its source")

# ---------------------------------------------------------------- struck-through text in notes
prs12 = Presentation()
tf12 = prs12.slides.add_slide(prs12.slide_layouts[5]).notes_slide.notes_text_frame
para = tf12.paragraphs[0]
for text, strike in [("教科書の", None), ("RNA", "sngStrike"), ("DNAプライマーゼ", None),
                     ("（旧称）", "dblStrike"), ("です。", None)]:
    run = para.add_run(); run.text = text
    if strike:
        run._r.get_or_add_rPr().set("strike", strike)
tf12.add_paragraph().text = "二段落目です。"
deck12 = os.path.join(d, "deck12.pptx"); prs12.save(deck12)
ws12 = os.path.join(d, "ws12"); os.makedirs(ws12)
pn.step_extract_notes(deck12, ws12, [1], "ja")
ok(read(os.path.join(ws12, "slide_1_ja.txt")) == "教科書のDNAプライマーゼです。\n二段落目です。",
   "single and double struck-through text is left out of the extracted note")
AudioSegment.silent(duration=300).export(os.path.join(ws12, "slide_1_ja.qwen3-1.7B.m4a"), format="ipod")
_, written12, _ = pn.step_pack_pptx(deck12, os.path.join(d, "out12.pptx"), ws12, [1], "ja", "qwen3-1.7B",
                                    update=True)
ok(written12 == [] and "RNA" in Presentation(os.path.join(d, "out12.pptx")).slides[0].notes_slide.notes_text_frame.text,
   "an unedited note with struck-through text is not rewritten (the strikethrough stays in the deck)")
_, written12, _ = pn.step_pack_pptx(deck12, os.path.join(d, "out12.pptx"), ws12, [1], "ja", "qwen3-1.7B")
ok(written12 == [], "the same text is not written again even without --update")
write(os.path.join(ws12, "slide_1_en.txt"), "This is the textbook's DNA primase.")
_, written12, _ = pn.step_pack_pptx(deck12, os.path.join(d, "out12b.pptx"), ws12, [1], None, "qwen3-1.7B",
                                    targets={"text"})
body12 = Presentation(os.path.join(d, "out12b.pptx")).slides[0].notes_slide.notes_text_frame._txBody
from lxml import etree as _et; xml12 = _et.tostring(body12, encoding="unicode")
secs12 = [s_ for s_ in pn.split_note_sections(Presentation(os.path.join(d, "out12b.pptx")).slides[0].notes_slide.notes_text_frame.text.split("\n")) if s_["lang"]]
ok(written12 == [1] and [s_["lang"] for s_ in secs12] == ["en", "ja"]
   and 'strike="sngStrike"' in xml12 and 'strike="dblStrike"' in xml12,
   "adding a language keeps the other part of the note as it was, formatting and struck-through text included")
ws12b = os.path.join(d, "ws12b"); os.makedirs(ws12b)
pn.step_extract_notes(os.path.join(d, "out12b.pptx"), ws12b, [1], "auto")
ok(read(os.path.join(ws12b, "slide_1_ja.txt")) == read(os.path.join(ws12, "slide_1_ja.txt"))
   and read(os.path.join(ws12b, "slide_1_en.txt")) == "This is the textbook's DNA primase.",
   "such a note is extracted again into the same texts")
import time as _time
_time.sleep(1.1)
write(os.path.join(ws12, "slide_1_ja.txt"), read(os.path.join(ws12, "slide_1_ja.txt")) + "直しました。")
_, written12, _ = pn.step_pack_pptx(os.path.join(d, "out12b.pptx"), os.path.join(d, "out12d.pptx"), ws12, [1], "ja",
                                    "qwen3-1.7B", targets={"text"}, update=True)
secs12 = [s_ for s_ in pn.split_note_sections(Presentation(os.path.join(d, "out12d.pptx")).slides[0].notes_slide.notes_text_frame.text.split("\n")) if s_["lang"]]
ok(written12 == [1] and [s_["lang"] for s_ in secs12] == ["ja", "en"],
   "the part written last goes on top; older parts move down")
_, written12, _ = pn.step_pack_pptx(os.path.join(d, "out12d.pptx"), os.path.join(d, "out12e.pptx"), ws12, [1], None,
                                    "qwen3-1.7B", targets={"text"})
ok(written12 == [] and pn._slide_note_raw(Presentation(os.path.join(d, "out12e.pptx")).slides[0])
   == pn._slide_note_raw(Presentation(os.path.join(d, "out12d.pptx")).slides[0]),
   "a pack that writes nothing leaves the order of the note as it is")
AudioSegment.silent(duration=300).export(os.path.join(ws12, "slide_1_en.qwen3-1.7B.m4a"), format="ipod")
logs.clear()
_, written12, _ = pn.step_pack_pptx(os.path.join(d, "out12d.pptx"), os.path.join(d, "out12f.pptx"), ws12, [1], None,
                                    "qwen3-1.7B", audio_lang="en")
note12f = Presentation(os.path.join(d, "out12f.pptx")).slides[0].notes_slide.notes_text_frame
secs12 = [s_ for s_ in pn.split_note_sections(note12f.text.split("\n")) if s_["lang"]]
xml12f = _et.tostring(note12f._txBody, encoding="unicode")
ok([s_["lang"] for s_ in secs12] == ["en", "ja"] and any("moved to the top" in m for m in logs)
   and written12 == [1],
   "the text of the language whose audio the slide plays is always on top (moved there unchanged)")
write(os.path.join(ws12, "slide_1_ja.txt"), read(os.path.join(ws12, "slide_1_ja.txt")) + "もう一度直しました。")
_, written12, _ = pn.step_pack_pptx(os.path.join(d, "out12f.pptx"), os.path.join(d, "out12g.pptx"), ws12, [1], None,
                                    "qwen3-1.7B", audio_lang="en", update=True)
secs12 = [s_ for s_ in pn.split_note_sections(Presentation(os.path.join(d, "out12g.pptx")).slides[0].notes_slide.notes_text_frame.text.split("\n")) if s_["lang"]]
ok(written12 == [1] and [s_["lang"] for s_ in secs12] == ["en", "ja"],
   "a text written for another language goes below the text of the audio")
ok(pn.record_audio_sources(ws12, [1], "ja", "qwen3-1.7B", None, 0) == [1]
   and pn.audio_is_older_than_text(ws12, 1, "ja", "qwen3-1.7B") is False
   and pn.audio_is_older_than_text(ws12, 1, "en", "qwen3-1.7B") is None,
   "synthesize records the text each audio was made from")
write(os.path.join(ws12, "slide_1_ja.txt"), read(os.path.join(ws12, "slide_1_ja.txt")) + "追加の文です。")
logs.clear()
pn.step_pack_pptx(deck12, os.path.join(d, "out12c.pptx"), ws12, [1], "ja", "qwen3-1.7B", targets={"audio"})
ok(pn.audio_is_older_than_text(ws12, 1, "ja", "qwen3-1.7B") is True
   and any("was edited after slide_1_ja.qwen3-1.7B.m4a was made" in m for m in logs),
   "pack says when a text was edited after its audio was made")

# ---------------------------------------------------------------- pack: notes and hand edits
def note_of(path, n=1):
    return Presentation(path).slides[n - 1].notes_slide.notes_text_frame.text

prs10 = Presentation()
for t in ["今日はDNAの話です。", "今日はRNAの話です。"]:
    prs10.slides.add_slide(prs10.slide_layouts[5]).notes_slide.notes_text_frame.text = t
deck10 = os.path.join(d, "deck10.pptx"); prs10.save(deck10)
ws10 = os.path.join(d, "ws10"); os.makedirs(ws10)
pn.step_extract_notes(deck10, ws10, [1], "auto")
pn.step_extract_notes(deck10, ws10, [2], "auto")
ok(sorted(pn._load_note_baseline(ws10)) == ["1", "2"],
   "extracting some slides keeps the record of the slides extracted before")
for n in (1, 2):
    AudioSegment.silent(duration=400 + 100 * n).export(os.path.join(ws10, f"slide_{n}_ja.qwen3-1.7B.m4a"), format="ipod")

# the documented workflow: edit the extracted text, synthesize, pack -> the note follows
write(os.path.join(ws10, "slide_1_ja.txt"), "今日はDNAの構造の話です。")
out10 = os.path.join(d, "out10.pptx")
logs.clear()
_, written10, mismatched10 = pn.step_pack_pptx(deck10, out10, ws10, [1, 2], "ja", "qwen3-1.7B")
ok(note_of(out10, 1) == "今日はDNAの話です。" and written10 == [] and mismatched10 == [1]
   and any("slide_1_ja.txt was edited in the workspace" in m for m in logs),
   "a text edited in the workspace is not written without --update or --overwrite, with a warning")
logs.clear()
_, written10, mismatched10 = pn.step_pack_pptx(deck10, out10, ws10, [1, 2], "ja", "qwen3-1.7B", update=True)
ok(note_of(out10, 1) == "今日はDNAの構造の話です。" and note_of(out10, 2) == "今日はRNAの話です。"
   and written10 == [1] and mismatched10 == [] and not any("WARNING" in m for m in logs if "Slide #1" in m),
   "with --update, text edited in the workspace is written into the note (the deck note was not touched)")

# target audio: the note is left, and the mismatch is reported
logs.clear()
_, written, mismatched = pn.step_pack_pptx(deck10, os.path.join(d, "out10a.pptx"), ws10, [1, 2], "ja", "qwen3-1.7B",
                                           targets={"audio"})
ok(note_of(os.path.join(d, "out10a.pptx"), 1) == "今日はDNAの話です。" and written == [] and mismatched == [1]
   and any("Slide #1: the note in the deck differs" in m for m in logs),
   "--data-type audio leaves the notes and warns where they no longer match the narration")

# a note edited in the deck after extract is not overwritten ...
prs10h = Presentation(out10)
prs10h.slides[0].notes_slide.notes_text_frame.text = "手で直したノート"
hand10 = os.path.join(d, "hand10.pptx"); prs10h.save(hand10)
write(os.path.join(ws10, "slide_1_ja.txt"), "今日はDNAの二重らせんの話です。")
logs.clear()
_, written, mismatched = pn.step_pack_pptx(hand10, os.path.join(d, "out10b.pptx"), ws10, [1], "ja", "qwen3-1.7B")
ok(note_of(os.path.join(d, "out10b.pptx"), 1) == "手で直したノート" and written == [] and mismatched == [1]
   and any("Slide #1: the note was edited in the deck after extract" in m for m in logs),
   "a note edited in the deck is protected, with a warning")
# ... unless --overwrite (forceupdate=True)
logs.clear()
_, written, _ = pn.step_pack_pptx(hand10, os.path.join(d, "out10c.pptx"), ws10, [1], "ja", "qwen3-1.7B",
                                  overwrite=True)
ok(note_of(os.path.join(d, "out10c.pptx"), 1) == "今日はDNAの二重らせんの話です。" and written == [1]
   and any("overwriting it (--overwrite)" in m for m in logs),
   "--overwrite overwrites it, with a warning")
# a note pack itself wrote is not mistaken for a hand edit
write(os.path.join(ws10, "slide_1_ja.txt"), "今日はDNAの複製の話です。")
_, written, _ = pn.step_pack_pptx(os.path.join(d, "out10c.pptx"), os.path.join(d, "out10d.pptx"), ws10, [1],
                                  "ja", "qwen3-1.7B", update=True)
ok(written == [1] and note_of(os.path.join(d, "out10d.pptx"), 1) == "今日はDNAの複製の話です。",
   "a note written by pack can be replaced by the next pack")

# --update leaves out what is already the same
logs.clear()
embedded, written, _ = pn.step_pack_pptx(os.path.join(d, "out10d.pptx"), os.path.join(d, "out10e.pptx"), ws10,
                                         [1], "ja", "qwen3-1.7B", update=True)
ok(embedded == 0 and written == [] and any("already has this audio" in m for m in logs),
   "--update leaves out the audio and the note that the deck already has")

# no record of extract: a hand edit cannot be ruled out
ws11 = os.path.join(d, "ws11"); os.makedirs(ws11)
write(os.path.join(ws11, "slide_1_ja.txt"), "別のテキスト")
_, written, _ = pn.step_pack_pptx(deck10, os.path.join(d, "out11.pptx"), ws11, [1], "ja", "qwen3-1.7B",
                                  targets={"text"})
ok(written == [] and note_of(os.path.join(d, "out11.pptx"), 1) == "今日はDNAの話です。",
   "without a record of extract, the note is not overwritten")
ok(pn.resolve_targets(None) == pn.resolve_targets("all") == {"audio", "text"}
   and pn.resolve_targets("text") == {"text"}, "pack target: omitted means all")

# ---------------------------------------------------------------- command line
# The command line is  pptx-narrator WS COMMAND [INPUT] [OUTPUT] [options].
# This section runs in a directory of its own, so that nothing is written beside the tests.
deck = os.path.abspath(deck)
cli_cwd = os.path.join(d, "cli"); os.makedirs(cli_cwd); os.chdir(cli_cwd)
os.makedirs(os.path.join(cli_cwd, "ws"))
for _lang in ("ja", "de", "nl"):
    write(os.path.join(cli_cwd, "ws", f"slide_1_{_lang}.txt"), "text")
write(os.path.join(cli_cwd, "ref.wav"), "not really audio")
write(os.path.join(cli_cwd, "ref.txt"), "transcript")
parser = pn.build_parser()
a = parser.parse_args(["ws", "translate", "--in_lang", "JA", "--out-lang", "zh_cn",
                       "--dict-file", "a.csv", "--dict_file", "b.csv"])
ok(a.workspace == "ws" and a.command == "translate" and a.in_lang == "ja" and a.out_lang == "zh-CN"
   and a.dict_file == ["a.csv", "b.csv"], "CLI normalization, aliases, repeatable --dict-file")

positionals = {"extract": [deck], "scan": ["d.csv"], "pack": [deck, "o.pptx"]}
ok([c for c in pn.COMMANDS if parser.parse_args(["ws", c] + positionals.get(c, [])).command != c] == [],
   "every pipeline step is a command of its own, after the workspace")
a = parser.parse_args(["ws", "pack", deck, "o.pptx", "--lang", "ja"])
ok(a.deck == deck and a.out_deck == "o.pptx" and a.lang == "ja"
   and parser.parse_args(["ws", "scan", "d.csv"]).dictionary == "d.csv"
   and parser.parse_args(["ws", "extract", deck]).deck == deck,
   "INPUT and OUTPUT are the files outside the workspace, given by position")

def expect_error(argv, text):
    buf = io.StringIO()
    try:
        with contextlib.redirect_stderr(buf):
            pn.main(argv)
    except SystemExit:
        pass
    ok(text in buf.getvalue(), f"CLI error: {text}")

expect_error(["ws", "extract"], "DECK")
expect_error(["ws", "pack", deck], "OUT")
expect_error(["ws", "synthesize", "--lang", "nl", "--engine", "qwen3",
              "--ref-wav", "ref.wav", "--ref-text", "ref.txt"], "Qwen3-TTS does not support 'nl'")
expect_error(["ws", "synthesize", "--lang", "de", "--engine", "gpt_sovits",
              "--ref-wav", "ref.wav", "--ref-text", "ref.txt"], "GPT-SoVITS does not support")
expect_error(["ws", "synthesize", "--lang", "ja", "--ref-wav", "no_such.wav", "--ref-text", "ref.txt"],
             "--ref-wav: no such file: no_such.wav")
expect_error(["ws", "scan", "d.csv"], "pptx-narrator WS scan: error:")
expect_error(["ws", "scan", "d.csv"], "usage: pptx-narrator WS scan")
expect_error(["ws", "translate", "--in-lang", "de", "--out-lang", "de"], "translate requires different")
expect_error(["ws", "translate", "--lang", "de"], "translate requires different")
expect_error(["ws", "translate", "--lang", "de", "--in-lang", "ja", "--out-lang", "en"],
             "--lang cannot be combined with --in-lang or --out-lang")
expect_error(["ws", "scan", "d.csv"], "scan requires --lang")
expect_error(["ws_fr", "extract", deck, "--in-lang", "fr"], "no note in fr")
expect_error(["no_such_ws", "scan", "d.csv", "--lang", "ja"], "workspace does not exist")
expect_error(["extract", deck], "the workspace comes first and the command second")
expect_error(["ws", "extrct", deck], "did you mean 'extract'")
write(os.path.join(cli_cwd, "terms.csv"), "string,replacement,type\n")
expect_error(["ws", "scan", "terms.csv", "--lang", "ja"], "already exists. Add --append")
expect_error(["ws", "scan", "terms.csv", "--lang", "ja", "--append", "--overwrite"],
             "--append and --overwrite cannot be combined")
expect_error(["ws", "pack", deck, deck, "--lang", "ja"], "OUT must differ from DECK")
shutil.copy(deck, os.path.join(cli_cwd, "made_before.pptx"))
expect_error(["ws", "pack", deck, "made_before.pptx", "--lang", "ja", "--data-type", "text"],
             "already exists. Add --update or --overwrite")
expect_error(["ws", "pack", deck, "o.pptx", "--lang", "ja", "--update", "--overwrite"],
             "--update and --overwrite cannot be combined")

# built-in defaults -> configuration file -> command line
args = parser.parse_args(["ws", "pack", deck, "o.pptx"])
eff = pn._merge_effective("pack", args, {}, parser)
ok(eff["remove_recorded"] == "all" and eff["slide_pause"] == 1.0 and eff["data_type"] == "all"
   and pn.RECORDED_CHOICES["none"] == (), "built-in defaults are used when nothing else sets a value")
eff = pn._merge_effective("pack", args, {"pack": {"slide_pause": 2.5}}, parser)
ok(eff["slide_pause"] == 2.5, "the configuration file overrides the built-in default")
eff = pn._merge_effective("pack", args, {"common": {"slide_pause": 3.5}}, parser)
ok(eff["slide_pause"] == 3.5, "a [common] section applies to every command")
args = parser.parse_args(["ws", "pack", deck, "o.pptx", "--slide-pause", "0.5"])
eff = pn._merge_effective("pack", args, {"pack": {"slide_pause": 2.5}}, parser)
ok(eff["slide_pause"] == 0.5, "the command line overrides the configuration file")
eff = pn._merge_effective("translate", parser.parse_args(["ws", "translate", "--dict-file", ""]),
                          {"translate": {"dict_file": ["from_config.csv"]}}, parser)
ok(eff["dict_file"] is None, "an empty value on the command line takes a configured value back")
eff = pn._merge_effective("pack", parser.parse_args(["ws", "pack", deck, "o.pptx"]),
                          {"pack": {"out": "x.pptx", "forceupdate": True, "workspace": "w"}}, parser)
ok("out" not in eff and "forceupdate" not in eff and "workspace" not in eff and eff["overwrite"] is False,
   "configuration keys of earlier versions are set aside, not acted on")
ok(parser.parse_args(["ws", "pack", deck, "o.pptx", "--data-type", "text"]).data_type == "text",
   "pack takes what to write as --data-type")
cfg_true = {"pack": {"update": True}}
ok(pn._merge_effective("pack", parser.parse_args(["ws", "pack", deck, "o.pptx"]), cfg_true, parser)["update"] is True
   and pn._merge_effective("pack", parser.parse_args(["ws", "pack", deck, "o.pptx", "--no-update"]),
                           cfg_true, parser)["update"] is False,
   "--no-... takes back a switch set in the configuration file")
ok(pn._normalize_config_keys({"a-b": {"c-d": 1}}) == {"a_b": {"c_d": 1}},
   "hyphenated keys in the configuration file are accepted")
ok(parser.parse_args(["ws", "scan", "d.csv", "--config", "x.toml"]).config == "x.toml",
   "--config may follow the command")

# --lang is --in-lang (and --out-lang) in one
def lang_of(argv, config=None):
    a = parser.parse_args(argv)
    return pn._resolve_lang(a.command, pn._merge_effective(a.command, a, config or {}, parser), a, parser)
e = lang_of(["ws", "synthesize", "--lang", "ja"])
ok(e["in_lang"] == "ja" and "lang" not in e, "--lang gives the language of a one-language command")
e = lang_of(["ws", "translate", "--lang", "ja"])
ok(e["in_lang"] == e["out_lang"] == "ja", "--lang sets --in-lang and --out-lang together")
e = lang_of(["ws", "synthesize", "--in-lang", "ja"], {"synthesize": {"lang": "en"}})
ok(e["in_lang"] == "ja", "--in-lang on the command line wins over lang in the configuration file")

# ------------------------------------------------ no implicit carry-over between runs
ok(not hasattr(pn, "_inherit_synthesis_settings") and not hasattr(pn, "_resolve_input"),
   "nothing is carried over implicitly from an earlier run")
wsi = os.path.join(cli_cwd, "wsi"); os.makedirs(wsi)
write(os.path.join(wsi, "slide_1_ja.txt"), "\u30c6\u30b9\u30c8")
AudioSegment.silent(duration=300).export(os.path.join(wsi, "slide_1_ja.qwen3-1.7B.m4a"), format="ipod")
write(os.path.join(wsi, "slide_1_ja.qwen3-1.7B.spoken.txt"), "\u30c6\u30b9\u30c8")
pn._save_state({"version": 1, "commands": {"synthesize": {
    "resolved_config": {"in_lang": "ja", "engine": "gpt_sovits", "model": "v2ProPlus"}}}},
    pn._workspace_path(wsi, pn.STATE_FILE))
a = parser.parse_args(["wsi", "pack", deck, "o.pptx", "--lang", "ja"])
ok(pn._merge_effective("pack", a, {}, parser)["engine"] == "qwen3",
   "pack does not take over the engine of the last synthesis")

AudioSegment.silent(duration=300).export(os.path.join(wsi, "slide_1_ja.v2ProPlus.m4a"), format="ipod")
ok(pn._audio_model_labels(wsi, "ja") == ["qwen3-1.7B", "v2ProPlus"],
   "the audio of a workspace is reported by model label")
expect_error(["wsi", "pack", deck, "o.pptx", "--lang", "ja"], "multiple audio model labels")

wse = os.path.join(cli_cwd, "wse"); os.makedirs(wse)
write(os.path.join(wse, "slide_1_ja.txt"), "\u30c6\u30b9\u30c8")
logs.clear()
with contextlib.redirect_stderr(io.StringIO()):
    pn.main(["wse", "pack", deck, "o2.pptx", "--lang", "ja", "--engine", "qwen3"])
ok(os.path.exists(os.path.join(cli_cwd, "o2.pptx")) and any("No ja audio for model 'qwen3-1.7B'" in m for m in logs),
   "pack without audio still writes the texts, and says the audio was not found")
AudioSegment.silent(duration=300).export(os.path.join(wse, "slide_1_en.qwen3-1.7B.m4a"), format="ipod")
AudioSegment.silent(duration=300).export(os.path.join(wse, "slide_1_ja.qwen3-1.7B.m4a"), format="ipod")
expect_error(["wse", "pack", deck, "o3.pptx"], "audio of several languages is in this workspace (en, ja)")
logs.clear()
with contextlib.redirect_stderr(io.StringIO()):
    pn.main(["wse", "pack", deck, "o3.pptx", "--data-type", "text"])
ok(os.path.exists(os.path.join(cli_cwd, "o3.pptx")), "without --lang, the texts of every language can be written")

# the record of a run: in the workspace, with the paths inside it relative to it
ws2 = os.path.join(cli_cwd, "ws2")
said = io.StringIO()
with contextlib.redirect_stderr(said):
    pn.main(["ws2", "extract", deck, "--lang", "ja"])
said = said.getvalue()
ok("Next, for example" in said and 'pptx-narrator ws2 synthesize --lang ja --ref-wav "ref.wav"' in said
   and said.index("Next, for example") < said.index("Report of extract")
   and "created in the workspace: " in said and "slide_1_ja.txt" in said,
   "a run ends with the commands that can come next and then a report of the files it wrote")
log2 = read(os.path.join(ws2, pn.LOG_FILE))
ok("==== pptx-narrator ws2 extract" in log2 and "note extracted to slide_1_ja.txt" in log2 and "Report of extract" in log2,
   "what a run logs and reports is also kept in the log file of the workspace")
hist, status = io.StringIO(), 0
with contextlib.redirect_stdout(hist):
    pn.main(["ws2", "history", "--dates"])
ok("extract " + deck in hist.getvalue() and "--in-lang ja" in hist.getvalue()
   and "file(s) written" in hist.getvalue(), "history lists the commands run in the workspace")
ok(os.path.exists(os.path.join(ws2, ".pptx_narrator_history.jsonl"))
   and not os.path.exists(os.path.join(cli_cwd, ".pptx_narrator_history.jsonl"))
   and not os.path.exists(pn._workspace_path(ws2, pn.STATE_FILE)),
   "the record of a run lives in the workspace, not in the current directory; no state file is written")
with open(os.path.join(ws2, ".pptx_narrator_history.jsonl"), encoding="utf-8") as f:
    last = json.loads(f.read().splitlines()[-1])
resolved = read(os.path.join(ws2, pn.RESOLVED_CONFIG_FILE))
ok(last["command"] == "extract" and last["paths"]["deck"] == deck and last["input"]["path"] == deck
   and "slide_1_ja.txt" in last["written"]["created"]
   and 'workspace = "."' in resolved and os.path.abspath(ws2) not in json.dumps(last["paths"]),
   "the deck outside the workspace is recorded by its absolute path, the workspace as '.'")
rec = pn._record_effective({"dict_file": [os.path.join(ws2, "d.csv"), "/abs/x.csv"],
                            "ref_wav": os.path.join(ws2, "voice", "ref.wav"), "slides": "1-3"}, ws2)
ok(rec["dict_file"] == ["d.csv", "/abs/x.csv"] and rec["ref_wav"] == os.path.join("voice", "ref.wav")
   and rec["slides"] == "1-3", "path settings inside the workspace are recorded relative to it")

# a failure during a run names the command and the cause, and is recorded as a failure
_orig_verify = pn.step_verify_audio
def _failing_verify(*a, **k):
    raise RuntimeError("the ASR model could not be loaded")
pn.step_verify_audio = _failing_verify
logs.clear()
try:
    with contextlib.redirect_stderr(io.StringIO()):
        pn.main(["wsi", "verify", "--lang", "ja", "--engine", "qwen3"])
    status = 0
except SystemExit as e:
    status = e.code
pn.step_verify_audio = _orig_verify
with open(os.path.join(wsi, ".pptx_narrator_history.jsonl"), encoding="utf-8") as f:
    last = json.loads(f.read().splitlines()[-1])
ok(status == 1 and any("verify failed: RuntimeError: the ASR model could not be loaded" in m for m in logs)
   and last["status"] == "failed" and "ASR model" in last["error"]
   and "Traceback" in read(os.path.join(wsi, pn.LOG_FILE)),
   "a failure names the command and its cause; the traceback goes to the log, the failure to the history")
pn._CONFIG_PATH = "/somewhere/pptx_narrator.toml"
logs.clear()
pn._merge_effective("pack", parser.parse_args(["ws", "pack", deck, "o.pptx"]), {"common": {"workspace": "w"}}, parser)
pn._CONFIG_PATH = None
ok(any("'workspace' in /somewhere/pptx_narrator.toml is no longer used (the workspace is now the first argument"
       in m and "Remove it from the file" in m for m in logs),
   "a key of an earlier version is reported with its file and what to do")
from lxml import etree as _et2
_p = _et2.fromstring('<a:p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:r><a:t>2026/9/26\u200bに発表</a:t></a:r>'
                     '<a:fld type="datetime1"><a:t>2026/9/26</a:t></a:fld><a:fld type="slidenum"><a:t>3</a:t></a:fld></a:p>')
ok(pn._paragraph_text(_p) == "2026/9/26\u200bに発表" and "2026/9/26\u200bに".translate(pn._ZERO_WIDTH) == "2026/9/26に",
   "date and slide-number fields are left out, a date the author typed is kept; zero-width characters are removed")

ok(not hasattr(parser.parse_args(["ws", "pack", deck, "o.pptx"]), "use_spoken_notes")
   and "--use-spoken-notes" not in open(os.path.join(os.path.dirname(__file__), "..", "README.md"),
                                        encoding="utf-8").read(),
   "the spoken text is never written into the notes")

ok(pn._rel_to_workspace("/w/ws/slide_1_ja.txt", "/w/ws") == "slide_1_ja.txt"
   and pn._rel_to_workspace("/w/ws", "/w/ws") == "."
   and pn._rel_to_workspace("/elsewhere/deck.pptx", "/w/ws") == "/elsewhere/deck.pptx",
   "a path inside the workspace is recorded relative to it")
snap = {"kind": "directory", "path": "/w/ws",
        "files": [{"path": "/w/ws/slide_1_ja.txt", "sha256": "x"}]}
ok(pn._relativize_snapshot(snap, "/w/ws")["files"][0]["path"] == "slide_1_ja.txt",
   "the files of a directory record are relative too")

ok(all("--slides" in [a.option_strings[0] for a in sp._actions if a.option_strings]
       for name, sp in [(n, sp) for act in parser._actions
                        if isinstance(act, argparse._SubParsersAction)
                        for n, sp in act.choices.items()]),
   "every command can be limited to a slide selection")
ok(pn._select_slides([1, 2, 3, 7], "2,7", parser) == [2, 7]
   and pn._select_slides([1, 2, 3], None, parser) == [1, 2, 3],
   "--slides narrows the slides of a workspace")
expect_error(["wse", "synthesize", "--lang", "ja", "--slides", "99",
              "--ref-wav", "ref.wav", "--ref-text", "ref.txt"], "no slide of this workspace matches")

def run_cli(argv):
    """Run the CLI and return (stdout, exit status)."""
    out = io.StringIO()
    status = 0
    try:
        with contextlib.redirect_stdout(out):
            pn.main(argv)
    except SystemExit as e:
        status = e.code if isinstance(e.code, int) else 1
    return out.getvalue(), status

top, status = run_cli(["--help"])
ok(status == 0 and top.count("extract") == 1,
   "--help lists every command once and is not an error")
ok(run_cli(["--version"])[0].strip().endswith(pn.__version__)
   and run_cli(["--version"])[1] == 0,
   "--version prints the version instead of the help")
ok(run_cli([])[1] != 0, "no command at all is still an error")
ok("WS COMMAND --help" in top, "the help says how to see the options of a command")

sub_help = run_cli(["ws", "verify", "--help"])[0]
ok("--in_lang" not in sub_help and "--in-lang" in sub_help and "--lang" in sub_help,
   "an option is listed once, under its hyphenated spelling")
ok(parser.parse_args(["ws", "verify", "--in_lang", "ja"]).in_lang == "ja",
   "the underscore spelling still works although it is not listed")
undocumented = []
for act in parser._actions:
    if isinstance(act, argparse._SubParsersAction):
        for name, sp in act.choices.items():
            ok(bool(sp.description), f"{name} describes itself in its help")
            undocumented += [f"{name} {a.option_strings[0]}" for a in sp._actions
                             if a.option_strings and a.help is None]
ok(undocumented == [], f"every option has help text (missing: {undocumented})")

ok(parser.parse_args(["ws", "synthesize", "--ref-wav", "a.wav", "--ref-text", "a.txt"]).ref_text == "a.txt",
   "the reference transcript option is --ref-text, matching --ref-wav")

ok(parser.parse_args(["ws", "synthesize", "--ref-w", "a.wav", "--ref-t", "a.txt"]).ref_wav == "a.wav"
   and parser.parse_args(["ws", "synthesize", "--ref-l", "ja"]).ref_lang == "ja"
   and parser.parse_args(["ws", "pack", deck, "o.pptx", "--data", "text"]).data_type == "text",
   "an option may be abbreviated as far as it stays unambiguous")
expect_error(["ws", "synthesize", "--ref-", "a"], "ambiguous option")

ok(parser.parse_args(["ws", "synthesize", "--dic", "x.csv"]).dict_file == ["x.csv"],
   "an abbreviation is not ambiguous against the option's own underscore spelling")
ok(parser.parse_args(["ws", "synthesize", "--dict-file=z.csv"]).dict_file == ["z.csv"],
   "--option=value works together with the underscore spelling")

print("ALL TESTS PASSED")
