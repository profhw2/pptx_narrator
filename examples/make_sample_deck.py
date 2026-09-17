#!/usr/bin/env python3
"""Build a small sample deck for trying PPTX-Narrator.

    python examples/make_sample_deck.py sample_lecture.pptx

The deck has three slides whose notes contain the kind of material the tool is
meant for: acronyms, gene and chemical names, and number-unit expressions, in
Japanese and in English. It is enough to try extraction, term scanning,
translation, synthesis, verification and packing.
"""
import sys

from pptx import Presentation
from pptx.util import Pt

SLIDES = [
    ("Genome editing",
     "本日は CRISPR-Cas9 によるゲノム編集について説明します。"
     "ヒトのゲノムサイズは約 3 Gbp で、その中の 1 か所を狙って切断します。"),
    ("PCR conditions",
     "PCR は 95 ℃ で 30 秒の熱変性から始めます。反応液量は 20 µL、"
     "プライマー濃度は 0.2 µM とします。"),
    ("mRNA vaccines",
     "This slide summarizes the mRNA vaccine platform. "
     "A lipid nanoparticle of about 100 nm carries the mRNA into the cell."),
]


def main(out_path="sample_lecture.pptx"):
    prs = Presentation()
    layout = prs.slide_layouts[5]  # title only
    for title, note in SLIDES:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        slide.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
        slide.notes_slide.notes_text_frame.text = note
    prs.save(out_path)
    print(f"Wrote {out_path} ({len(SLIDES)} slides with notes)")


if __name__ == "__main__":
    main(*sys.argv[1:])
